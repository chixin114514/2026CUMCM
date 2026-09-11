from __future__ import annotations

import argparse
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
BUILD_DIR = ROOT / "build/deliverables"
BUILD_DIR.mkdir(parents=True, exist_ok=True)


def load_air():
    ws = load_workbook(DATA_DIR / "attachment1.xlsx", data_only=True, read_only=True).active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return (
        np.array([float(r[0]) for r in rows]),
        np.array([float(r[1]) for r in rows]),
        np.array([float(r[2]) for r in rows]),
    )


def load_radius():
    ws = load_workbook(DATA_DIR / "attachment2.xlsx", data_only=True, read_only=True).active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return (
        np.array([float(r[0]) for r in rows]),
        0.01 * np.array([float(r[1]) for r in rows]),
    )


T_DATA, TA_DATA, CA_DATA = load_air()
STABLE = (T_DATA >= 9000.0) & (T_DATA <= 14400.0)
TA_STEADY = float(TA_DATA[STABLE].mean())
CA_STEADY = float(CA_DATA[STABLE].mean())
T_RADIUS, RADIUS_DATA = load_radius()

R0 = 0.02
H = 25.0
HM = 8.0e-7
TARGET = 0.15


def ambient(t):
    if t <= T_DATA[-1]:
        return float(np.interp(t, T_DATA, TA_DATA)), float(np.interp(t, T_DATA, CA_DATA))
    return TA_STEADY, CA_STEADY


def radius(t, shrinking):
    if not shrinking:
        return R0
    if t <= T_RADIUS[-1]:
        return float(np.interp(t, T_RADIUS, RADIUS_DATA))
    return float(RADIUS_DATA[-1])


def thomas(lower, diag, upper, rhs):
    n = len(diag)
    cpv = np.empty(n - 1)
    dpv = np.empty(n)
    cpv[0] = upper[0] / diag[0]
    dpv[0] = rhs[0] / diag[0]
    for i in range(1, n - 1):
        den = diag[i] - lower[i - 1] * cpv[i - 1]
        cpv[i] = upper[i] / den
        dpv[i] = (rhs[i] - lower[i - 1] * dpv[i - 1]) / den
    den = diag[-1] - lower[-1] * cpv[-1]
    dpv[-1] = (rhs[-1] - lower[-1] * dpv[-2]) / den
    ans = np.empty(n)
    ans[-1] = dpv[-1]
    for i in range(n - 2, -1, -1):
        ans[i] = dpv[i] - cpv[i] * ans[i + 1]
    return ans


def harmonic(values):
    return 2.0 * values[:-1] * values[1:] / (values[:-1] + values[1:])


def make_operator(face_property, volumes_x, face_areas_x, capacity, transfer, r_now, dx):
    conductance = face_areas_x * face_property / dx
    physical_volumes = r_now * r_now * volumes_x
    lower = conductance / (capacity[1:] * physical_volumes[1:])
    upper = conductance / (capacity[:-1] * physical_volumes[:-1])
    diag = np.zeros(len(capacity))
    diag[:-1] -= upper
    diag[1:] -= lower
    surface_area = 2.0 * np.pi * r_now
    boundary = surface_area * transfer / (capacity[-1] * physical_volumes[-1])
    diag[-1] -= boundary
    return lower, diag, upper, boundary


def apply_operator(lower, diag, upper, field):
    ans = diag * field
    ans[:-1] += upper * field[1:]
    ans[1:] += lower * field[:-1]
    return ans


def matrix_from_operator(lower, diag, upper, dt):
    return -0.5 * dt * lower, 1.0 - 0.5 * dt * diag, -0.5 * dt * upper


def grid_geometry(n):
    dx = 1.0 / n
    x = np.linspace(0.0, 1.0, n + 1)
    edge = np.empty(n + 2)
    edge[0] = 0.0
    edge[1:-1] = 0.5 * (x[:-1] + x[1:])
    edge[-1] = 1.0
    volumes_x = np.pi * (edge[1:] ** 2 - edge[:-1] ** 2)
    face_areas_x = 2.0 * np.pi * 0.5 * (x[:-1] + x[1:])
    return dx, x, volumes_x, face_areas_x


def sample_fixed(field, x, positions_m):
    return np.interp(positions_m / R0, x, field).tolist()


def sample_shrinking(field, x, r_now, positions_m):
    row = []
    for r in positions_m:
        if r <= r_now + 1.0e-12:
            row.append(float(np.interp(r / r_now, x, field)))
        else:
            row.append(None)
    row.append(float(field[-1]))
    return row


def solve_problem1():
    n = 400
    dt = 0.25
    dx, x, volumes_x, face_areas_x = grid_geometry(n)
    positions = np.arange(0.0, 0.0200001, 0.001)
    output_times = np.arange(1.0, 1800.0 + 1.0, 1.0)
    temp_rows, moisture_rows = [], []

    rho, cp, k = 820.0, 2600.0, 0.36
    alpha = k / (rho * cp)
    temp = np.full(n + 1, 28.0)
    moisture = np.full(n + 1, 2.55)

    lo_t, di_t, up_t, b_t = make_operator(
        np.full(n, alpha), volumes_x, face_areas_x,
        np.ones(n + 1), H / (rho * cp), R0, dx
    )
    mat_t = matrix_from_operator(lo_t, di_t, up_t, dt)

    def diff(c):
        return 7.0e-9 * np.exp(-0.89 / np.maximum(c, 1.0e-12))

    out_index = 0
    steps = int(round(1800.0 / dt))
    for step in range(1, steps + 1):
        t0, t1 = (step - 1) * dt, step * dt
        ta0, ca0 = ambient(t0)
        ta1, ca1 = ambient(t1)

        rhs_t = temp + 0.5 * dt * apply_operator(lo_t, di_t, up_t, temp)
        rhs_t[-1] += 0.5 * dt * b_t * (ta0 + ta1)
        temp = thomas(*mat_t, rhs_t)

        d0 = harmonic(diff(moisture))
        lo0, di0, up0, b0 = make_operator(
            d0, volumes_x, face_areas_x, np.ones(n + 1), HM, R0, dx
        )
        rhs_c = moisture + 0.5 * dt * apply_operator(lo0, di0, up0, moisture)
        rhs_c[-1] += 0.5 * dt * b0 * ca0
        guess = moisture.copy()
        for _ in range(30):
            dg = harmonic(diff(guess))
            lo, di, up, b = make_operator(
                dg, volumes_x, face_areas_x, np.ones(n + 1), HM, R0, dx
            )
            rhs = rhs_c.copy()
            rhs[-1] += 0.5 * dt * b * ca1
            new = thomas(*matrix_from_operator(lo, di, up, dt), rhs)
            if np.max(np.abs(new - guess)) < 1.0e-11:
                guess = new
                break
            guess = 0.7 * new + 0.3 * guess
        moisture = guess

        if out_index < len(output_times) and abs(t1 - output_times[out_index]) < 1.0e-9:
            temp_rows.append(sample_fixed(temp, x, positions))
            moisture_rows.append(sample_fixed(moisture, x, positions))
            out_index += 1

    result = {
        "problem": 1,
        "time_s": output_times.tolist(),
        "radius_cm": (positions * 100.0).tolist(),
        "temperature_C": temp_rows,
        "moisture_kgkg": moisture_rows,
        "solver": {"N": n, "dt_s": dt},
    }
    (BUILD_DIR / "problem1.json").write_text(json.dumps(result), encoding="utf-8")
    print("wrote problem1.json", flush=True)


def p3_rho(c): return 650.0 + 128.0 * c
def p3_cp(c): return 1450.0 + 2736.0 * c / (c + 1.0)
def p3_k(c): return 0.21 + 0.38 * c / (c + 1.0)
def p3_d(t, c): return 2.4e-3 * np.exp(-0.45 / np.maximum(c, 1e-12)) * np.exp(-3850.0 / (t + 273.15))
def p4_rho(c): return 760.0 + 90.0 * c
def p4_cp(c): return 1850.0 + 2150.0 * c / (c + 1.0)
def p4_k(c): return 0.12 + 0.20 * c / (c + 1.0)
def p4_d(t, c): return 4.2e-4 * np.exp(-0.30 / np.maximum(c, 1e-12)) * np.exp(-3850.0 / (t + 273.15))


def solve_coupled(problem, n, dt_pre, dt_long, end_time=None, shrinking=False):
    dx, x, volumes_x, face_areas_x = grid_geometry(n)
    positions = np.arange(0.0, 0.0200001, 0.001)
    rho = p4_rho if problem == 4 else p3_rho
    cp = p4_cp if problem == 4 else p3_cp
    conductivity = p4_k if problem == 4 else p3_k
    diffusivity = p4_d if problem == 4 else p3_d

    if problem == 2:
        output_times = np.arange(1.0, end_time + 1.0, 1.0)
        interval = 1.0
    else:
        output_times = []
        interval = 60.0
    next_output = interval

    temp = np.full(n + 1, 28.0)
    moisture = np.full(n + 1, 2.55)
    t = 0.0
    times, temp_rows, moisture_rows, radius_rows = [], [], [], []
    previous_center = moisture[0]

    limit = end_time if end_time is not None else 240.0 * 3600.0
    while t < limit - 1.0e-12:
        dt = dt_pre if t < 14400.0 else dt_long
        if t < 14400.0 < t + dt:
            dt = 14400.0 - t
        if end_time is not None and t + dt > end_time:
            dt = end_time - t
        t0, t1 = t, t + dt
        r0, r1 = radius(t0, shrinking), radius(t1, shrinking)
        ta0, ca0 = ambient(t0)
        ta1, ca1 = ambient(t1)

        cap0 = rho(moisture) * cp(moisture)
        lo_t0, di_t0, up_t0, b_t0 = make_operator(
            harmonic(conductivity(moisture)), volumes_x, face_areas_x,
            cap0, H, r0, dx
        )
        lo_c0, di_c0, up_c0, b_c0 = make_operator(
            harmonic(diffusivity(temp, moisture)), volumes_x, face_areas_x,
            np.ones(n + 1), HM, r0, dx
        )
        rhs_t0 = temp + 0.5 * dt * apply_operator(lo_t0, di_t0, up_t0, temp)
        rhs_t0[-1] += 0.5 * dt * b_t0 * ta0
        rhs_c0 = moisture + 0.5 * dt * apply_operator(lo_c0, di_c0, up_c0, moisture)
        rhs_c0[-1] += 0.5 * dt * b_c0 * ca0

        temp_guess = temp.copy()
        moisture_guess = moisture.copy()
        for iteration in range(1, 41):
            cap1 = rho(moisture_guess) * cp(moisture_guess)
            lo_t, di_t, up_t, b_t = make_operator(
                harmonic(conductivity(moisture_guess)), volumes_x, face_areas_x,
                cap1, H, r1, dx
            )
            rhs_t = rhs_t0.copy()
            rhs_t[-1] += 0.5 * dt * b_t * ta1
            temp_new = thomas(*matrix_from_operator(lo_t, di_t, up_t, dt), rhs_t)

            lo_c, di_c, up_c, b_c = make_operator(
                harmonic(diffusivity(temp_new, moisture_guess)), volumes_x,
                face_areas_x, np.ones(n + 1), HM, r1, dx
            )
            rhs_c = rhs_c0.copy()
            rhs_c[-1] += 0.5 * dt * b_c * ca1
            moisture_new = thomas(*matrix_from_operator(lo_c, di_c, up_c, dt), rhs_c)
            err_t = float(np.max(np.abs(temp_new - temp_guess)))
            err_c = float(np.max(np.abs(moisture_new - moisture_guess)))
            temp_guess, moisture_guess = temp_new, moisture_new
            if err_t < 1.0e-9 and err_c < 1.0e-11:
                break
        else:
            raise RuntimeError(f"problem {problem}: Picard failed at t={t1}")

        previous_t, previous_temp, previous_moisture = t, temp, moisture
        previous_center = moisture[0]
        t, temp, moisture = t1, temp_new, moisture_new

        crossed = end_time is None and moisture[0] <= TARGET
        if crossed:
            crossing_fraction = (previous_center - TARGET) / (previous_center - moisture[0])
            crossing_time = previous_t + crossing_fraction * (t - previous_t)
            output_limit = crossing_time
        else:
            output_limit = t

        while next_output <= output_limit + 1.0e-9:
            f = (next_output - previous_t) / (t - previous_t)
            temp_out = previous_temp + f * (temp - previous_temp)
            moisture_out = previous_moisture + f * (moisture - previous_moisture)
            r_out = radius(next_output, shrinking)
            times.append(float(next_output))
            if shrinking:
                moisture_rows.append(sample_shrinking(moisture_out, x, r_out, positions))
                radius_rows.append(r_out * 100.0)
            else:
                temp_rows.append(sample_fixed(temp_out, x, positions))
                moisture_rows.append(sample_fixed(moisture_out, x, positions))
            next_output += interval

        if crossed:
            f = crossing_fraction
            end_t = crossing_time
            end_temp = previous_temp + f * (temp - previous_temp)
            end_moisture = previous_moisture + f * (moisture - previous_moisture)
            end_moisture[0] = TARGET
            if not times or end_t - times[-1] > 1.0e-7:
                times.append(float(end_t))
                if shrinking:
                    end_r = radius(end_t, True)
                    moisture_rows.append(sample_shrinking(end_moisture, x, end_r, positions))
                    radius_rows.append(end_r * 100.0)
                else:
                    temp_rows.append(sample_fixed(end_temp, x, positions))
                    moisture_rows.append(sample_fixed(end_moisture, x, positions))
            break

    result = {
        "problem": problem,
        "time_s": times,
        "radius_cm": (positions * 100.0).tolist(),
        "moisture_kgkg": moisture_rows,
        "solver": {"N": n, "dt_pre_s": dt_pre, "dt_long_s": dt_long},
    }
    if temp_rows:
        result["temperature_C"] = temp_rows
    if radius_rows:
        result["current_radius_cm"] = radius_rows
    if end_time is None:
        result["end_time_s"] = times[-1]
        result["end_time_h"] = times[-1] / 3600.0
    path = BUILD_DIR / f"problem{problem}.json"
    path.write_text(json.dumps(result), encoding="utf-8")
    print(f"wrote {path.name}: {len(times)} rows", flush=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("problem", type=int, choices=[1, 2, 3, 4])
    args = parser.parse_args()
    if args.problem == 1:
        solve_problem1()
    elif args.problem == 2:
        solve_coupled(2, n=400, dt_pre=0.5, dt_long=0.5, end_time=10800.0)
    elif args.problem == 3:
        solve_coupled(3, n=1600, dt_pre=2.0, dt_long=10.0)
    else:
        solve_coupled(4, n=800, dt_pre=2.0, dt_long=10.0, shrinking=True)


if __name__ == "__main__":
    main()
