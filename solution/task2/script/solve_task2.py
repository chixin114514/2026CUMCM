"""A题问题二：动态物性圆柱径向热质耦合求解。

本脚本严格实现 ``handoff/task2_handoff.md`` 规定的守恒型径向有限体积
模型：空间界面物性采用严格正值调和平均，时间方向采用 Crank--Nicolson，
并在每个时间步使用同一半时间层状态进行同步 Picard 迭代。内部网格含
``r=0`` 和 ``r=R`` 两个节点、两端为半控制体积；正式输出为每 1 s、每
0.1 cm。
"""

from __future__ import annotations

import argparse
import csv
import time as wall_time
from copy import copy
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.linalg import solve_banded


# 几何、初始状态和输出约定（内部长度统一使用 m）。
RADIUS_M = 0.02
LENGTH_M = 0.25  # 一维径向守恒式中约去，保留作题目参数记录。
# handoff/task2_handoff.md 的正式基准：N=1600 个径向区间、1601 个节点。
BASE_DR_M = 1.25e-5
INITIAL_TEMPERATURE_C = 28.0
INITIAL_MOISTURE_KGKG = 2.55
FINAL_TIME_S = 10800.0
OUTPUT_DT_S = 1.0
OUTPUT_RADIUS_CM = np.arange(0.0, 2.0000001, 0.1)

# 问题一沿用的边界传递参数。
HEAT_TRANSFER_W_M2_K = 25.0
MASS_TRANSFER_M_S = 8.0e-7

# Picard 收敛参数。绝对阈值足以保证四位小数输出稳定，同时避免无意义地
# 将迭代压到机器精度。
PICARD_TOL_T_C = 1.0e-8
PICARD_TOL_C_KGKG = 1.0e-10
PICARD_MAX_ITER = 50
LINEAR_RESIDUAL_TOL = 1.0e-9

DEFAULT_DATA_PATH = (
    Path(__file__).resolve().parents[3] / "problem" / "附件" / "附件1.xlsx"
)
DEFAULT_TEMPLATE_PATH = (
    Path(__file__).resolve().parents[3]
    / "problem"
    / "附件"
    / "附件3"
    / "result2.xlsx"
)
DEFAULT_XLSX_PATH = Path(__file__).resolve().parents[1] / "answer" / "result2.xlsx"
DEFAULT_CSV_PATH = (
    Path(__file__).resolve().parents[1] / "answer" / "task2_all_results.csv"
)
DEFAULT_TABLE3_PATH = (
    Path(__file__).resolve().parents[1] / "answer" / "table3_temperature.csv"
)
DEFAULT_TABLE4_PATH = (
    Path(__file__).resolve().parents[1] / "answer" / "table4_moisture.csv"
)
DEFAULT_DIAGNOSTICS_JSON = (
    Path(__file__).resolve().parents[1] / "answer" / "task2_diagnostics.json"
)
DEFAULT_DIAGNOSTICS_CSV = (
    Path(__file__).resolve().parents[1] / "answer" / "task2_diagnostics.csv"
)


def load_air_data(
    path: str | Path = DEFAULT_DATA_PATH,
    required_end_s: float = FINAL_TIME_S,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """读取附件1的全部烘房边界数据，供求解区间内分段线性插值。"""

    frame = pd.read_excel(Path(path))
    column_map = {str(column).strip(): column for column in frame.columns}
    required = ("时间", "温度", "水分浓度")
    missing = [column for column in required if column not in column_map]
    if missing:
        raise ValueError(f"附件1缺少必需列: {missing}; 实际列为 {list(frame.columns)!r}")

    selected = frame[[column_map[column] for column in required]].copy()
    selected.columns = required
    for column in required:
        selected[column] = pd.to_numeric(selected[column], errors="raise")
    selected = selected.sort_values("时间").drop_duplicates("时间", keep="last")

    time_s = selected["时间"].to_numpy(dtype=float)
    air_temperature_C = selected["温度"].to_numpy(dtype=float)
    air_moisture_kgkg = selected["水分浓度"].to_numpy(dtype=float)
    if time_s.size < 2 or np.any(np.diff(time_s) <= 0.0):
        raise ValueError("附件1时间必须包含至少两个严格递增时间点")
    if time_s[0] > 1.0e-10 or time_s[-1] < required_end_s - 1.0e-10:
        raise ValueError(f"附件1烘房边界数据未覆盖 0--{required_end_s:g} s")
    if not np.all(np.isfinite(np.column_stack((time_s, air_temperature_C, air_moisture_kgkg)))):
        raise ValueError("附件1烘房边界数据包含非有限数值")
    return time_s, air_temperature_C, air_moisture_kgkg


def control_volume_geometry(
    radius_m: float = RADIUS_M, dr_m: float = BASE_DR_M
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """返回含两端半控制体积的节点、面和径向测度。"""

    if radius_m <= 0.0 or dr_m <= 0.0:
        raise ValueError("半径和空间步长必须为正")
    n_intervals = int(round(radius_m / dr_m))
    if not np.isclose(n_intervals * dr_m, radius_m, rtol=0.0, atol=1.0e-12):
        raise ValueError("当前网格要求半径可由整数个 dr 划分")
    radii_m = np.linspace(0.0, radius_m, n_intervals + 1)
    faces_m = 0.5 * (radii_m[:-1] + radii_m[1:])
    radial_measure_m2 = np.empty_like(radii_m)
    radial_measure_m2[0] = 0.5 * faces_m[0] ** 2
    radial_measure_m2[1:-1] = 0.5 * (faces_m[1:] ** 2 - faces_m[:-1] ** 2)
    radial_measure_m2[-1] = 0.5 * (radius_m**2 - faces_m[-1] ** 2)
    return radii_m, faces_m, radial_measure_m2


def material_properties(
    temperature_C: np.ndarray, moisture_kgkg: np.ndarray
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """按当前 Picard 状态计算 rho、cp、k 和 D（SI 单位）。"""

    temperature_C = np.asarray(temperature_C, dtype=float)
    moisture_kgkg = np.asarray(moisture_kgkg, dtype=float)
    if (
        not np.all(np.isfinite(temperature_C))
        or not np.all(np.isfinite(moisture_kgkg))
        or np.any(moisture_kgkg <= 0.0)
    ):
        raise ValueError("动态物性计算要求温度和水分为有限值且水分严格为正")
    concentration = moisture_kgkg
    temperature_K = temperature_C + 273.15
    if np.any(temperature_K <= 0.0):
        raise ValueError("动态物性计算遇到非正 Kelvin 温度")

    rho = 650.0 + 128.0 * concentration
    cp = 1450.0 + 2736.0 * concentration / (concentration + 1.0)
    k = 0.21 + 0.38 * concentration / (concentration + 1.0)
    diffusivity = 2.4e-3 * np.exp(-0.45 / concentration) * np.exp(
        -3850.0 / temperature_K
    )
    return rho, cp, k, diffusivity


def _harmonic_mean_positive(
    left: np.ndarray, right: np.ndarray, name: str
) -> np.ndarray:
    """计算严格正值数组的界面调和平均。"""

    if (
        np.any(~np.isfinite(left))
        or np.any(~np.isfinite(right))
        or np.any(left <= 0.0)
        or np.any(right <= 0.0)
    ):
        raise ValueError(f"{name} 界面调和平均要求两侧物性均为有限正值")
    result = 2.0 * left * right / (left + right)
    if np.any(~np.isfinite(result)) or np.any(result <= 0.0):
        raise ValueError(f"{name} 界面调和平均未得到严格正值")
    return result


def _solve_tridiagonal(
    lower: np.ndarray, diagonal: np.ndarray, upper: np.ndarray, rhs: np.ndarray
) -> tuple[np.ndarray, float]:
    """解三对角线性系统，矩阵系数均为当前 Picard 冻结值。"""

    n = diagonal.size
    banded = np.zeros((3, n), dtype=float)
    banded[0, 1:] = upper
    banded[1, :] = diagonal
    banded[2, :-1] = lower
    solution = solve_banded((1, 1), banded, rhs, check_finite=False)
    residual = diagonal * solution - rhs
    if n > 1:
        residual[:-1] += upper * solution[1:]
        residual[1:] += lower * solution[:-1]
    scale = max(
        float(np.max(np.abs(rhs))),
        float(np.max(np.abs(diagonal * solution))),
        float(np.max(np.abs(upper * solution[1:]))) if n > 1 else 0.0,
        float(np.max(np.abs(lower * solution[:-1]))) if n > 1 else 0.0,
        1.0e-30,
    )
    normalized_residual = float(np.max(np.abs(residual)) / scale)
    if not np.isfinite(normalized_residual) or normalized_residual > LINEAR_RESIDUAL_TOL:
        raise RuntimeError(
            "三对角线性系统残差异常: "
            f"normalized_inf_residual={normalized_residual:.3e}, "
            f"tol={LINEAR_RESIDUAL_TOL:.3e}"
        )
    return solution, normalized_residual


def _solve_cn_field(
    old_field: np.ndarray,
    storage: np.ndarray,
    face_property: np.ndarray,
    faces_m: np.ndarray,
    dr_m: float,
    dt_s: float,
    boundary_conductance: float,
    air_value: float,
) -> tuple[np.ndarray, float]:
    """用冻结的半时间层物性求解一个守恒 FVM 的 CN 场。

    令 ``K`` 为正的扩散/Robin 离散矩阵，空间算子为 ``-K u+b``。
    对 ``S (u_new-u_old)/dt = (-K u+b)`` 使用 CN 后，线性系统为

    ``(S+dt*K/2)u_new=(S-dt*K/2)u_old+dt*b``。

    其中 ``air_value`` 已由调用方取为当前时间步的半时间层环境值。
    """

    old_field = np.asarray(old_field, dtype=float)
    storage = np.asarray(storage, dtype=float)
    face_property = np.asarray(face_property, dtype=float)
    faces_m = np.asarray(faces_m, dtype=float)
    if old_field.ndim != 1 or storage.shape != old_field.shape:
        raise ValueError("CN 场和储存项的数组形状不一致")
    if face_property.size != old_field.size - 1 or faces_m.shape != face_property.shape:
        raise ValueError("CN 界面物性或界面半径数组长度错误")
    if dt_s <= 0.0 or dr_m <= 0.0 or boundary_conductance < 0.0:
        raise ValueError("CN 时间步、空间步和边界传递系数必须有效")
    if (
        np.any(~np.isfinite(old_field))
        or np.any(~np.isfinite(storage))
        or np.any(~np.isfinite(face_property))
        or np.any(storage <= 0.0)
    ):
        raise ValueError("CN 线性系统输入包含非有限值或非正储存项")

    conductance = face_property * faces_m / dr_m
    if (
        np.any(~np.isfinite(conductance))
        or np.any(conductance <= 0.0)
        or not np.isfinite(air_value)
    ):
        raise ValueError("CN 界面通量系数必须为有限正值")

    # K 的对角线：中心和表面均只含一个内部界面，表面再叠加 Robin 项。
    diagonal_K = np.empty_like(storage)
    diagonal_K[0] = conductance[0]
    diagonal_K[1:-1] = conductance[:-1] + conductance[1:]
    diagonal_K[-1] = conductance[-1] + boundary_conductance

    half_dt = 0.5 * dt_s
    diagonal = storage + half_dt * diagonal_K
    lower = -half_dt * conductance
    upper = -half_dt * conductance

    # (S-dt*K/2)u_old 的三对角乘积；表面 Robin 源项在两端环境值相同，
    # 因而 CN 的旧/新半权重合并为 dt*b_mid。
    rhs = (storage - half_dt * diagonal_K) * old_field
    rhs[:-1] += half_dt * conductance * old_field[1:]
    rhs[1:] += half_dt * conductance * old_field[:-1]
    rhs[-1] += dt_s * boundary_conductance * air_value
    return _solve_tridiagonal(lower, diagonal, upper, rhs)


def _picard_step(
    old_temperature_C: np.ndarray,
    old_moisture_kgkg: np.ndarray,
    time_new_s: float,
    dt_s: float,
    radial_measure_m2: np.ndarray,
    faces_m: np.ndarray,
    dr_m: float,
    air_temperature_C: float,
    air_moisture_kgkg: float,
    tol_temperature_C: float = PICARD_TOL_T_C,
    tol_moisture_kgkg: float = PICARD_TOL_C_KGKG,
    max_iter: int = PICARD_MAX_ITER,
) -> tuple[np.ndarray, np.ndarray, int, float, float, float]:
    """使用当前时间步内的同步半时间层 Picard 外迭代求 T、C 自洽解。"""

    # time_new_s 同时用于异常信息，明确边界取值属于当前时间步。
    guess_temperature_C = old_temperature_C.copy()
    guess_moisture_kgkg = old_moisture_kgkg.copy()
    last_delta_temperature = np.inf
    last_delta_moisture = np.inf
    max_linear_residual = 0.0

    for iteration in range(1, max_iter + 1):
        # 两个场严格使用同一组半时间层状态，避免由求解顺序造成耦合偏差。
        midpoint_temperature_C = 0.5 * (
            old_temperature_C + guess_temperature_C
        )
        midpoint_moisture_kgkg = 0.5 * (
            old_moisture_kgkg + guess_moisture_kgkg
        )
        rho, cp, conductivity, diffusivity = material_properties(
            midpoint_temperature_C, midpoint_moisture_kgkg
        )
        face_k = _harmonic_mean_positive(
            conductivity[:-1], conductivity[1:], "导热系数"
        )
        face_D = _harmonic_mean_positive(
            diffusivity[:-1], diffusivity[1:], "扩散系数"
        )

        new_temperature_C, temperature_residual = _solve_cn_field(
            old_temperature_C,
            rho * cp * radial_measure_m2,
            face_k,
            faces_m,
            dr_m,
            dt_s,
            RADIUS_M * HEAT_TRANSFER_W_M2_K,
            air_temperature_C,
        )
        new_moisture_kgkg, moisture_residual = _solve_cn_field(
            old_moisture_kgkg,
            radial_measure_m2,
            face_D,
            faces_m,
            dr_m,
            dt_s,
            RADIUS_M * MASS_TRANSFER_M_S,
            air_moisture_kgkg,
        )
        max_linear_residual = max(
            max_linear_residual, temperature_residual, moisture_residual
        )
        last_delta_temperature = float(
            np.max(np.abs(new_temperature_C - guess_temperature_C))
        )
        last_delta_moisture = float(
            np.max(np.abs(new_moisture_kgkg - guess_moisture_kgkg))
        )
        if not np.all(np.isfinite(new_moisture_kgkg)) or np.min(
            new_moisture_kgkg
        ) <= 0.0:
            raise RuntimeError("CN Picard 步得到非正或非有限水分浓度")
        if (
            last_delta_temperature <= tol_temperature_C
            and last_delta_moisture <= tol_moisture_kgkg
        ):
            return (
                new_temperature_C,
                new_moisture_kgkg,
                iteration,
                last_delta_temperature,
                last_delta_moisture,
                max_linear_residual,
            )
        guess_temperature_C = new_temperature_C
        guess_moisture_kgkg = new_moisture_kgkg

    raise RuntimeError(
        "问题二 Picard 迭代未收敛: "
        f"t={time_new_s:g}s, iter={max_iter}, "
        f"max|dT|={last_delta_temperature:.3e}, "
        f"max|dC|={last_delta_moisture:.3e}"
    )


def _output_indices(radii_m: np.ndarray) -> np.ndarray:
    """取得 0--2 cm 每 0.1 cm 输出位置对应的内部节点。"""

    target_m = OUTPUT_RADIUS_CM / 100.0
    indices = np.array([int(np.argmin(np.abs(radii_m - r))) for r in target_m])
    if not np.allclose(radii_m[indices], target_m, atol=1.0e-12, rtol=0.0):
        raise ValueError("当前内部网格不能精确包含题目要求的 0.1 cm 输出位置")
    return indices


def solve_task2(
    data_path: str | Path = DEFAULT_DATA_PATH,
    t_end_s: float = FINAL_TIME_S,
    time_step_s: float = 1.0,
    output_dt_s: float = OUTPUT_DT_S,
    dr_m: float = BASE_DR_M,
    progress: bool = False,
) -> dict[str, Any]:
    """求解问题二并返回输出网格上的完整数值场。"""

    if t_end_s <= 0.0 or time_step_s <= 0.0 or output_dt_s <= 0.0:
        raise ValueError("终止时间、时间步和输出间隔必须为正")
    n_steps_float = t_end_s / time_step_s
    n_outputs_float = t_end_s / output_dt_s
    n_steps = int(round(n_steps_float))
    n_outputs = int(round(n_outputs_float))
    if not np.isclose(n_steps_float, n_steps, atol=1.0e-10):
        raise ValueError("终止时间必须由整数个求解时间步组成")
    if not np.isclose(n_outputs_float, n_outputs, atol=1.0e-10):
        raise ValueError("终止时间必须由整数个输出间隔组成")
    ratio = output_dt_s / time_step_s
    output_stride = int(round(ratio))
    if not np.isclose(ratio, output_stride, atol=1.0e-10):
        raise ValueError("为避免额外插值，output_dt_s 必须是 time_step_s 的整数倍")

    air_time_s, air_temperature, air_moisture = load_air_data(
        data_path, required_end_s=t_end_s
    )
    radii_m, faces_m, radial_measure_m2 = control_volume_geometry(RADIUS_M, dr_m)
    output_indices = _output_indices(radii_m)
    n_nodes = radii_m.size
    n_outputs_total = n_outputs + 1
    output_times = output_dt_s * np.arange(n_outputs_total, dtype=float)
    temperature_output = np.empty((n_outputs_total, output_indices.size), dtype=float)
    moisture_output = np.empty_like(temperature_output)
    air_temperature_output = np.interp(output_times, air_time_s, air_temperature)
    air_moisture_output = np.interp(output_times, air_time_s, air_moisture)
    initial_temperature = np.full(n_nodes, INITIAL_TEMPERATURE_C, dtype=float)
    initial_moisture = np.full(n_nodes, INITIAL_MOISTURE_KGKG, dtype=float)
    current_temperature = initial_temperature.copy()
    current_moisture = initial_moisture.copy()
    temperature_output[0] = current_temperature[output_indices]
    moisture_output[0] = current_moisture[output_indices]
    iteration_counts = np.empty(n_steps, dtype=int)
    picard_delta_temperature = np.empty(n_steps, dtype=float)
    picard_delta_moisture = np.empty(n_steps, dtype=float)
    linear_residuals = np.empty(n_steps, dtype=float)
    moisture_balance_residuals = np.empty(n_steps, dtype=float)
    energy_balance_residuals = np.empty(n_steps, dtype=float)
    moisture_boundary_fluxes = np.empty(n_steps, dtype=float)
    energy_boundary_fluxes = np.empty(n_steps, dtype=float)
    air_temperature_interp = lambda t: float(np.interp(t, air_time_s, air_temperature))
    air_moisture_interp = lambda t: float(np.interp(t, air_time_s, air_moisture))

    for step in range(1, n_steps + 1):
        time_new_s = step * time_step_s
        time_mid_s = time_new_s - 0.5 * time_step_s
        previous_temperature = current_temperature
        previous_moisture = current_moisture
        air_temperature_mid = air_temperature_interp(time_mid_s)
        air_moisture_mid = air_moisture_interp(time_mid_s)
        (
            current_temperature,
            current_moisture,
            iteration_counts[step - 1],
            picard_delta_temperature[step - 1],
            picard_delta_moisture[step - 1],
            linear_residuals[step - 1],
        ) = _picard_step(
            current_temperature,
            current_moisture,
            time_new_s,
            time_step_s,
            radial_measure_m2,
            faces_m,
            dr_m,
            air_temperature_mid,
            air_moisture_mid,
        )
        # 质量守恒：约去 2πL 后，全域水分储量变化应等于 CN 加权的表面质流。
        moisture_flux = RADIUS_M * MASS_TRANSFER_M_S * (
            air_moisture_mid
            - 0.5 * (previous_moisture[-1] + current_moisture[-1])
        )
        moisture_boundary_fluxes[step - 1] = moisture_flux
        moisture_balance_residuals[step - 1] = float(
            radial_measure_m2 @ (current_moisture - previous_moisture)
        ) / time_step_s - moisture_flux
        # 能量守恒：按同一半时间层物性冻结储存项，全局焓变应等于表面热流。
        midpoint_temperature = 0.5 * (previous_temperature + current_temperature)
        midpoint_moisture = 0.5 * (previous_moisture + current_moisture)
        rho_mid, cp_mid, _, _ = material_properties(midpoint_temperature, midpoint_moisture)
        energy_flux = RADIUS_M * HEAT_TRANSFER_W_M2_K * (
            air_temperature_mid
            - 0.5 * (previous_temperature[-1] + current_temperature[-1])
        )
        energy_boundary_fluxes[step - 1] = energy_flux
        energy_balance_residuals[step - 1] = float(
            (rho_mid * cp_mid * radial_measure_m2) @ (current_temperature - previous_temperature)
        ) / time_step_s - energy_flux
        if step % output_stride == 0:
            output_index = step // output_stride
            temperature_output[output_index] = current_temperature[output_indices]
            moisture_output[output_index] = current_moisture[output_indices]
        if progress and (step == 1 or step == n_steps or step % max(1, n_steps // 10) == 0):
            print(f"  已完成 {time_new_s:.1f}/{t_end_s:.1f} s")

    return {
        "time_s": output_times,
        "radius_cm": OUTPUT_RADIUS_CM.copy(),
        "temperature_C": temperature_output,
        "moisture_kgkg": moisture_output,
        "air_temperature_C": air_temperature_output,
        "air_moisture_kgkg": air_moisture_output,
        "radius_m_internal": radii_m,
        "time_step_s": float(time_step_s),
        "dr_m": float(dr_m),
        "iteration_counts": iteration_counts,
        "picard_delta_temperature": picard_delta_temperature,
        "picard_delta_moisture": picard_delta_moisture,
        "linear_residuals": linear_residuals,
        "moisture_balance_residuals": moisture_balance_residuals,
        "energy_balance_residuals": energy_balance_residuals,
        "moisture_boundary_fluxes": moisture_boundary_fluxes,
        "energy_boundary_fluxes": energy_boundary_fluxes,
        "max_linear_residual": float(np.max(linear_residuals)) if n_steps else 0.0,
        "center_inner_face_radius_m": float(0.0),
        "center_inner_boundary_coefficient": float(0.0),
        "air_temperature_last_C": air_temperature_interp(t_end_s),
        "air_moisture_last_kgkg": air_moisture_interp(t_end_s),
    }


def validate_solution(solution: dict[str, Any]) -> dict[str, Any]:
    """执行最低限度的数值和物理 sanity check，并返回摘要。"""

    time_s = np.asarray(solution["time_s"], dtype=float)
    radius_cm = np.asarray(solution["radius_cm"], dtype=float)
    temperature = np.asarray(solution["temperature_C"], dtype=float)
    moisture = np.asarray(solution["moisture_kgkg"], dtype=float)
    air_temperature = np.asarray(solution["air_temperature_C"], dtype=float)
    air_moisture = np.asarray(solution["air_moisture_kgkg"], dtype=float)
    if time_s.size != 10801 or radius_cm.size != 21:
        raise ValueError("完整结果必须为 10801 个时间点和 21 个半径点")
    if temperature.shape != (time_s.size, radius_cm.size):
        raise ValueError("温度输出数组形状错误")
    if moisture.shape != temperature.shape:
        raise ValueError("温度和水分输出数组形状不一致")
    if air_temperature.shape != time_s.shape or air_moisture.shape != time_s.shape:
        raise ValueError("烘房边界输出数组形状错误")
    if not np.all(np.isfinite(np.column_stack((temperature.ravel(), moisture.ravel())))):
        raise ValueError("结果包含 NaN 或 Inf")
    if not np.all(np.isfinite(np.column_stack((air_temperature, air_moisture)))):
        raise ValueError("烘房边界结果包含 NaN 或 Inf")
    if np.min(moisture) < -1.0e-10:
        raise ValueError(f"结果包含负水分浓度: min={np.min(moisture):.6g}")
    if not np.allclose(temperature[0], INITIAL_TEMPERATURE_C, atol=1.0e-12):
        raise ValueError("初始全径向温度不符合题设")
    if not np.allclose(moisture[0], INITIAL_MOISTURE_KGKG, atol=1.0e-12):
        raise ValueError("初始全径向水分不符合题设")
    if not np.isclose(time_s[0], 0.0, atol=1.0e-12) or not np.isclose(
        time_s[-1], FINAL_TIME_S, atol=1.0e-10
    ):
        raise ValueError("结果时间范围必须完整覆盖 0--10800 s")
    if not np.all(np.diff(time_s) > 0.0):
        raise ValueError("输出时间没有严格递增")
    if not np.allclose(np.diff(time_s), OUTPUT_DT_S, atol=1.0e-10):
        raise ValueError("输出时间间隔不是 1 s")
    if not np.isclose(radius_cm[0], 0.0) or not np.isclose(radius_cm[-1], 2.0):
        raise ValueError("输出半径范围不符合 0--2 cm")

    final_temperature = temperature[-1]
    final_moisture = moisture[-1]
    if final_temperature[-1] <= final_temperature[0]:
        raise ValueError("末时刻表面温度未高于中心，检查热边界符号")
    if final_moisture[-1] >= final_moisture[0]:
        raise ValueError("末时刻表面水分未低于中心，检查传质边界符号")
    if np.max(temperature) > max(temperature[-1, -1], 60.0) + 20.0:
        raise ValueError("温度出现明显数值爆炸")
    if np.max(np.abs(np.diff(temperature, axis=0))) > 5.0:
        raise ValueError("相邻输出时刻温度出现明显跳变")
    if np.max(np.abs(np.diff(moisture, axis=0))) > 0.2:
        raise ValueError("相邻输出时刻水分出现明显跳变")

    heat_boundary_driver = air_temperature - temperature[:, -1]
    moisture_boundary_driver = air_moisture - moisture[:, -1]
    heat_inward_fraction = float(np.mean(heat_boundary_driver[1:] >= -1.0e-8))
    moisture_outward_fraction = float(np.mean(moisture_boundary_driver[1:] <= 1.0e-8))
    if heat_boundary_driver[1] < -1.0e-8 or heat_boundary_driver[-1] < -1.0e-8:
        raise ValueError("表面热边界方向异常：起始或末时刻未体现烘房向药材传热")
    if moisture_boundary_driver[1] > 1.0e-8 or moisture_boundary_driver[-1] > 1.0e-8:
        raise ValueError("表面传质边界方向异常：起始或末时刻未体现药材向外失水")
    if heat_inward_fraction < 0.95 or moisture_outward_fraction < 0.95:
        raise ValueError("表面边界方向异常：主导时段未体现热进入、水分流出")

    iterations = np.asarray(solution["iteration_counts"], dtype=int)
    if iterations.size and np.max(iterations) > PICARD_MAX_ITER:
        raise ValueError("Picard 迭代次数超出上限")
    picard_delta_temperature = np.asarray(
        solution["picard_delta_temperature"], dtype=float
    )
    picard_delta_moisture = np.asarray(solution["picard_delta_moisture"], dtype=float)
    linear_residuals = np.asarray(solution["linear_residuals"], dtype=float)
    if (
        picard_delta_temperature.shape != iterations.shape
        or picard_delta_moisture.shape != iterations.shape
        or linear_residuals.shape != iterations.shape
    ):
        raise ValueError("Picard 或线性残差记录长度错误")
    if not np.all(np.isfinite(linear_residuals)) or np.max(linear_residuals) > LINEAR_RESIDUAL_TOL:
        raise ValueError("三对角线性系统残差超过容差")
    if np.any(picard_delta_temperature > PICARD_TOL_T_C) or np.any(
        picard_delta_moisture > PICARD_TOL_C_KGKG
    ):
        raise ValueError("存在未达到收敛阈值的 Picard 时间步")
    center_face_radius = float(solution["center_inner_face_radius_m"])
    center_boundary_coefficient = float(solution["center_inner_boundary_coefficient"])
    if not np.isclose(center_face_radius, 0.0, atol=1.0e-15) or not np.isclose(
        center_boundary_coefficient, 0.0, atol=1.0e-30
    ):
        raise ValueError("中心内侧控制面未体现零通量离散")
    return {
        "time_end_s": float(time_s[-1]),
        "temperature_min_C": float(np.min(temperature)),
        "temperature_max_C": float(np.max(temperature)),
        "moisture_min_kgkg": float(np.min(moisture)),
        "moisture_max_kgkg": float(np.max(moisture)),
        "max_picard_iterations": int(np.max(iterations)) if iterations.size else 0,
        "mean_picard_iterations": float(np.mean(iterations)) if iterations.size else 0.0,
        "max_linear_residual": float(np.max(linear_residuals))
        if linear_residuals.size
        else 0.0,
        "center_surface_temperature_gap_C": float(final_temperature[-1] - final_temperature[0]),
        "center_surface_moisture_gap_kgkg": float(final_moisture[0] - final_moisture[-1]),
        "center_inner_face_radius_m": center_face_radius,
        "center_inner_boundary_coefficient": center_boundary_coefficient,
        "center_flux_zero_discretisation": bool(
            np.isclose(center_face_radius, 0.0, atol=1.0e-15)
            and np.isclose(center_boundary_coefficient, 0.0, atol=1.0e-30)
        ),
        "minimum_heat_boundary_driver_C": float(np.min(heat_boundary_driver[1:])),
        "heat_inward_fraction": heat_inward_fraction,
        "maximum_moisture_boundary_driver_kgkg": float(
            np.max(moisture_boundary_driver[1:])
        ),
        "moisture_outward_fraction": moisture_outward_fraction,
    }


def write_result_xlsx(
    solution: dict[str, Any],
    output_path: str | Path = DEFAULT_XLSX_PATH,
    template_path: str | Path = DEFAULT_TEMPLATE_PATH,
) -> None:
    """以 result2.xlsx 为模板，写入两张完整结果表。"""

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(Path(template_path))
    required_sheets = ("温度", "水分浓度")
    if tuple(workbook.sheetnames) != required_sheets:
        raise ValueError(f"result2.xlsx 模板工作表不符合要求: {workbook.sheetnames}")

    time_s = np.asarray(solution["time_s"], dtype=float)
    radius_cm = np.asarray(solution["radius_cm"], dtype=float)
    temperature = np.asarray(solution["temperature_C"], dtype=float)
    moisture = np.asarray(solution["moisture_kgkg"], dtype=float)
    if time_s.size != 10801 or radius_cm.size != 21:
        raise ValueError("最终 xlsx 必须为 10801 个时间点和 21 个半径点")

    for sheet_name, values in (("温度", temperature), ("水分浓度", moisture)):
        worksheet = workbook[sheet_name]
        # 模板的两张表字体略有不同；在删除示例数据行前分别保存三类
        # 参考样式，确保扩展出的完整结果仍与对应模板一致。
        header_style = copy(worksheet.cell(row=1, column=2)._style)
        time_style = copy(worksheet.cell(row=2, column=1)._style)
        value_style = copy(worksheet.cell(row=2, column=2)._style)
        if worksheet.max_row > 1:
            worksheet.delete_rows(2, worksheet.max_row - 1)
        # 保留模板左上角的文字；距离表头按题目要求改为完整的 0--2 cm。
        for column_index, radius in enumerate(radius_cm, start=2):
            cell = worksheet.cell(
                row=1, column=column_index, value=round(float(radius), 1)
            )
            cell._style = copy(header_style)
            cell.number_format = "0.0"
        for row_index, time_value in enumerate(time_s, start=2):
            time_cell = worksheet.cell(row=row_index, column=1, value=int(round(time_value)))
            time_cell._style = copy(time_style)
            time_cell.number_format = "0"
            for column_index, value in enumerate(values[row_index - 2], start=2):
                cell = worksheet.cell(
                    row=row_index, column=column_index, value=float(np.round(value, 4))
                )
                cell._style = copy(value_style)
                cell.number_format = "0.0000"
    workbook.save(output_path)


def write_result_csv(
    solution: dict[str, Any], output_path: str | Path = DEFAULT_CSV_PATH
) -> None:
    """写入问题二全部时间/半径结果的 long-format CSV。"""

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    time_s = np.asarray(solution["time_s"], dtype=float)
    radius_cm = np.asarray(solution["radius_cm"], dtype=float)
    temperature = np.asarray(solution["temperature_C"], dtype=float)
    moisture = np.asarray(solution["moisture_kgkg"], dtype=float)
    with output_path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            ["time_s", "time_h", "radius_cm", "temperature_C", "moisture_kgkg"]
        )
        for time_index, time_value in enumerate(time_s):
            for radius_index, radius_value in enumerate(radius_cm):
                writer.writerow(
                    [
                        int(round(time_value)),
                        f"{time_value / 3600.0:.6f}",
                        f"{radius_value:.1f}",
                        f"{temperature[time_index, radius_index]:.4f}",
                        f"{moisture[time_index, radius_index]:.4f}",
                    ]
                )


def key_tables(solution: dict[str, Any]) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """提取表3、表4的 6 个时间和 5 个半径结果。"""

    target_times = np.array([1800, 3600, 5400, 7200, 9000, 10800], dtype=float)
    time_s = np.asarray(solution["time_s"], dtype=float)
    indices = np.array([int(np.argmin(np.abs(time_s - t))) for t in target_times])
    if not np.allclose(time_s[indices], target_times, atol=1.0e-10):
        raise ValueError("完整结果中缺少表3/表4要求的时间点")
    temperature = np.round(np.asarray(solution["temperature_C"])[indices], 4)
    moisture = np.round(np.asarray(solution["moisture_kgkg"])[indices], 4)
    return target_times, temperature, moisture


def write_table_csv(
    solution: dict[str, Any],
    temperature_path: str | Path = DEFAULT_TABLE3_PATH,
    moisture_path: str | Path = DEFAULT_TABLE4_PATH,
) -> tuple[Path, Path]:
    """把表 3、表 4 写成 6 行 × 5 个关键半径的宽表 CSV。"""

    target_times, temperature, moisture = key_tables(solution)
    radius_columns = [0.0, 0.5, 1.0, 1.5, 2.0]
    radius_indices = np.array([0, 5, 10, 15, 20], dtype=int)
    header = ["time_h"] + [f"{value:.1f}cm" for value in radius_columns]
    for values, path in (
        (temperature, temperature_path),
        (moisture, moisture_path),
    ):
        destination = Path(path)
        destination.parent.mkdir(parents=True, exist_ok=True)
        with destination.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle, lineterminator="\n")
            writer.writerow(header)
            for time_value, row in zip(target_times, values):
                writer.writerow(
                    [f"{time_value / 3600.0:.1f}"]
                    + [f"{row[index]:.4f}" for index in radius_indices]
                )
    return Path(temperature_path), Path(moisture_path)


def diagnostics_summary(solution: dict[str, Any]) -> dict[str, Any]:
    """汇总问题二 Picard 统计、物理约束与守恒残差。"""

    time_s = np.asarray(solution["time_s"], dtype=float)
    radius_cm = np.asarray(solution["radius_cm"], dtype=float)
    temperature = np.asarray(solution["temperature_C"], dtype=float)
    moisture = np.asarray(solution["moisture_kgkg"], dtype=float)
    air_temperature = np.asarray(solution["air_temperature_C"], dtype=float)
    air_moisture = np.asarray(solution["air_moisture_kgkg"], dtype=float)
    iterations = np.asarray(solution["iteration_counts"], dtype=int)
    delta_temperature = np.asarray(solution["picard_delta_temperature"], dtype=float)
    delta_moisture = np.asarray(solution["picard_delta_moisture"], dtype=float)
    linear_residuals = np.asarray(solution["linear_residuals"], dtype=float)
    mass_residuals = np.asarray(solution["moisture_balance_residuals"], dtype=float)
    energy_residuals = np.asarray(solution["energy_balance_residuals"], dtype=float)
    mass_fluxes = np.asarray(solution["moisture_boundary_fluxes"], dtype=float)
    energy_fluxes = np.asarray(solution["energy_boundary_fluxes"], dtype=float)
    dt_value = float(solution["time_step_s"])

    radius_m = radius_cm * 1.0e-2
    weights = radius_m if radius_m[0] > 0.0 else radius_m + 1.0e-30
    trapezoid = getattr(np, "trapezoid", np.trapz)
    volume_average = trapezoid(
        moisture * weights[None, :], radius_m, axis=1
    ) / trapezoid(weights, radius_m)

    heat_driver = air_temperature - temperature[:, -1]
    moisture_driver = air_moisture - moisture[:, -1]

    def _relative(residuals: np.ndarray, fluxes: np.ndarray) -> tuple[float, float]:
        scale = np.maximum.reduce(
            [
                np.abs(residuals),
                np.abs(fluxes),
                np.full_like(residuals, np.finfo(float).tiny),
            ]
        )
        relative = np.abs(residuals) / scale
        index = int(np.argmax(np.abs(residuals)))
        return float(np.max(relative)), float((index + 1) * dt_value)

    mass_relative, mass_worst_time = _relative(mass_residuals, mass_fluxes)
    energy_relative, energy_worst_time = _relative(energy_residuals, energy_fluxes)
    return {
        "task": 2,
        "caliber": {
            "spatial_game": "守恒型径向有限体积（含中心半控制体）",
            "temporal": "Crank-Nicolson",
            "interface_property": "调和平均",
            "nonlinear": "同步半时间层 Picard",
            "n_intervals": int(round(RADIUS_M / float(solution["dr_m"]))),
            "dr_m": float(solution["dr_m"]),
            "dt_s": dt_value,
            "n_steps": int(iterations.size),
            "output_shape": [int(moisture.shape[0]), int(moisture.shape[1])],
            "picard_tol_temperature_C": float(PICARD_TOL_T_C),
            "picard_tol_moisture_kgkg": float(PICARD_TOL_C_KGKG),
            "linear_residual_tol": float(LINEAR_RESIDUAL_TOL),
        },
        "picard": {
            "max_iterations": int(np.max(iterations)),
            "mean_iterations": float(np.mean(iterations)),
            "max_final_delta_temperature_C": float(np.max(delta_temperature)),
            "max_final_delta_moisture_kgkg": float(np.max(delta_moisture)),
            "converged_steps": int(
                np.sum(
                    (delta_temperature <= PICARD_TOL_T_C)
                    & (delta_moisture <= PICARD_TOL_C_KGKG)
                )
            ),
            "failed_steps": int(
                np.sum(
                    (delta_temperature > PICARD_TOL_T_C)
                    | (delta_moisture > PICARD_TOL_C_KGKG)
                )
            ),
        },
        "linear_solver": {
            "max_normalized_inf_residual": float(np.max(linear_residuals)),
            "tol": float(LINEAR_RESIDUAL_TOL),
        },
        "physics": {
            "temperature_min_C": float(np.min(temperature)),
            "temperature_max_C": float(np.max(temperature)),
            "moisture_min_kgkg": float(np.min(moisture)),
            "moisture_max_kgkg": float(np.max(moisture)),
            "moisture_positive": bool(np.min(moisture) > 0.0),
            "finite": bool(
                np.all(np.isfinite(temperature)) and np.all(np.isfinite(moisture))
            ),
            "min_heat_boundary_driver_C": float(np.min(heat_driver[1:])),
            "max_moisture_boundary_driver_kgkg": float(np.max(moisture_driver[1:])),
            "final_surface_center_temperature_gap_C": float(
                temperature[-1, -1] - temperature[-1, 0]
            ),
            "final_center_surface_moisture_gap_kgkg": float(
                moisture[-1, 0] - moisture[-1, -1]
            ),
            "volume_average_moisture_initial_kgkg": float(volume_average[0]),
            "volume_average_moisture_final_kgkg": float(volume_average[-1]),
            "volume_average_moisture_monotone_decreasing": bool(
                np.all(np.diff(volume_average) <= 1.0e-12)
            ),
        },
        "conservation": {
            "moisture": {
                "quantity": "约去 2πL 的全域水分储量 Σ V_i C_i",
                "max_abs_balance_residual": float(np.max(np.abs(mass_residuals))),
                "max_relative_balance_residual": mass_relative,
                "worst_time_s": mass_worst_time,
                "boundary_flux_integral": float(np.sum(mass_fluxes) * dt_value),
            },
            "energy": {
                "quantity": "半时间层物性冻结下的全局焓变 Σ V_i (ρc_p)_mid T_i",
                "max_abs_balance_residual": float(np.max(np.abs(energy_residuals))),
                "max_relative_balance_residual": energy_relative,
                "worst_time_s": energy_worst_time,
                "boundary_flux_integral": float(np.sum(energy_fluxes) * dt_value),
            },
        },
    }


def write_diagnostics(
    solution: dict[str, Any],
    json_path: str | Path = DEFAULT_DIAGNOSTICS_JSON,
    csv_path: str | Path = DEFAULT_DIAGNOSTICS_CSV,
) -> tuple[Path, Path]:
    """把逐步 Picard/线性/守恒日志写为 CSV，并写汇总 JSON。"""

    import json

    summary = diagnostics_summary(solution)
    json_destination = Path(json_path)
    json_destination.parent.mkdir(parents=True, exist_ok=True)
    json_destination.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    step_count = np.asarray(solution["iteration_counts"]).size
    dt_value = float(solution["time_step_s"])
    iterations = np.asarray(solution["iteration_counts"], dtype=int)
    delta_temperature = np.asarray(solution["picard_delta_temperature"], dtype=float)
    delta_moisture = np.asarray(solution["picard_delta_moisture"], dtype=float)
    linear_residuals = np.asarray(solution["linear_residuals"], dtype=float)
    mass_residuals = np.asarray(solution["moisture_balance_residuals"], dtype=float)
    energy_residuals = np.asarray(solution["energy_balance_residuals"], dtype=float)
    csv_destination = Path(csv_path)
    csv_destination.parent.mkdir(parents=True, exist_ok=True)
    with csv_destination.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(
            (
                "step",
                "time_s",
                "picard_iterations",
                "picard_delta_temperature_C",
                "picard_delta_moisture_kgkg",
                "linear_normalized_inf_residual",
                "moisture_balance_residual",
                "energy_balance_residual",
            )
        )
        for step in range(step_count):
            writer.writerow(
                (
                    step + 1,
                    f"{(step + 1) * dt_value:.6f}",
                    int(iterations[step]),
                    f"{delta_temperature[step]:.6e}",
                    f"{delta_moisture[step]:.6e}",
                    f"{linear_residuals[step]:.6e}",
                    f"{mass_residuals[step]:.6e}",
                    f"{energy_residuals[step]:.6e}",
                )
            )
    return json_destination, csv_destination


def _key_point_values(
    solution: dict[str, Any],
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """提取表3/表4的时间和五个关键径向位置。"""

    target_times = np.array([1800, 3600, 5400, 7200, 9000, 10800], dtype=float)
    time_s = np.asarray(solution["time_s"], dtype=float)
    time_indices = np.array(
        [int(np.argmin(np.abs(time_s - target))) for target in target_times]
    )
    if not np.allclose(time_s[time_indices], target_times, atol=1.0e-10):
        raise ValueError("完整结果中缺少敏感性检查要求的时间点")
    key_radius_indices = np.array([0, 5, 10, 15, 20], dtype=int)
    return (
        target_times,
        OUTPUT_RADIUS_CM[key_radius_indices],
        np.asarray(solution["temperature_C"])[
            time_indices[:, None], key_radius_indices[None, :]
        ],
        np.asarray(solution["moisture_kgkg"])[
            time_indices[:, None], key_radius_indices[None, :]
        ],
    )


def time_step_sensitivity(
    base_solution: dict[str, Any] | None = None,
    data_path: str | Path = DEFAULT_DATA_PATH,
) -> dict[str, Any]:
    """以正式 ``dt=1 s`` 结果为基准，比较 ``dt=0.5 s``。

    两次计算均使用正式 ``N=1600`` 网格和 1 s 输出，在表3/表4的时空
    位置比较原始值，作为 handoff §21.3 的时间步减半校核。
    """

    if base_solution is not None and np.isclose(
        float(base_solution["time_step_s"]), 1.0
    ) and np.isclose(float(base_solution["dr_m"]), BASE_DR_M):
        reference_solution = base_solution
    else:
        reference_solution = solve_task2(
            data_path=data_path,
            t_end_s=FINAL_TIME_S,
            time_step_s=1.0,
            output_dt_s=1.0,
            dr_m=BASE_DR_M,
            progress=True,
        )
    comparison_solution = solve_task2(
        data_path=data_path,
        t_end_s=FINAL_TIME_S,
        time_step_s=0.5,
        output_dt_s=1.0,
        dr_m=BASE_DR_M,
        progress=True,
    )
    validate_solution(comparison_solution)
    key_times_s, key_radius_cm, reference_temperature, reference_moisture = (
        _key_point_values(reference_solution)
    )
    _, _, comparison_temperature, comparison_moisture = _key_point_values(
        comparison_solution
    )
    temperature_difference = np.abs(comparison_temperature - reference_temperature)
    moisture_difference = np.abs(comparison_moisture - reference_moisture)
    temperature_position = np.unravel_index(
        int(np.argmax(temperature_difference)), temperature_difference.shape
    )
    moisture_position = np.unravel_index(
        int(np.argmax(moisture_difference)), moisture_difference.shape
    )
    return {
        "reference_solution": reference_solution,
        "comparison_solution": comparison_solution,
        "reference_time_step_s": 1.0,
        "comparison_time_step_s": 0.5,
        "dr_m": BASE_DR_M,
        "key_times_s": key_times_s,
        "key_radius_cm": key_radius_cm,
        "temperature_max_abs_difference_C": float(np.max(temperature_difference)),
        "temperature_difference_position": (
            float(key_times_s[temperature_position[0]]),
            float(key_radius_cm[temperature_position[1]]),
        ),
        "moisture_max_abs_difference_kgkg": float(np.max(moisture_difference)),
        "moisture_difference_position": (
            float(key_times_s[moisture_position[0]]),
            float(key_radius_cm[moisture_position[1]]),
        ),
    }


def space_grid_sensitivity(
    base_solution: dict[str, Any] | None = None,
    data_path: str | Path = DEFAULT_DATA_PATH,
) -> dict[str, Any]:
    """执行 ``N=400 -> 800 -> 1600`` 的逐级网格敏感性检查。

    所有网格使用 handoff 规定的 CN、同步 Picard、1 s 时间步和 1 s 输出。
    返回相邻网格以及粗网格到正式精细网格的表3/表4关键点最大差异。
    """

    if base_solution is not None and np.isclose(
        float(base_solution["time_step_s"]), 1.0
    ) and np.isclose(float(base_solution["dr_m"]), BASE_DR_M):
        fine_solution = base_solution
    else:
        fine_solution = solve_task2(
            data_path=data_path,
            t_end_s=FINAL_TIME_S,
            time_step_s=1.0,
            output_dt_s=1.0,
            dr_m=BASE_DR_M,
            progress=True,
        )
    middle_solution = solve_task2(
        data_path=data_path,
        t_end_s=FINAL_TIME_S,
        time_step_s=1.0,
        output_dt_s=1.0,
        dr_m=RADIUS_M / 800.0,
        progress=True,
    )
    coarse_solution = solve_task2(
        data_path=data_path,
        t_end_s=FINAL_TIME_S,
        time_step_s=1.0,
        output_dt_s=1.0,
        dr_m=RADIUS_M / 400.0,
        progress=True,
    )
    validate_solution(middle_solution)
    validate_solution(coarse_solution)

    key_times_s, key_radius_cm, fine_temperature, fine_moisture = _key_point_values(
        fine_solution
    )
    _, _, middle_temperature, middle_moisture = _key_point_values(middle_solution)
    _, _, coarse_temperature, coarse_moisture = _key_point_values(coarse_solution)

    def compare(
        left_temperature: np.ndarray,
        left_moisture: np.ndarray,
        right_temperature: np.ndarray,
        right_moisture: np.ndarray,
    ) -> dict[str, Any]:
        temperature_difference = np.abs(left_temperature - right_temperature)
        moisture_difference = np.abs(left_moisture - right_moisture)
        temperature_position = np.unravel_index(
            int(np.argmax(temperature_difference)), temperature_difference.shape
        )
        moisture_position = np.unravel_index(
            int(np.argmax(moisture_difference)), moisture_difference.shape
        )
        return {
            "temperature_max_abs_difference_C": float(np.max(temperature_difference)),
            "temperature_difference_position": (
                float(key_times_s[temperature_position[0]]),
                float(key_radius_cm[temperature_position[1]]),
            ),
            "moisture_max_abs_difference_kgkg": float(np.max(moisture_difference)),
            "moisture_difference_position": (
                float(key_times_s[moisture_position[0]]),
                float(key_radius_cm[moisture_position[1]]),
            ),
        }

    n400_vs_n800 = compare(
        coarse_temperature,
        coarse_moisture,
        middle_temperature,
        middle_moisture,
    )
    n800_vs_n1600 = compare(
        middle_temperature,
        middle_moisture,
        fine_temperature,
        fine_moisture,
    )
    n400_vs_n1600 = compare(
        coarse_temperature,
        coarse_moisture,
        fine_temperature,
        fine_moisture,
    )
    return {
        "fine_solution": fine_solution,
        "middle_solution": middle_solution,
        "coarse_solution": coarse_solution,
        "fine_n_intervals": 1600,
        "middle_n_intervals": 800,
        "coarse_n_intervals": 400,
        "fine_dr_m": BASE_DR_M,
        "middle_dr_m": RADIUS_M / 800.0,
        "coarse_dr_m": RADIUS_M / 400.0,
        "time_step_s": 1.0,
        "key_times_s": key_times_s,
        "key_radius_cm": key_radius_cm,
        "n400_vs_n800": n400_vs_n800,
        "n800_vs_n1600": n800_vs_n1600,
        "n400_vs_n1600": n400_vs_n1600,
    }


def _print_key_tables(solution: dict[str, Any]) -> None:
    target_times, temperature, moisture = key_tables(solution)
    print("表3 温度 / °C（行依次为 0.5--3.0 h，列为 0--2 cm）")
    print("时间/h," + ",".join(f"{r:.1f} cm" for r in OUTPUT_RADIUS_CM[[0, 5, 10, 15, 20]]))
    for time_value, row in zip(target_times, temperature):
        print(f"{time_value / 3600.0:.1f}," + ",".join(f"{x:.4f}" for x in row[[0, 5, 10, 15, 20]]))
    print("表4 水分浓度 / kg/kg（行依次为 0.5--3.0 h，列为 0--2 cm）")
    print("时间/h," + ",".join(f"{r:.1f} cm" for r in OUTPUT_RADIUS_CM[[0, 5, 10, 15, 20]]))
    for time_value, row in zip(target_times, moisture):
        print(f"{time_value / 3600.0:.1f}," + ",".join(f"{x:.4f}" for x in row[[0, 5, 10, 15, 20]]))


def main() -> None:
    parser = argparse.ArgumentParser(description="求解 A题问题二")
    parser.add_argument("--t-end", type=float, default=FINAL_TIME_S)
    parser.add_argument("--time-step", type=float, default=1.0)
    parser.add_argument("--output-dt", type=float, default=OUTPUT_DT_S)
    parser.add_argument("--no-write", action="store_true", help="只求解和检查，不写答案文件")
    parser.add_argument("--skip-sensitivity", action="store_true")
    parser.add_argument("--data-path", type=Path, default=DEFAULT_DATA_PATH)
    args = parser.parse_args()

    started = wall_time.perf_counter()
    print(
        f"问题二求解开始: t_end={args.t_end:g}s, dt={args.time_step:g}s, "
        f"dr={BASE_DR_M:g}m"
    )
    solution = solve_task2(
        data_path=args.data_path,
        t_end_s=args.t_end,
        time_step_s=args.time_step,
        output_dt_s=args.output_dt,
        progress=True,
    )
    summary = validate_solution(solution)
    print(f"基本数值检查: {summary}")
    if not args.no_write:
        write_result_xlsx(solution)
        write_result_csv(solution)
        table3_path, table4_path = write_table_csv(solution)
        json_path, diagnostics_path = write_diagnostics(solution)
        diagnostics = diagnostics_summary(solution)
        print(f"已生成: {DEFAULT_XLSX_PATH}")
        print(f"已生成: {DEFAULT_CSV_PATH}")
        print(f"已生成: {table3_path}")
        print(f"已生成: {table4_path}")
        print(f"诊断日志: {json_path}, {diagnostics_path}")
        print(
            "Picard: max={} mean={:.3f}; 守恒最大相对残差: 水分={:.3e}, 能量={:.3e}".format(
                diagnostics["picard"]["max_iterations"],
                diagnostics["picard"]["mean_iterations"],
                diagnostics["conservation"]["moisture"]["max_relative_balance_residual"],
                diagnostics["conservation"]["energy"]["max_relative_balance_residual"],
            )
        )
    if (
        not args.skip_sensitivity
        and np.isclose(args.t_end, FINAL_TIME_S)
        and np.isclose(args.time_step, 1.0)
    ):
        sensitivity = time_step_sensitivity(solution, data_path=args.data_path)
        print(
            "时间步敏感性（基准 dt=1 s，对比 dt=0.5 s，dr=0.00125 cm，表3/表4点）: "
            f"max|dT|={sensitivity['temperature_max_abs_difference_C']:.6e} °C "
            f"at (t,r)={sensitivity['temperature_difference_position']}; "
            f"max|dC|={sensitivity['moisture_max_abs_difference_kgkg']:.6e} kg/kg "
            f"at (t,r)={sensitivity['moisture_difference_position']}"
        )
        grid_sensitivity = space_grid_sensitivity(solution, data_path=args.data_path)
        print(
            "空间网格敏感性（N=400→800→1600，dt=1 s，表3/表4点）: "
            f"N400-N800: max|dT|={grid_sensitivity['n400_vs_n800']['temperature_max_abs_difference_C']:.6e} °C, "
            f"max|dC|={grid_sensitivity['n400_vs_n800']['moisture_max_abs_difference_kgkg']:.6e} kg/kg; "
            f"N800-N1600: max|dT|={grid_sensitivity['n800_vs_n1600']['temperature_max_abs_difference_C']:.6e} °C, "
            f"max|dC|={grid_sensitivity['n800_vs_n1600']['moisture_max_abs_difference_kgkg']:.6e} kg/kg"
        )
        _print_key_tables(solution)
    print(f"总耗时: {wall_time.perf_counter() - started:.2f} s")


if __name__ == "__main__":
    main()
