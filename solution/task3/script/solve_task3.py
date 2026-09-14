"""A题问题三：固定半径圆柱的长期热质耦合烘干求解器。

本脚本独立实现 handoff/task3_handoff.md 规定的守恒径向有限体积模型：
界面物性使用调和平均，时间方向使用 Crank--Nicolson，非线性物性在
半时间层通过 Picard 迭代更新。前 14400 s 采用附件1分段线性边界，
之后采用附件1稳定段(9000--14400 s)均值。程序从 t=0 连续积分，
对完整内部水分场检查阈值，并按 handoff 要求从跨越步前态局部二分重积分，
返回严格满足 max(C)<0.15 的右端点。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import time as wall_time
from copy import copy
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.linalg import solve_banded


# 题设几何、初值和问题二/三沿用的传递参数。
RADIUS_M = 0.02
LENGTH_M = 0.25
INITIAL_TEMPERATURE_C = 28.0
INITIAL_MOISTURE_KGKG = 2.55
HEAT_TRANSFER_W_M2_K = 25.0
MASS_TRANSFER_M_S = 8.0e-7

# handoff/task3_handoff.md 的最终离散规格。
N_DEFAULT = 1600
TIME_STEP_EARLY_S = 1.0
TIME_STEP_LONG_DEFAULT_S = 10.0
OUTPUT_DT_S = 60.0
STABLE_START_S = 9000.0
AIR_DATA_END_S = 14400.0
TABLE5_DT_S = 6.0 * 3600.0
DRYING_THRESHOLD_KGKG = 0.15
DEFAULT_MAX_TIME_S = 7.0 * 24.0 * 3600.0

PICARD_TOL_T_C = 1.0e-8
PICARD_TOL_C_KGKG = 1.0e-10
PICARD_MAX_ITER = 50
LINEAR_RESIDUAL_TOL = 1.0e-9

REPO_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_DATA_PATH = REPO_ROOT / "problem" / "附件" / "附件1.xlsx"
DEFAULT_TEMPLATE_PATH = REPO_ROOT / "problem" / "附件" / "附件3" / "result3.xlsx"
DEFAULT_ANSWER_DIR = Path(__file__).resolve().parents[1] / "answer"
DEFAULT_XLSX_PATH = DEFAULT_ANSWER_DIR / "result3.xlsx"
DEFAULT_CSV_PATH = DEFAULT_ANSWER_DIR / "task3_all_results.csv"
DEFAULT_TASK2_CSV_PATH = REPO_ROOT / "solution" / "task2" / "answer" / "task2_all_results.csv"
DEFAULT_HANDOFF_PATH = REPO_ROOT / "handoff" / "task3_handoff.md"
DEFAULT_DIAGNOSTICS_JSON = DEFAULT_ANSWER_DIR / "task3_diagnostics.json"
DEFAULT_DIAGNOSTICS_CSV = DEFAULT_ANSWER_DIR / "task3_diagnostics.csv"

OUTPUT_RADIUS_CM = np.arange(0.0, 2.0000001, 0.1)
KEY_RADIUS_INDICES = np.array([0, 5, 10, 15, 20], dtype=int)


def load_air_data(
    path: str | Path = DEFAULT_DATA_PATH,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """读取附件1的时间、烘房温度和烘房水分浓度。"""

    frame = pd.read_excel(Path(path))
    columns = {str(column).strip(): column for column in frame.columns}
    required = ("时间", "温度", "水分浓度")
    missing = [name for name in required if name not in columns]
    if missing:
        raise ValueError(f"附件1缺少必需列: {missing}")
    selected = frame[[columns[name] for name in required]].copy()
    selected.columns = required
    for column in required:
        selected[column] = pd.to_numeric(selected[column], errors="raise")
    selected = selected.sort_values("时间").drop_duplicates("时间", keep="last")
    time_s = selected["时间"].to_numpy(dtype=float)
    temperature_C = selected["温度"].to_numpy(dtype=float)
    moisture_kgkg = selected["水分浓度"].to_numpy(dtype=float)
    values = np.column_stack((time_s, temperature_C, moisture_kgkg))
    if time_s.size < 2 or np.any(np.diff(time_s) <= 0.0):
        raise ValueError("附件1时间必须包含至少两个严格递增时间点")
    if time_s[0] > 1.0e-10 or time_s[-1] < AIR_DATA_END_S - 1.0e-10:
        raise ValueError("附件1未覆盖到14400 s")
    if not np.all(np.isfinite(values)):
        raise ValueError("附件1数据包含NaN或Inf")
    return time_s, temperature_C, moisture_kgkg


def stable_air_means(
    time_s: np.ndarray,
    temperature_C: np.ndarray,
    moisture_kgkg: np.ndarray,
) -> tuple[float, float, int]:
    """实时计算附件1在9000--14400 s（含端点）的统计均值。"""

    stable = (time_s >= STABLE_START_S - 1.0e-10) & (
        time_s <= AIR_DATA_END_S + 1.0e-10
    )
    if not np.any(stable):
        raise ValueError("附件1没有9000--14400 s稳定段数据")
    stable_temperature = float(np.mean(temperature_C[stable]))
    stable_moisture = float(np.mean(moisture_kgkg[stable]))
    return stable_temperature, stable_moisture, int(np.count_nonzero(stable))


def air_at(
    time_value_s: float,
    time_s: np.ndarray,
    temperature_C: np.ndarray,
    moisture_kgkg: np.ndarray,
    stable_temperature_C: float,
    stable_moisture_kgkg: float,
) -> tuple[float, float]:
    """附件范围内分段线性；14400 s 后使用稳定段均值。"""

    if time_value_s <= AIR_DATA_END_S + 1.0e-10:
        clipped = min(max(float(time_value_s), float(time_s[0])), AIR_DATA_END_S)
        return (
            float(np.interp(clipped, time_s, temperature_C)),
            float(np.interp(clipped, time_s, moisture_kgkg)),
        )
    return stable_temperature_C, stable_moisture_kgkg


def fixed_geometry(
    n_intervals: int,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, float]:
    """构造固定半径圆柱的径向节点、面和半控制体积测度。"""

    if n_intervals < 2:
        raise ValueError("径向区间数至少为2")
    dr_m = RADIUS_M / float(n_intervals)
    radii_m = np.linspace(0.0, RADIUS_M, n_intervals + 1)
    faces_m = 0.5 * (radii_m[:-1] + radii_m[1:])
    measure = np.empty_like(radii_m)
    measure[0] = 0.5 * faces_m[0] ** 2
    measure[1:-1] = 0.5 * (faces_m[1:] ** 2 - faces_m[:-1] ** 2)
    measure[-1] = 0.5 * (RADIUS_M**2 - faces_m[-1] ** 2)
    if not np.isclose(np.sum(measure), 0.5 * RADIUS_M**2, atol=1.0e-18):
        raise ValueError("径向控制体测度未覆盖整个圆柱截面")
    return radii_m, faces_m, measure, dr_m


def output_indices(radii_m: np.ndarray) -> np.ndarray:
    """返回0--2 cm、每0.1 cm输出位置对应的内部节点。"""

    target_m = OUTPUT_RADIUS_CM / 100.0
    indices = np.array([int(np.argmin(np.abs(radii_m - r))) for r in target_m])
    if not np.allclose(radii_m[indices], target_m, rtol=0.0, atol=1.0e-14):
        raise ValueError("当前N不能精确包含0.1 cm输出位置")
    return indices


def material_properties(
    temperature_C: np.ndarray,
    moisture_kgkg: np.ndarray,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """按附录3计算动态物性，扩散系数温度使用K。"""

    temperature_C = np.asarray(temperature_C, dtype=float)
    moisture_kgkg = np.asarray(moisture_kgkg, dtype=float)
    concentration = np.maximum(moisture_kgkg, 1.0e-12)
    temperature_K = temperature_C + 273.15
    if np.any(~np.isfinite(temperature_K)) or np.any(temperature_K <= 0.0):
        raise ValueError("物性计算遇到非正或非有限Kelvin温度")
    rho = 650.0 + 128.0 * concentration
    cp = 1450.0 + 2736.0 * concentration / (concentration + 1.0)
    conductivity = 0.21 + 0.38 * concentration / (concentration + 1.0)
    diffusivity = 2.4e-3 * np.exp(-0.45 / concentration) * np.exp(
        -3850.0 / temperature_K
    )
    return rho, cp, conductivity, diffusivity


def harmonic_mean(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    """返回正物性的界面调和平均。"""

    denominator = left + right
    if np.any(denominator <= 0.0) or not np.all(np.isfinite(denominator)):
        raise ValueError("界面物性非正或非有限")
    return 2.0 * left * right / denominator


def solve_tridiagonal(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    rhs: np.ndarray,
) -> np.ndarray:
    """求解三对角线性系统。"""

    n = diagonal.size
    banded = np.zeros((3, n), dtype=float)
    banded[0, 1:] = upper
    banded[1] = diagonal
    banded[2, :-1] = lower
    result = solve_banded((1, 1), banded, rhs, check_finite=False)
    if not np.all(np.isfinite(result)):
        raise RuntimeError("三对角线性系统求解得到NaN或Inf")
    return result


def solve_cn_field(
    old_field: np.ndarray,
    storage: np.ndarray,
    face_property: np.ndarray,
    faces_m: np.ndarray,
    dr_m: float,
    dt_s: float,
    boundary_conductance: float,
    air_value: float,
) -> tuple[np.ndarray, float]:
    """用冻结的半时间层物性求解一个守恒FVM的CN场。

    令 K 为正的扩散/Robin离散矩阵，空间算子为 ``-K u+b``。
    因而 CN 方程为

        (S+dt*K/2)u_new=(S-dt*K/2)u_old+dt*b_mid,

    其中边界值明确取 t+dt/2。
    """

    conductance = face_property * faces_m / dr_m
    diagonal_K = np.empty_like(storage)
    diagonal_K[0] = conductance[0]
    diagonal_K[1:-1] = conductance[:-1] + conductance[1:]
    diagonal_K[-1] = conductance[-1] + boundary_conductance

    half_dt = 0.5 * dt_s
    diagonal = storage + half_dt * diagonal_K
    lower = -half_dt * conductance
    upper = -half_dt * conductance

    # (S - dt*K/2) old 的三对角乘积；表面Robin源项为 dt*b_mid。
    rhs = (storage - half_dt * diagonal_K) * old_field
    rhs[:-1] += half_dt * conductance * old_field[1:]
    rhs[1:] += half_dt * conductance * old_field[:-1]
    rhs[-1] += dt_s * boundary_conductance * air_value
    result = solve_tridiagonal(lower, diagonal, upper, rhs)
    residual = diagonal * result - rhs
    if result.size > 1:
        residual[:-1] += upper * result[1:]
        residual[1:] += lower * result[:-1]
    scale = max(
        float(np.max(np.abs(rhs))),
        float(np.max(np.abs(diagonal * result))),
        float(np.max(np.abs(upper * result[1:]))) if result.size > 1 else 0.0,
        float(np.max(np.abs(lower * result[:-1]))) if result.size > 1 else 0.0,
        1.0e-30,
    )
    normalized_residual = float(np.max(np.abs(residual)) / scale)
    if (
        not np.isfinite(normalized_residual)
        or normalized_residual > LINEAR_RESIDUAL_TOL
    ):
        raise RuntimeError(
            "CN 三对角线性系统残差异常: "
            f"normalized_inf_residual={normalized_residual:.3e}, "
            f"tol={LINEAR_RESIDUAL_TOL:.3e}"
        )
    return result, normalized_residual


def picard_step(
    old_temperature_C: np.ndarray,
    old_moisture_kgkg: np.ndarray,
    faces_m: np.ndarray,
    measure: np.ndarray,
    dr_m: float,
    dt_s: float,
    air_temperature_C: float,
    air_moisture_kgkg: float,
) -> tuple[np.ndarray, np.ndarray, int, float, float, float, float]:
    """一个时间步的CN热质耦合Picard迭代。"""

    guess_temperature_C = old_temperature_C.copy()
    guess_moisture_kgkg = old_moisture_kgkg.copy()
    last_delta_temperature = np.inf
    last_delta_moisture = np.inf
    max_temperature_residual = 0.0
    max_moisture_residual = 0.0

    for iteration in range(1, PICARD_MAX_ITER + 1):
        # 物性严格由 old 与当前 new guess 的半时间层状态更新。
        midpoint_temperature_C = 0.5 * (old_temperature_C + guess_temperature_C)
        midpoint_moisture_kgkg = 0.5 * (old_moisture_kgkg + guess_moisture_kgkg)
        rho, cp, conductivity, diffusivity = material_properties(
            midpoint_temperature_C, midpoint_moisture_kgkg
        )
        midpoint_k = harmonic_mean(conductivity[:-1], conductivity[1:])
        midpoint_D = harmonic_mean(diffusivity[:-1], diffusivity[1:])

        new_temperature_C, temperature_residual = solve_cn_field(
            old_temperature_C,
            rho * cp * measure,
            midpoint_k,
            faces_m,
            dr_m,
            dt_s,
            RADIUS_M * HEAT_TRANSFER_W_M2_K,
            air_temperature_C,
        )
        new_moisture_kgkg, moisture_residual = solve_cn_field(
            old_moisture_kgkg,
            measure,
            midpoint_D,
            faces_m,
            dr_m,
            dt_s,
            RADIUS_M * MASS_TRANSFER_M_S,
            air_moisture_kgkg,
        )
        max_temperature_residual = max(max_temperature_residual, temperature_residual)
        max_moisture_residual = max(max_moisture_residual, moisture_residual)
        if np.min(new_moisture_kgkg) < -1.0e-10:
            raise RuntimeError("CN Picard步得到明显负水分浓度")
        new_moisture_kgkg = np.maximum(new_moisture_kgkg, 0.0)
        last_delta_temperature = float(
            np.max(np.abs(new_temperature_C - guess_temperature_C))
        )
        last_delta_moisture = float(
            np.max(np.abs(new_moisture_kgkg - guess_moisture_kgkg))
        )
        if (
            last_delta_temperature <= PICARD_TOL_T_C
            and last_delta_moisture <= PICARD_TOL_C_KGKG
        ):
            return (
                new_temperature_C,
                new_moisture_kgkg,
                iteration,
                last_delta_temperature,
                last_delta_moisture,
                max_temperature_residual,
                max_moisture_residual,
            )
        guess_temperature_C = new_temperature_C
        guess_moisture_kgkg = new_moisture_kgkg

    raise RuntimeError(
        "CN Picard迭代未收敛: "
        f"dt={dt_s:g}s, iter={PICARD_MAX_ITER}, "
        f"max|dT|={last_delta_temperature:.3e}, "
        f"max|dC|={last_delta_moisture:.3e}"
    )


def locate_threshold_by_bisection(
    base_time_s: float,
    base_temperature_C: np.ndarray,
    base_moisture_kgkg: np.ndarray,
    upper_time_s: float,
    upper_temperature_C: np.ndarray,
    upper_moisture_kgkg: np.ndarray,
    air_time_s: np.ndarray,
    air_temperature_C: np.ndarray,
    air_moisture_kgkg: np.ndarray,
    stable_temperature_C: float,
    stable_moisture_kgkg: float,
    faces_m: np.ndarray,
    measure: np.ndarray,
    dr_m: float,
    width_tol_s: float = 0.1,
) -> tuple[
    float,
    float,
    np.ndarray,
    np.ndarray,
    np.ndarray,
    np.ndarray,
    float,
    float,
]:
    """按 handoff 第17节从固定跨越前状态局部二分阈值时刻。

    每个中点都从同一 ``(base_temperature_C, base_moisture_kgkg)`` 重新
    积分，而不是把历史结果或线性插值当作终止场。返回严格满足阈值的
    右端点及其状态，同时返回最终左端点状态和两端最大水分。
    """

    if width_tol_s <= 0.0:
        raise ValueError("阈值二分宽度必须为正")
    threshold = DRYING_THRESHOLD_KGKG
    lower_time = float(base_time_s)
    lower_temperature = np.asarray(base_temperature_C, dtype=float).copy()
    lower_moisture = np.asarray(base_moisture_kgkg, dtype=float).copy()
    upper_time = float(upper_time_s)
    upper_temperature = np.asarray(upper_temperature_C, dtype=float).copy()
    upper_moisture = np.asarray(upper_moisture_kgkg, dtype=float).copy()
    lower_max = float(np.max(lower_moisture))
    upper_max = float(np.max(upper_moisture))
    if lower_max < threshold or upper_max >= threshold:
        raise ValueError("阈值二分初始括号不满足左端>=阈值且右端<阈值")

    while upper_time - lower_time > width_tol_s:
        midpoint_time = 0.5 * (lower_time + upper_time)
        midpoint_dt = midpoint_time - base_time_s
        midpoint_air = air_at(
            base_time_s + 0.5 * midpoint_dt,
            air_time_s,
            air_temperature_C,
            air_moisture_kgkg,
            stable_temperature_C,
            stable_moisture_kgkg,
        )
        midpoint_temperature, midpoint_moisture, _, _, _, _, _ = picard_step(
            base_temperature_C,
            base_moisture_kgkg,
            faces_m,
            measure,
            dr_m,
            midpoint_dt,
            midpoint_air[0],
            midpoint_air[1],
        )
        midpoint_max = float(np.max(midpoint_moisture))
        if midpoint_max < threshold:
            upper_time = midpoint_time
            upper_temperature = midpoint_temperature
            upper_moisture = midpoint_moisture
            upper_max = midpoint_max
        else:
            lower_time = midpoint_time
            lower_temperature = midpoint_temperature
            lower_moisture = midpoint_moisture
            lower_max = midpoint_max

    if not lower_max >= threshold or not upper_max < threshold:
        raise RuntimeError("阈值二分结束时未保持合法的左右括号")
    return (
        upper_time,
        lower_time,
        lower_temperature,
        lower_moisture,
        upper_temperature,
        upper_moisture,
        lower_max,
        upper_max,
    )


def is_multiple(time_value_s: float, interval_s: float) -> bool:
    """判断浮点时间是否为给定间隔的整数倍。"""

    quotient = time_value_s / interval_s
    return abs(quotient - round(quotient)) <= 1.0e-9


def solve_task3(
    data_path: str | Path = DEFAULT_DATA_PATH,
    n_intervals: int = N_DEFAULT,
    time_step_long_s: float = TIME_STEP_LONG_DEFAULT_S,
    output_dt_s: float = OUTPUT_DT_S,
    max_time_s: float = DEFAULT_MAX_TIME_S,
    progress: bool = False,
    collect_first3: bool = True,
) -> dict[str, Any]:
    """从 t=0 连续积分，返回规则采样、终止括号和验证所需状态。"""

    if time_step_long_s <= 0.0 or output_dt_s <= 0.0 or max_time_s <= 0.0:
        raise ValueError("时间步、输出间隔和时间上限必须为正")
    if not is_multiple(output_dt_s, time_step_long_s) or not is_multiple(
        output_dt_s, TIME_STEP_EARLY_S
    ):
        raise ValueError("60 s输出间隔必须同时整除两段时间步")
    air_time_s, air_temperature, air_moisture = load_air_data(data_path)
    stable_temperature_C, stable_moisture_kgkg, stable_count = stable_air_means(
        air_time_s, air_temperature, air_moisture
    )
    radii_m, faces_m, measure, dr_m = fixed_geometry(n_intervals)
    indices = output_indices(radii_m)
    key_indices = indices[KEY_RADIUS_INDICES]
    n_nodes = radii_m.size

    current_temperature = np.full(n_nodes, INITIAL_TEMPERATURE_C, dtype=float)
    current_moisture = np.full(n_nodes, INITIAL_MOISTURE_KGKG, dtype=float)
    current_time_s = 0.0

    sampled_times: list[float] = [0.0]
    sampled_moisture: list[np.ndarray] = [current_moisture[indices].copy()]
    first3 = (
        np.full((10801, indices.size), np.nan, dtype=float)
        if collect_first3
        else None
    )
    first3_temperature = (
        np.full((10801, indices.size), np.nan, dtype=float)
        if collect_first3
        else None
    )
    if first3 is not None:
        assert first3_temperature is not None
        first3_temperature[0] = current_temperature[indices]
        first3[0] = current_moisture[indices]
    table5_times: list[float] = []
    table5_values: list[np.ndarray] = []
    step_times_s: list[float] = []
    step_dt_s: list[float] = []
    iteration_counts: list[int] = []
    delta_temperature: list[float] = []
    delta_moisture: list[float] = []
    temperature_linear_residuals: list[float] = []
    moisture_linear_residuals: list[float] = []
    temperature_balance_residuals: list[float] = []
    moisture_balance_residuals: list[float] = []
    temperature_boundary_fluxes: list[float] = []
    moisture_boundary_fluxes: list[float] = []

    previous_max = float(np.max(current_moisture))
    max_positive_step_change = -np.inf
    max_positive_average_change = -np.inf
    max_abs_step_change = 0.0
    max_radial_increase = -np.inf
    volume_measure = float(np.sum(measure))
    previous_average = float(np.sum(measure * current_moisture) / volume_measure)
    volume_average_history = [previous_average]
    min_moisture = float(np.min(current_moisture))
    max_moisture = float(np.max(current_moisture))
    min_temperature = float(np.min(current_temperature))
    max_temperature = float(np.max(current_temperature))
    early_steps = 0
    long_steps = 0

    termination_time_s: float | None = None
    termination_lower_time_s: float | None = None
    termination_upper_time_s: float | None = None
    termination_lower_temperature: np.ndarray | None = None
    termination_lower_moisture: np.ndarray | None = None
    termination_upper_temperature: np.ndarray | None = None
    termination_upper_moisture: np.ndarray | None = None
    termination_lower_max = np.nan
    termination_upper_max = np.nan
    termination_crossing_base_time_s: float | None = None
    termination_crossing_base_temperature: np.ndarray | None = None
    termination_crossing_base_moisture: np.ndarray | None = None
    termination_crossing_bracket_width_s = np.nan
    termination_bisection_iterations = 0

    max_steps_guard = int(np.ceil(max_time_s / TIME_STEP_EARLY_S)) + 1
    for _ in range(max_steps_guard):
        if current_time_s >= max_time_s - 1.0e-10:
            break
        if current_time_s < AIR_DATA_END_S - 1.0e-10:
            dt_s = min(TIME_STEP_EARLY_S, AIR_DATA_END_S - current_time_s)
            early_steps += 1
        else:
            dt_s = time_step_long_s
            long_steps += 1
        time_new_s = current_time_s + dt_s
        midpoint_air = air_at(
            current_time_s + 0.5 * dt_s,
            air_time_s,
            air_temperature,
            air_moisture,
            stable_temperature_C,
            stable_moisture_kgkg,
        )
        old_temperature = current_temperature
        old_moisture = current_moisture
        (
            current_temperature,
            current_moisture,
            iterations,
            step_delta_temperature,
            step_delta_moisture,
            temperature_linear_residual,
            moisture_linear_residual,
        ) = picard_step(
            old_temperature,
            old_moisture,
            faces_m,
            measure,
            dr_m,
            dt_s,
            midpoint_air[0],
            midpoint_air[1],
        )
        iteration_counts.append(iterations)
        delta_temperature.append(step_delta_temperature)
        delta_moisture.append(step_delta_moisture)
        step_times_s.append(float(time_new_s))
        step_dt_s.append(float(dt_s))
        temperature_linear_residuals.append(float(temperature_linear_residual))
        moisture_linear_residuals.append(float(moisture_linear_residual))

        change = current_moisture - old_moisture
        average = float(np.sum(measure * current_moisture) / volume_measure)
        max_positive_step_change = max(max_positive_step_change, float(np.max(change)))
        max_positive_average_change = max(
            max_positive_average_change, average - previous_average
        )
        max_abs_step_change = max(max_abs_step_change, float(np.max(np.abs(change))))
        max_radial_increase = max(max_radial_increase, float(np.max(np.diff(current_moisture))))
        previous_average = average
        volume_average_history.append(average)

        # 质量和能量守恒残差与 Task1/2/4 使用同一约去 2πL 定义。
        moisture_flux = RADIUS_M * MASS_TRANSFER_M_S * (
            midpoint_air[1]
            - 0.5 * (old_moisture[-1] + current_moisture[-1])
        )
        moisture_boundary_fluxes.append(float(moisture_flux))
        moisture_balance_residuals.append(
            float(np.sum(measure * (current_moisture - old_moisture))) / dt_s
            - float(moisture_flux)
        )
        midpoint_temperature = 0.5 * (old_temperature + current_temperature)
        midpoint_moisture = 0.5 * (old_moisture + current_moisture)
        rho_mid, cp_mid, _, _ = material_properties(
            midpoint_temperature, midpoint_moisture
        )
        temperature_flux = RADIUS_M * HEAT_TRANSFER_W_M2_K * (
            midpoint_air[0]
            - 0.5 * (old_temperature[-1] + current_temperature[-1])
        )
        temperature_boundary_fluxes.append(float(temperature_flux))
        temperature_balance_residuals.append(
            float(
                np.sum(
                    rho_mid
                    * cp_mid
                    * measure
                    * (current_temperature - old_temperature)
                )
            )
            / dt_s
            - float(temperature_flux)
        )
        min_moisture = min(min_moisture, float(np.min(current_moisture)))
        max_moisture = max(max_moisture, float(np.max(current_moisture)))
        min_temperature = min(min_temperature, float(np.min(current_temperature)))
        max_temperature = max(max_temperature, float(np.max(current_temperature)))

        if first3 is not None and time_new_s <= 10800.0 + 1.0e-10:
            assert first3_temperature is not None
            first3_temperature[int(round(time_new_s))] = current_temperature[indices]
            first3[int(round(time_new_s))] = current_moisture[indices]

        current_max = float(np.max(current_moisture))
        crossing = previous_max >= DRYING_THRESHOLD_KGKG and current_max < DRYING_THRESHOLD_KGKG
        if crossing:
            termination_crossing_base_time_s = float(current_time_s)
            termination_crossing_base_temperature = old_temperature.copy()
            termination_crossing_base_moisture = old_moisture.copy()
            termination_crossing_bracket_width_s = float(dt_s)
            termination_bisection_iterations = max(
                0, int(np.ceil(np.log2(dt_s / 0.1)))
            )
            (
                termination_time_s,
                termination_lower_time_s,
                termination_lower_temperature,
                termination_lower_moisture,
                termination_upper_temperature,
                termination_upper_moisture,
                termination_lower_max,
                termination_upper_max,
            ) = locate_threshold_by_bisection(
                base_time_s=current_time_s,
                base_temperature_C=old_temperature,
                base_moisture_kgkg=old_moisture,
                upper_time_s=time_new_s,
                upper_temperature_C=current_temperature,
                upper_moisture_kgkg=current_moisture,
                air_time_s=air_time_s,
                air_temperature_C=air_temperature,
                air_moisture_kgkg=air_moisture,
                stable_temperature_C=stable_temperature_C,
                stable_moisture_kgkg=stable_moisture_kgkg,
                faces_m=faces_m,
                measure=measure,
                dr_m=dr_m,
            )
            termination_upper_time_s = termination_time_s
            # 当前跨越步终点若在烘干时刻之后，不写入60 s结果网格。
            if is_multiple(time_new_s, output_dt_s) and time_new_s <= termination_time_s + 1.0e-9:
                sampled_times.append(float(time_new_s))
                sampled_moisture.append(current_moisture[indices].copy())
            break

        if is_multiple(time_new_s, output_dt_s):
            sampled_times.append(float(time_new_s))
            sampled_moisture.append(current_moisture[indices].copy())
        if is_multiple(time_new_s, TABLE5_DT_S):
            table5_times.append(float(time_new_s))
            table5_values.append(current_moisture[key_indices].copy())

        current_time_s = time_new_s
        previous_max = current_max
        if progress and (is_multiple(time_new_s, TABLE5_DT_S) or time_new_s <= 1.0):
            print(
                f"  已推进 {time_new_s / 3600.0:.4f} h，"
                f"max C={current_max:.10f}，Picard={iterations}"
            )
    else:
        raise RuntimeError("超过时间步保护上限仍未达到烘干阈值")

    if termination_time_s is None:
        raise RuntimeError(
            f"在{max_time_s / 3600.0:g} h内未达到全场阈值 "
            f"max C < {DRYING_THRESHOLD_KGKG:g}"
        )
    assert termination_lower_temperature is not None
    assert termination_lower_moisture is not None
    assert termination_upper_temperature is not None
    assert termination_upper_moisture is not None
    assert termination_upper_time_s is not None
    assert termination_crossing_base_time_s is not None
    assert termination_crossing_base_temperature is not None
    assert termination_crossing_base_moisture is not None

    # handoff 第17节规定以严格低于阈值的二分右端点作为终止状态。
    event_temperature = termination_upper_temperature.copy()
    event_moisture = termination_upper_moisture.copy()
    event_dt_s = float(termination_time_s - termination_crossing_base_time_s)
    if event_dt_s <= 0.0:
        raise RuntimeError("阈值事件右端点没有位于跨越步之后")
    event_air = air_at(
        termination_crossing_base_time_s + 0.5 * event_dt_s,
        air_time_s,
        air_temperature,
        air_moisture,
        stable_temperature_C,
        stable_moisture_kgkg,
    )
    event_moisture_flux = RADIUS_M * MASS_TRANSFER_M_S * (
        event_air[1]
        - 0.5 * (termination_crossing_base_moisture[-1] + event_moisture[-1])
    )
    event_moisture_balance = (
        float(
            np.sum(
                measure
                * (event_moisture - termination_crossing_base_moisture)
            )
        )
        / event_dt_s
        - float(event_moisture_flux)
    )
    event_midpoint_temperature = 0.5 * (
        termination_crossing_base_temperature + event_temperature
    )
    event_midpoint_moisture = 0.5 * (
        termination_crossing_base_moisture + event_moisture
    )
    event_rho, event_cp, _, _ = material_properties(
        event_midpoint_temperature, event_midpoint_moisture
    )
    event_temperature_flux = RADIUS_M * HEAT_TRANSFER_W_M2_K * (
        event_air[0]
        - 0.5
        * (termination_crossing_base_temperature[-1] + event_temperature[-1])
    )
    event_temperature_balance = (
        float(
            np.sum(
                event_rho
                * event_cp
                * measure
                * (event_temperature - termination_crossing_base_temperature)
            )
        )
        / event_dt_s
        - float(event_temperature_flux)
    )

    return {
        "n_intervals": int(n_intervals),
        "dr_m": float(dr_m),
        "time_step_early_s": float(TIME_STEP_EARLY_S),
        "time_step_long_s": float(time_step_long_s),
        "output_dt_s": float(output_dt_s),
        "radii_m_internal": radii_m,
        "key_indices": key_indices,
        "time_s": np.asarray(sampled_times, dtype=float),
        "radius_cm": OUTPUT_RADIUS_CM.copy(),
        "moisture_kgkg": np.asarray(sampled_moisture, dtype=float),
        "first3_temperature_C": first3_temperature,
        "first3_moisture_kgkg": first3,
        "table5_times_s": np.asarray(table5_times, dtype=float),
        "table5_values_kgkg": np.asarray(table5_values, dtype=float)
        if table5_values
        else np.empty((0, KEY_RADIUS_INDICES.size), dtype=float),
        "termination_time_s": float(termination_time_s),
        "termination_lower_time_s": float(termination_lower_time_s),
        "termination_upper_time_s": float(termination_upper_time_s),
        "termination_lower_temperature_C": termination_lower_temperature,
        "termination_lower_moisture_kgkg": termination_lower_moisture,
        "termination_upper_temperature_C": termination_upper_temperature,
        "termination_upper_moisture_kgkg": termination_upper_moisture,
        "termination_event_temperature_C": event_temperature,
        "termination_event_moisture_kgkg": event_moisture,
        "termination_lower_max_kgkg": float(termination_lower_max),
        "termination_upper_max_kgkg": float(termination_upper_max),
        "termination_crossing_base_time_s": float(termination_crossing_base_time_s),
        "termination_crossing_base_temperature_C": termination_crossing_base_temperature,
        "termination_crossing_base_moisture_kgkg": termination_crossing_base_moisture,
        "termination_crossing_bracket_width_s": float(
            termination_crossing_bracket_width_s
        ),
        "termination_bisection_iterations": int(termination_bisection_iterations),
        "termination_event_moisture_balance_residual": float(
            event_moisture_balance
        ),
        "termination_event_temperature_balance_residual": float(
            event_temperature_balance
        ),
        "termination_event_moisture_boundary_flux": float(event_moisture_flux),
        "termination_event_temperature_boundary_flux": float(event_temperature_flux),
        "stable_temperature_C": stable_temperature_C,
        "stable_moisture_kgkg": stable_moisture_kgkg,
        "stable_count": stable_count,
        "iteration_counts": np.asarray(iteration_counts, dtype=int),
        "picard_delta_temperature": np.asarray(delta_temperature, dtype=float),
        "picard_delta_moisture": np.asarray(delta_moisture, dtype=float),
        "step_times_s": np.asarray(step_times_s, dtype=float),
        "step_dt_s": np.asarray(step_dt_s, dtype=float),
        "temperature_linear_residuals": np.asarray(
            temperature_linear_residuals, dtype=float
        ),
        "moisture_linear_residuals": np.asarray(
            moisture_linear_residuals, dtype=float
        ),
        "linear_residuals": np.maximum(
            np.asarray(temperature_linear_residuals, dtype=float),
            np.asarray(moisture_linear_residuals, dtype=float),
        ),
        "temperature_balance_residuals": np.asarray(
            temperature_balance_residuals, dtype=float
        ),
        "moisture_balance_residuals": np.asarray(
            moisture_balance_residuals, dtype=float
        ),
        "temperature_boundary_fluxes": np.asarray(
            temperature_boundary_fluxes, dtype=float
        ),
        "moisture_boundary_fluxes": np.asarray(
            moisture_boundary_fluxes, dtype=float
        ),
        "volume_average_history": np.asarray(volume_average_history, dtype=float),
        "air_time_s": air_time_s.copy(),
        "air_temperature_C": air_temperature.copy(),
        "air_moisture_kgkg": air_moisture.copy(),
        "early_steps": int(early_steps),
        "long_steps": int(long_steps),
        "min_internal_moisture_kgkg": float(min_moisture),
        "max_internal_moisture_kgkg": float(max_moisture),
        "min_temperature_C": float(min_temperature),
        "max_temperature_C": float(max_temperature),
        "max_positive_step_change_kgkg": float(max_positive_step_change),
        "max_positive_average_change_kgkg": float(max_positive_average_change),
        "max_abs_step_change_kgkg": float(max_abs_step_change),
        "max_radial_increase_kgkg": float(max_radial_increase),
    }


def validate_solution(
    solution: dict[str, Any],
    require_first3: bool = True,
) -> dict[str, Any]:
    """执行终止、有限性、采样、单调性和CN结果的最低检查。"""

    times = np.asarray(solution["time_s"], dtype=float)
    radii = np.asarray(solution["radius_cm"], dtype=float)
    moisture = np.asarray(solution["moisture_kgkg"], dtype=float)
    lower_C = np.asarray(solution["termination_lower_moisture_kgkg"], dtype=float)
    upper_C = np.asarray(solution["termination_upper_moisture_kgkg"], dtype=float)
    event_C = np.asarray(solution["termination_event_moisture_kgkg"], dtype=float)
    lower_T = np.asarray(solution["termination_lower_temperature_C"], dtype=float)
    upper_T = np.asarray(solution["termination_upper_temperature_C"], dtype=float)
    event_T = np.asarray(solution["termination_event_temperature_C"], dtype=float)
    base_T = np.asarray(
        solution["termination_crossing_base_temperature_C"], dtype=float
    )
    base_C = np.asarray(
        solution["termination_crossing_base_moisture_kgkg"], dtype=float
    )
    arrays = (
        times,
        radii,
        moisture,
        lower_C,
        upper_C,
        event_C,
        lower_T,
        upper_T,
        event_T,
        base_T,
        base_C,
    )
    if not all(np.all(np.isfinite(array)) for array in arrays):
        raise ValueError("结果包含NaN或Inf")
    if moisture.shape != (times.size, radii.size):
        raise ValueError("规则采样水分场形状错误")
    if not all(
        array.shape == base_C.shape
        for array in (base_T, lower_T, upper_T, event_T, lower_C, upper_C, event_C)
    ):
        raise ValueError("终止事件剖面数组形状错误")
    if times.size < 1 or not np.isclose(times[0], 0.0, atol=1.0e-12):
        raise ValueError("输出时间未从0 s开始")
    if times.size > 1 and not np.allclose(np.diff(times), OUTPUT_DT_S, atol=1.0e-8):
        raise ValueError("输出时间间隔不是60 s")
    termination_time = float(solution["termination_time_s"])
    expected_last = np.floor(termination_time / OUTPUT_DT_S + 1.0e-12) * OUTPUT_DT_S
    if not np.isclose(times[-1], expected_last, atol=1.0e-8):
        raise ValueError("规则采样未止于烘干结束前最后一个60 s网格")
    expected_radii = np.arange(0.0, 2.0000001, 0.1)
    if radii.size != 21 or not np.allclose(radii, expected_radii, atol=1.0e-12):
        raise ValueError("输出半径不是0--2 cm每0.1 cm")
    if np.min(moisture) < -1.0e-10 or np.min(lower_C) < -1.0e-10 or np.min(upper_C) < -1.0e-10:
        raise ValueError("水分场出现明显负值")

    threshold = DRYING_THRESHOLD_KGKG
    lower_max = float(solution["termination_lower_max_kgkg"])
    upper_max = float(solution["termination_upper_max_kgkg"])
    event_max = float(np.max(event_C))
    lower_time = float(solution["termination_lower_time_s"])
    upper_time = float(solution["termination_upper_time_s"])
    if not lower_time < upper_time:
        raise ValueError("阈值二分左右端点时间顺序错误")
    if upper_time - lower_time > 0.1 + 1.0e-10:
        raise ValueError("阈值二分最终区间宽度超过0.1 s")
    if not np.isclose(float(solution["termination_time_s"]), upper_time, atol=1.0e-12):
        raise ValueError("终止时间不是阈值二分严格右端点")
    if not lower_max >= threshold:
        raise ValueError("终止前层没有保持max C >= 0.15")
    if not upper_max < threshold:
        raise ValueError("终止后层没有满足全场max C < 0.15")
    if not np.isclose(lower_max, np.max(lower_C), atol=1.0e-12):
        raise ValueError("终止前max C不是完整内部场最大值")
    if not np.isclose(upper_max, np.max(upper_C), atol=1.0e-12):
        raise ValueError("终止后max C不是完整内部场最大值")
    if not event_max < threshold:
        raise ValueError(f"二分右端点场max C未严格低于阈值: {event_max:.12g}")

    if require_first3:
        first3_temperature = solution["first3_temperature_C"]
        if first3_temperature is None or not np.all(np.isfinite(first3_temperature)):
            raise ValueError("未完整记录0--3 h连续温度场")
        first3 = solution["first3_moisture_kgkg"]
        if first3 is None or not np.all(np.isfinite(first3)):
            raise ValueError("未完整记录0--3 h连续状态")
    sampled_delta = np.diff(moisture, axis=0)
    max_positive_sampled = float(np.max(sampled_delta)) if sampled_delta.size else 0.0
    if np.min(event_C) < -1.0e-10:
        raise ValueError("二分右端点场出现明显负值")
    iterations = np.asarray(solution["iteration_counts"], dtype=int)
    step_times = np.asarray(solution["step_times_s"], dtype=float)
    step_dt = np.asarray(solution["step_dt_s"], dtype=float)
    temperature_linear = np.asarray(
        solution["temperature_linear_residuals"], dtype=float
    )
    moisture_linear = np.asarray(
        solution["moisture_linear_residuals"], dtype=float
    )
    linear_residuals = np.asarray(solution["linear_residuals"], dtype=float)
    temperature_balance = np.asarray(
        solution["temperature_balance_residuals"], dtype=float
    )
    moisture_balance = np.asarray(
        solution["moisture_balance_residuals"], dtype=float
    )
    temperature_fluxes = np.asarray(
        solution["temperature_boundary_fluxes"], dtype=float
    )
    moisture_fluxes = np.asarray(solution["moisture_boundary_fluxes"], dtype=float)
    expected_step_shape = iterations.shape
    step_arrays = (
        step_times,
        step_dt,
        temperature_linear,
        moisture_linear,
        linear_residuals,
        temperature_balance,
        moisture_balance,
        temperature_fluxes,
        moisture_fluxes,
    )
    if any(array.shape != expected_step_shape for array in step_arrays):
        raise ValueError("逐时间步诊断数组长度不一致")
    if iterations.size == 0 or int(np.max(iterations)) > PICARD_MAX_ITER:
        raise ValueError("Picard迭代记录错误")
    if not all(np.all(np.isfinite(array)) for array in step_arrays):
        raise ValueError("逐时间步诊断包含NaN或Inf")
    if np.any(step_dt <= 0.0) or not np.all(np.diff(step_times) > 0.0):
        raise ValueError("逐时间步时间记录错误")
    if np.max(linear_residuals) > LINEAR_RESIDUAL_TOL:
        raise ValueError("三对角线性系统残差超过容差")
    delta_temperature = np.asarray(
        solution["picard_delta_temperature"], dtype=float
    )
    delta_moisture = np.asarray(solution["picard_delta_moisture"], dtype=float)
    if delta_temperature.shape != expected_step_shape or delta_moisture.shape != expected_step_shape:
        raise ValueError("Picard误差记录长度错误")
    if np.any(delta_temperature > PICARD_TOL_T_C):
        raise ValueError("存在未达到温度Picard收敛阈值的时间步")
    if np.any(delta_moisture > PICARD_TOL_C_KGKG):
        raise ValueError("存在未达到水分Picard收敛阈值的时间步")

    return {
        "n_intervals": int(solution["n_intervals"]),
        "dr_cm": float(solution["dr_m"]) * 100.0,
        "time_step_early_s": float(solution["time_step_early_s"]),
        "time_step_long_s": float(solution["time_step_long_s"]),
        "termination_time_s": termination_time,
        "termination_time_h": termination_time / 3600.0,
        "termination_lower_time_s": lower_time,
        "termination_upper_time_s": upper_time,
        "termination_bracket_width_s": upper_time - lower_time,
        "termination_lower_max_kgkg": lower_max,
        "termination_upper_max_kgkg": upper_max,
        "termination_event_max_kgkg": event_max,
        "termination_event_max_index": int(np.argmax(event_C)),
        "termination_event_center_kgkg": float(event_C[0]),
        "termination_event_surface_kgkg": float(event_C[-1]),
        "min_internal_moisture_kgkg": float(solution["min_internal_moisture_kgkg"]),
        "max_internal_moisture_kgkg": float(solution["max_internal_moisture_kgkg"]),
        "min_temperature_C": float(solution["min_temperature_C"]),
        "max_temperature_C": float(solution["max_temperature_C"]),
        "max_positive_step_change_kgkg": float(solution["max_positive_step_change_kgkg"]),
        "max_positive_average_change_kgkg": float(
            solution["max_positive_average_change_kgkg"]
        ),
        "max_positive_sampled_change_kgkg": max_positive_sampled,
        "max_abs_step_change_kgkg": float(solution["max_abs_step_change_kgkg"]),
        "max_radial_increase_kgkg": float(solution["max_radial_increase_kgkg"]),
        "max_picard_iterations": int(np.max(iterations)),
        "mean_picard_iterations": float(np.mean(iterations)),
        "sample_count": int(times.size),
        "internal_node_count": int(upper_C.size),
        "stable_temperature_C": float(solution["stable_temperature_C"]),
        "stable_moisture_kgkg": float(solution["stable_moisture_kgkg"]),
        "stable_count": int(solution["stable_count"]),
        "early_steps": int(solution["early_steps"]),
        "long_steps": int(solution["long_steps"]),
    }


def compare_with_task2(
    solution: dict[str, Any],
    task2_csv_path: str | Path = DEFAULT_TASK2_CSV_PATH,
) -> dict[str, Any]:
    """将本轮N=1600 CN结果的0--3 h 1 s温度/水分场与Task2对照。"""

    first3_temperature = solution["first3_temperature_C"]
    first3_moisture = solution["first3_moisture_kgkg"]
    if (
        first3_temperature is None
        or first3_moisture is None
        or not np.all(np.isfinite(first3_temperature))
        or not np.all(np.isfinite(first3_moisture))
    ):
        raise ValueError("缺少问题三0--3 h连续场")
    frame = pd.read_csv(Path(task2_csv_path))
    required = {"time_s", "radius_cm", "temperature_C", "moisture_kgkg"}
    if not required.issubset(frame.columns):
        raise ValueError("问题二CSV缺少比较字段")
    expected_rows = 10801 * 21
    if len(frame) != expected_rows:
        raise ValueError(f"问题二CSV记录数为{len(frame)}，不是{expected_rows}")
    task2_times = frame["time_s"].to_numpy(dtype=float)
    task2_radii = frame["radius_cm"].to_numpy(dtype=float)
    expected_times = np.repeat(np.arange(10801, dtype=float), 21)
    expected_radii = np.tile(OUTPUT_RADIUS_CM, 10801)
    if not np.array_equal(task2_times, expected_times):
        raise ValueError("问题二CSV时间顺序不是0--10800 s逐秒")
    if not np.allclose(task2_radii, expected_radii, atol=1.0e-12):
        raise ValueError("问题二CSV半径顺序不是0--2 cm每0.1 cm")
    task2_temperature = frame["temperature_C"].to_numpy(dtype=float).reshape((10801, 21))
    task2_moisture = frame["moisture_kgkg"].to_numpy(dtype=float).reshape((10801, 21))
    raw_temperature_diff = np.abs(first3_temperature - task2_temperature)
    raw_moisture_diff = np.abs(first3_moisture - task2_moisture)
    rounded_temperature_diff = np.abs(
        np.round(first3_temperature, 4) - np.round(task2_temperature, 4)
    )
    rounded_moisture_diff = np.abs(
        np.round(first3_moisture, 4) - np.round(task2_moisture, 4)
    )
    return {
        "rows_compared": expected_rows,
        "max_abs_difference_raw_temperature_C": float(np.max(raw_temperature_diff)),
        "max_abs_difference_4dp_temperature_C": float(
            np.max(rounded_temperature_diff)
        ),
        "max_abs_difference_raw_moisture_kgkg": float(np.max(raw_moisture_diff)),
        "max_abs_difference_4dp_moisture_kgkg": float(
            np.max(rounded_moisture_diff)
        ),
    }


def table5_from_solution(solution: dict[str, Any]) -> tuple[np.ndarray, np.ndarray]:
    """返回表5常规6 h行和阈值二分右端点行。"""

    regular_times = np.asarray(solution["table5_times_s"], dtype=float)
    regular_values = np.asarray(solution["table5_values_kgkg"], dtype=float)
    event = np.asarray(solution["termination_event_moisture_kgkg"], dtype=float)
    event_indices = np.asarray(solution["key_indices"], dtype=int)
    event_values = event[event_indices]
    if regular_values.ndim != 2 or regular_values.shape[1] != KEY_RADIUS_INDICES.size:
        raise ValueError("表5常规行形状错误")
    return regular_times, np.vstack((regular_values, event_values))


def print_table5(solution: dict[str, Any]) -> None:
    """打印论文表5：6 h行至结束前，以及阈值二分右端点行。"""

    regular_times, rows = table5_from_solution(solution)
    termination_time_h = float(solution["termination_time_s"]) / 3600.0
    print("表5 水分浓度 / kg/kg（列为0、0.5、1.0、1.5、2.0 cm）")
    print("时间/h,0 cm,0.5 cm,1.0 cm,1.5 cm,2.0 cm")
    for index, time_value in enumerate(regular_times):
        print(
            f"{time_value / 3600.0:.1f},"
            + ",".join(f"{float(value):.4f}" for value in rows[index])
        )
    print(
        f"烘干结束时间 ({termination_time_h:.6f} h),"
        + ",".join(f"{float(value):.4f}" for value in rows[-1])
    )


def write_result_xlsx(
    solution: dict[str, Any],
    output_path: str | Path = DEFAULT_XLSX_PATH,
    template_path: str | Path = DEFAULT_TEMPLATE_PATH,
) -> None:
    """以result3模板为基础写入60 s水分结果。"""

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(Path(template_path))
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError(f"result3模板工作表错误: {workbook.sheetnames}")
    worksheet = workbook["Sheet1"]
    times = np.asarray(solution["time_s"], dtype=float)
    radii = np.asarray(solution["radius_cm"], dtype=float)
    values = np.asarray(solution["moisture_kgkg"], dtype=float)
    header_style = copy(worksheet.cell(row=1, column=2)._style)
    time_style = copy(worksheet.cell(row=2, column=1)._style)
    value_style = copy(worksheet.cell(row=2, column=2)._style)
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)
    for column, radius in enumerate(radii, start=2):
        cell = worksheet.cell(row=1, column=column, value=round(float(radius), 1))
        cell._style = copy(header_style)
        cell.number_format = "0.0"
    for row, time_value in enumerate(times, start=2):
        time_cell = worksheet.cell(row=row, column=1, value=int(round(time_value)))
        time_cell._style = copy(time_style)
        time_cell.number_format = "0"
        for column, value in enumerate(values[row - 2], start=2):
            cell = worksheet.cell(
                row=row,
                column=column,
                value=float(f"{float(value):.4f}"),
            )
            cell._style = copy(value_style)
            cell.number_format = "0.0000"
    workbook.save(output_path)


def write_result_csv(
    solution: dict[str, Any],
    output_path: str | Path = DEFAULT_CSV_PATH,
) -> None:
    """写入题目要求的三列long-format CSV。"""

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    times = np.asarray(solution["time_s"], dtype=float)
    radii = np.asarray(solution["radius_cm"], dtype=float)
    values = np.asarray(solution["moisture_kgkg"], dtype=float)
    with output_path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(["time_s", "radius_cm", "moisture_kgkg"])
        for time_index, time_value in enumerate(times):
            for radius_index, radius_value in enumerate(radii):
                writer.writerow(
                    [
                        int(round(time_value)),
                        f"{radius_value:.1f}",
                        f"{values[time_index, radius_index]:.4f}",
                    ]
                )


def validate_written_outputs(
    solution: dict[str, Any],
    xlsx_path: str | Path = DEFAULT_XLSX_PATH,
    csv_path: str | Path = DEFAULT_CSV_PATH,
) -> dict[str, Any]:
    """回读CSV/XLSX，检查结构和全部对应水分值。"""

    expected = np.round(np.asarray(solution["moisture_kgkg"], dtype=float), 4)
    expected_times = np.asarray(solution["time_s"], dtype=int)
    expected_radii = np.asarray(solution["radius_cm"], dtype=float)
    frame = pd.read_csv(Path(csv_path))
    expected_columns = ["time_s", "radius_cm", "moisture_kgkg"]
    if list(frame.columns) != expected_columns:
        raise ValueError(f"CSV字段错误: {list(frame.columns)}")
    expected_rows = expected_times.size * expected_radii.size
    if len(frame) != expected_rows:
        raise ValueError(f"CSV行数为{len(frame)}，不是{expected_rows}")
    csv_times = frame["time_s"].to_numpy(dtype=int)
    csv_radii = frame["radius_cm"].to_numpy(dtype=float)
    csv_values = frame["moisture_kgkg"].to_numpy(dtype=float).reshape(expected.shape)
    if not np.array_equal(csv_times, np.repeat(expected_times, expected_radii.size)):
        raise ValueError("CSV时间结构错误")
    if not np.allclose(csv_radii, np.tile(expected_radii, expected_times.size), atol=1.0e-12):
        raise ValueError("CSV半径结构错误")

    workbook = load_workbook(Path(xlsx_path), read_only=True, data_only=True)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError(f"XLSX工作表错误: {workbook.sheetnames}")
    worksheet = workbook["Sheet1"]
    if (worksheet.max_row, worksheet.max_column) != (
        expected_times.size + 1,
        expected_radii.size + 1,
    ):
        raise ValueError("XLSX尺寸错误")
    rows = list(
        worksheet.iter_rows(
            min_row=1,
            max_row=worksheet.max_row,
            min_col=1,
            max_col=worksheet.max_column,
            values_only=True,
        )
    )
    xlsx_times = np.asarray([row[0] for row in rows[1:]], dtype=int)
    xlsx_radii = np.asarray(rows[0][1:], dtype=float)
    xlsx_values = np.asarray([row[1:] for row in rows[1:]], dtype=float)
    if not np.array_equal(xlsx_times, expected_times) or not np.allclose(
        xlsx_radii, expected_radii, atol=1.0e-12
    ):
        raise ValueError("XLSX时间或半径结构错误")
    xlsx_values = np.round(xlsx_values, 4)
    if not np.array_equal(csv_values, xlsx_values):
        raise ValueError("CSV与XLSX对应值不一致")
    if not np.array_equal(csv_values, expected):
        raise ValueError("答案文件与求解器采样值不一致")
    return {
        "csv_rows": int(len(frame)),
        "xlsx_shape": (int(worksheet.max_row), int(worksheet.max_column)),
        "max_abs_csv_xlsx_difference_kgkg": float(np.max(np.abs(csv_values - xlsx_values))),
        "max_abs_file_solver_difference_kgkg": float(np.max(np.abs(csv_values - expected))),
    }


def sha256_file(path: str | Path) -> dict[str, Any]:
    """返回用于复现审计的文件大小、路径和SHA-256。"""

    destination = Path(path)
    if not destination.is_file():
        raise FileNotFoundError(f"诊断哈希目标不存在: {destination}")
    digest = hashlib.sha256()
    with destination.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return {
        "path": str(destination.resolve()),
        "bytes": int(destination.stat().st_size),
        "sha256": digest.hexdigest(),
    }


def _profile_summary(
    radii_m: np.ndarray,
    temperature_C: np.ndarray,
    moisture_kgkg: np.ndarray,
) -> dict[str, Any]:
    """将一个完整内部事件剖面转为可直接持久化的JSON对象。"""

    temperature = np.asarray(temperature_C, dtype=float)
    moisture = np.asarray(moisture_kgkg, dtype=float)
    radii = np.asarray(radii_m, dtype=float)
    if not (radii.shape == temperature.shape == moisture.shape):
        raise ValueError("事件剖面半径、温度和水分数组形状不一致")
    return {
        "radius_m": radii.tolist(),
        "radius_cm": (radii * 100.0).tolist(),
        "temperature_C": temperature.tolist(),
        "moisture_kgkg": moisture.tolist(),
        "temperature_min_C": float(np.min(temperature)),
        "temperature_max_C": float(np.max(temperature)),
        "moisture_min_kgkg": float(np.min(moisture)),
        "moisture_max_kgkg": float(np.max(moisture)),
        "moisture_max_index": int(np.argmax(moisture)),
        "moisture_max_radius_m": float(radii[int(np.argmax(moisture))]),
        "max_radial_increase_kgkg": float(np.max(np.diff(moisture))),
    }


def diagnostics_summary(
    solution: dict[str, Any],
    data_path: str | Path = DEFAULT_DATA_PATH,
    task2_csv_path: str | Path = DEFAULT_TASK2_CSV_PATH,
    handoff_path: str | Path = DEFAULT_HANDOFF_PATH,
    solver_path: str | Path = Path(__file__),
    task2_comparison: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """汇总问题三 Picard、物理约束、守恒残差和精确阈值事件证据。"""

    time_s = np.asarray(solution["time_s"], dtype=float)
    radius_cm = np.asarray(solution["radius_cm"], dtype=float)
    moisture = np.asarray(solution["moisture_kgkg"], dtype=float)
    radii_m = np.asarray(solution["radii_m_internal"], dtype=float)
    event_temperature = np.asarray(
        solution["termination_event_temperature_C"], dtype=float
    )
    event_moisture = np.asarray(
        solution["termination_event_moisture_kgkg"], dtype=float
    )
    lower_temperature = np.asarray(
        solution["termination_lower_temperature_C"], dtype=float
    )
    lower_moisture = np.asarray(
        solution["termination_lower_moisture_kgkg"], dtype=float
    )
    upper_temperature = np.asarray(
        solution["termination_upper_temperature_C"], dtype=float
    )
    upper_moisture = np.asarray(
        solution["termination_upper_moisture_kgkg"], dtype=float
    )
    iterations = np.asarray(solution["iteration_counts"], dtype=int)
    delta_temperature = np.asarray(solution["picard_delta_temperature"], dtype=float)
    delta_moisture = np.asarray(solution["picard_delta_moisture"], dtype=float)
    step_times = np.asarray(solution["step_times_s"], dtype=float)
    step_dt = np.asarray(solution["step_dt_s"], dtype=float)
    temperature_linear = np.asarray(
        solution["temperature_linear_residuals"], dtype=float
    )
    moisture_linear = np.asarray(
        solution["moisture_linear_residuals"], dtype=float
    )
    linear_residuals = np.asarray(solution["linear_residuals"], dtype=float)
    temperature_balance = np.asarray(
        solution["temperature_balance_residuals"], dtype=float
    )
    moisture_balance = np.asarray(
        solution["moisture_balance_residuals"], dtype=float
    )
    temperature_fluxes = np.asarray(
        solution["temperature_boundary_fluxes"], dtype=float
    )
    moisture_fluxes = np.asarray(solution["moisture_boundary_fluxes"], dtype=float)
    volume_average = np.asarray(solution["volume_average_history"], dtype=float)
    air_time = np.asarray(solution["air_time_s"], dtype=float)
    air_temperature = np.asarray(solution["air_temperature_C"], dtype=float)
    air_moisture = np.asarray(solution["air_moisture_kgkg"], dtype=float)

    sample_air = np.asarray(
        [
            air_at(
                float(value),
                air_time,
                air_temperature,
                air_moisture,
                float(solution["stable_temperature_C"]),
                float(solution["stable_moisture_kgkg"]),
            )
            for value in time_s
        ],
        dtype=float,
    )
    # Task3只保存60 s水分场；边界水分方向可由其表面值直接核验。
    sample_moisture_driver = sample_air[:, 1] - moisture[:, -1]
    event_air = air_at(
        float(solution["termination_time_s"]),
        air_time,
        air_temperature,
        air_moisture,
        float(solution["stable_temperature_C"]),
        float(solution["stable_moisture_kgkg"]),
    )
    first3_temperature = solution["first3_temperature_C"]
    if first3_temperature is not None:
        first3_temperature = np.asarray(first3_temperature, dtype=float)
        first3_air = np.asarray(
            [
                air_at(
                    float(value),
                    air_time,
                    air_temperature,
                    air_moisture,
                    float(solution["stable_temperature_C"]),
                    float(solution["stable_moisture_kgkg"]),
                )
                for value in np.arange(first3_temperature.shape[0], dtype=float)
            ],
            dtype=float,
        )
        first3_heat_driver = first3_air[:, 0] - first3_temperature[:, -1]
    else:
        first3_heat_driver = np.empty(0, dtype=float)
    volume_measure = float(np.sum(fixed_geometry(int(solution["n_intervals"]))[2]))
    event_average = float(np.sum(fixed_geometry(int(solution["n_intervals"]))[2] * event_moisture) / volume_measure)
    source_hashes = {
        "solver": sha256_file(solver_path),
        "handoff": sha256_file(handoff_path),
        "air_attachment": sha256_file(data_path),
        "task2_result_csv": sha256_file(task2_csv_path),
    }
    generated_outputs: dict[str, Any] = {}
    for label, path in (
        ("result3_xlsx", DEFAULT_XLSX_PATH),
        ("task3_result_csv", DEFAULT_CSV_PATH),
    ):
        if Path(path).is_file():
            generated_outputs[label] = sha256_file(path)

    def _relative(
        residuals: np.ndarray, fluxes: np.ndarray
    ) -> tuple[float, float]:
        if residuals.size == 0:
            return 0.0, 0.0
        # Use one characteristic boundary-flux scale for the whole run.
        # A per-step denominator reports a misleading relative value of 1
        # whenever both the flux and roundoff residual are near zero.
        scale = max(
            float(np.max(np.abs(fluxes))),
            float(np.max(np.abs(residuals))),
            np.finfo(float).tiny,
        )
        relative = np.abs(residuals) / scale
        index = int(np.argmax(np.abs(residuals)))
        return float(np.max(relative)), float(step_times[index])

    moisture_relative, moisture_worst_time = _relative(
        moisture_balance, moisture_fluxes
    )
    temperature_relative, temperature_worst_time = _relative(
        temperature_balance, temperature_fluxes
    )
    event_profile = _profile_summary(radii_m, event_temperature, event_moisture)
    left_profile = _profile_summary(radii_m, lower_temperature, lower_moisture)
    right_profile = _profile_summary(radii_m, upper_temperature, upper_moisture)
    event_max_index = int(np.argmax(event_moisture))
    iteration_values, iteration_counts_hist = np.unique(iterations, return_counts=True)
    comparison = task2_comparison or solution.get("task2_comparison") or {}
    return {
        "task": 3,
        "source_hashes": source_hashes,
        "input_sources": source_hashes,
        "generated_outputs": generated_outputs,
        "caliber": {
            "spatial_scheme": "守恒型径向有限体积（中心/表面半控制体积）",
            "temporal_scheme": "Crank-Nicolson",
            "interface_property": "调和平均",
            "nonlinear_scheme": "同步半时间层 Picard",
            "n_intervals": int(solution["n_intervals"]),
            "internal_node_count": int(radii_m.size),
            "dr_m": float(solution["dr_m"]),
            "dt_early_s": float(solution["time_step_early_s"]),
            "dt_long_s": float(solution["time_step_long_s"]),
            "output_dt_s": float(solution["output_dt_s"]),
            "step_count": int(iterations.size),
            "early_steps": int(solution["early_steps"]),
            "long_steps": int(solution["long_steps"]),
            "sample_count": int(time_s.size),
            "output_shape": [int(moisture.shape[0]), int(moisture.shape[1])],
            "first3_temperature_shape": list(
                np.asarray(solution["first3_temperature_C"]).shape
            ),
            "first3_moisture_shape": list(
                np.asarray(solution["first3_moisture_kgkg"]).shape
            ),
            "picard_tol_temperature_C": float(PICARD_TOL_T_C),
            "picard_tol_moisture_kgkg": float(PICARD_TOL_C_KGKG),
            "linear_residual_tol": float(LINEAR_RESIDUAL_TOL),
            "threshold_kgkg": float(DRYING_THRESHOLD_KGKG),
            "stable_start_s": float(STABLE_START_S),
            "air_data_end_s": float(AIR_DATA_END_S),
        },
        "inputs": {
            "stable_temperature_C": float(solution["stable_temperature_C"]),
            "stable_moisture_kgkg": float(solution["stable_moisture_kgkg"]),
            "stable_count": int(solution["stable_count"]),
            "initial_temperature_C": float(INITIAL_TEMPERATURE_C),
            "initial_moisture_kgkg": float(INITIAL_MOISTURE_KGKG),
            "radius_m": float(RADIUS_M),
            "length_m": float(LENGTH_M),
        },
        "picard": {
            "max_iterations": int(np.max(iterations)),
            "mean_iterations": float(np.mean(iterations)),
            "iteration_histogram": {
                str(int(value)): int(count)
                for value, count in zip(iteration_values, iteration_counts_hist)
            },
            "max_final_delta_temperature_C": float(np.max(delta_temperature)),
            "max_final_delta_moisture_kgkg": float(np.max(delta_moisture)),
            "converged_steps": int(
                np.sum(
                    (delta_temperature <= PICARD_TOL_T_C)
                    & (delta_moisture <= PICARD_TOL_C_KGKG)
                )
            ),
            "failed_steps": int(
                np.sum(
                    (delta_temperature > PICARD_TOL_T_C)
                    | (delta_moisture > PICARD_TOL_C_KGKG)
                )
            ),
            "bisection_reintegrations": int(
                solution["termination_bisection_iterations"]
            ),
        },
        "linear_solver": {
            "max_temperature_normalized_inf_residual": float(
                np.max(temperature_linear)
            ),
            "max_moisture_normalized_inf_residual": float(
                np.max(moisture_linear)
            ),
            "max_normalized_inf_residual": float(np.max(linear_residuals)),
            "tol": float(LINEAR_RESIDUAL_TOL),
        },
        "physics": {
            "finite": bool(
                all(
                    np.all(np.isfinite(array))
                    for array in (
                        moisture,
                        event_temperature,
                        event_moisture,
                        lower_temperature,
                        lower_moisture,
                        upper_temperature,
                        upper_moisture,
                    )
                )
            ),
            "temperature_min_C": float(solution["min_temperature_C"]),
            "temperature_max_C": float(solution["max_temperature_C"]),
            "moisture_min_kgkg": float(solution["min_internal_moisture_kgkg"]),
            "moisture_max_kgkg": float(solution["max_internal_moisture_kgkg"]),
            "moisture_nonnegative": bool(
                solution["min_internal_moisture_kgkg"] >= -1.0e-10
            ),
            "moisture_bounded_by_initial": bool(
                solution["max_internal_moisture_kgkg"]
                <= INITIAL_MOISTURE_KGKG + 1.0e-10
            ),
            "center_inner_face_radius_m": 0.0,
            "center_zero_flux_discretisation": True,
            "volume_average_moisture_initial_kgkg": float(volume_average[0]),
            "volume_average_moisture_final_step_kgkg": float(volume_average[-1]),
            "volume_average_moisture_event_kgkg": event_average,
            "volume_average_moisture_monotone_decreasing": bool(
                np.all(np.diff(volume_average) <= 1.0e-12)
            ),
            "max_positive_step_change_kgkg": float(
                solution["max_positive_step_change_kgkg"]
            ),
            "max_positive_average_change_kgkg": float(
                solution["max_positive_average_change_kgkg"]
            ),
            "max_radial_increase_kgkg": float(solution["max_radial_increase_kgkg"]),
            "radial_moisture_nonincreasing": bool(
                solution["max_radial_increase_kgkg"] <= 1.0e-10
            ),
            "event_radial_moisture_nonincreasing": bool(
                np.max(np.diff(event_moisture)) <= 1.0e-10
            ),
            "sampled_moisture_boundary_driver_max_kgkg": float(
                np.max(sample_moisture_driver[1:])
            ),
            "sampled_moisture_outward_fraction": float(
                np.mean(sample_moisture_driver[1:] <= 1.0e-8)
            ),
            "temperature_boundary_driver_min_first3_C": float(
                np.min(first3_heat_driver[1:])
            )
            if first3_heat_driver.size > 1
            else 0.0,
            "temperature_inward_fraction_first3": float(
                np.mean(first3_heat_driver[1:] >= -1.0e-8)
            )
            if first3_heat_driver.size > 1
            else 0.0,
            "event_temperature_boundary_driver_C": float(
                event_air[0] - event_temperature[-1]
            ),
            "event_temperature_center_surface_gap_C": float(
                event_temperature[-1] - event_temperature[0]
            ),
            "event_air_temperature_C": float(event_air[0]),
            "event_air_moisture_kgkg": float(event_air[1]),
        },
        "conservation": {
            "coordinate_measure": "约去2πL；Σ V_i C_i和Σ V_i(ρc_p)_mid T_i",
            "moisture": {
                "max_abs_balance_residual": float(np.max(np.abs(moisture_balance))),
                "max_relative_balance_residual": moisture_relative,
                "worst_time_s": moisture_worst_time,
                "boundary_flux_integral": float(np.sum(moisture_fluxes * step_dt)),
                "event_max_abs_balance_residual": float(
                    abs(solution["termination_event_moisture_balance_residual"])
                ),
            },
            "energy": {
                "max_abs_balance_residual": float(
                    np.max(np.abs(temperature_balance))
                ),
                "max_relative_balance_residual": temperature_relative,
                "worst_time_s": temperature_worst_time,
                "boundary_flux_integral": float(
                    np.sum(temperature_fluxes * step_dt)
                ),
                "event_max_abs_balance_residual": float(
                    abs(solution["termination_event_temperature_balance_residual"])
                ),
            },
        },
        "termination": {
            "event_state_definition": "严格满足max(C)<0.15的二分右端点t_R；非跨步线性插值",
            "time_s": float(solution["termination_time_s"]),
            "time_h": float(solution["termination_time_s"]) / 3600.0,
            "lower_time_s": float(solution["termination_lower_time_s"]),
            "upper_time_s": float(solution["termination_upper_time_s"]),
            "bracket_width_s": float(
                solution["termination_upper_time_s"]
                - solution["termination_lower_time_s"]
            ),
            "crossing_base_time_s": float(solution["termination_crossing_base_time_s"]),
            "crossing_step_width_s": float(
                solution["termination_crossing_bracket_width_s"]
            ),
            "bisection_iterations": int(solution["termination_bisection_iterations"]),
            "threshold_kgkg": float(DRYING_THRESHOLD_KGKG),
            "lower_max_kgkg": float(solution["termination_lower_max_kgkg"]),
            "upper_max_kgkg": float(solution["termination_upper_max_kgkg"]),
            "event_max_kgkg": float(np.max(event_moisture)),
            "event_max_index": event_max_index,
            "event_max_radius_m": float(radii_m[event_max_index]),
            "event_center_kgkg": float(event_moisture[0]),
            "event_surface_kgkg": float(event_moisture[-1]),
            "left_profile": left_profile,
            "right_profile": right_profile,
            "event_profile": event_profile,
            "event_moisture_boundary_flux": float(
                solution["termination_event_moisture_boundary_flux"]
            ),
            "event_temperature_boundary_flux": float(
                solution["termination_event_temperature_boundary_flux"]
            ),
        },
        "task2_consistency": comparison,
    }


def write_diagnostics(
    solution: dict[str, Any],
    json_path: str | Path = DEFAULT_DIAGNOSTICS_JSON,
    csv_path: str | Path = DEFAULT_DIAGNOSTICS_CSV,
    data_path: str | Path = DEFAULT_DATA_PATH,
    task2_csv_path: str | Path = DEFAULT_TASK2_CSV_PATH,
    handoff_path: str | Path = DEFAULT_HANDOFF_PATH,
    solver_path: str | Path = Path(__file__),
    task2_comparison: dict[str, Any] | None = None,
) -> tuple[Path, Path]:
    """写入逐主步诊断CSV及含事件剖面的汇总JSON。"""

    summary = diagnostics_summary(
        solution,
        data_path=data_path,
        task2_csv_path=task2_csv_path,
        handoff_path=handoff_path,
        solver_path=solver_path,
        task2_comparison=task2_comparison,
    )
    json_destination = Path(json_path)
    json_destination.parent.mkdir(parents=True, exist_ok=True)
    json_destination.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    step_times = np.asarray(solution["step_times_s"], dtype=float)
    step_dt = np.asarray(solution["step_dt_s"], dtype=float)
    iterations = np.asarray(solution["iteration_counts"], dtype=int)
    delta_temperature = np.asarray(solution["picard_delta_temperature"], dtype=float)
    delta_moisture = np.asarray(solution["picard_delta_moisture"], dtype=float)
    temperature_linear = np.asarray(
        solution["temperature_linear_residuals"], dtype=float
    )
    moisture_linear = np.asarray(
        solution["moisture_linear_residuals"], dtype=float
    )
    linear_residuals = np.asarray(solution["linear_residuals"], dtype=float)
    temperature_balance = np.asarray(
        solution["temperature_balance_residuals"], dtype=float
    )
    moisture_balance = np.asarray(
        solution["moisture_balance_residuals"], dtype=float
    )
    temperature_fluxes = np.asarray(
        solution["temperature_boundary_fluxes"], dtype=float
    )
    moisture_fluxes = np.asarray(solution["moisture_boundary_fluxes"], dtype=float)
    csv_destination = Path(csv_path)
    csv_destination.parent.mkdir(parents=True, exist_ok=True)
    with csv_destination.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(
            (
                "step",
                "time_s",
                "dt_s",
                "picard_iterations",
                "picard_delta_temperature_C",
                "picard_delta_moisture_kgkg",
                "temperature_linear_normalized_inf_residual",
                "moisture_linear_normalized_inf_residual",
                "linear_normalized_inf_residual",
                "temperature_boundary_flux",
                "temperature_balance_residual",
                "moisture_boundary_flux",
                "moisture_balance_residual",
            )
        )
        for index in range(iterations.size):
            writer.writerow(
                (
                    index + 1,
                    f"{step_times[index]:.6f}",
                    f"{step_dt[index]:.6f}",
                    int(iterations[index]),
                    f"{delta_temperature[index]:.6e}",
                    f"{delta_moisture[index]:.6e}",
                    f"{temperature_linear[index]:.6e}",
                    f"{moisture_linear[index]:.6e}",
                    f"{linear_residuals[index]:.6e}",
                    f"{temperature_fluxes[index]:.6e}",
                    f"{temperature_balance[index]:.6e}",
                    f"{moisture_fluxes[index]:.6e}",
                    f"{moisture_balance[index]:.6e}",
                )
            )
    return json_destination, csv_destination


def compare_sensitivity_tables(
    base: dict[str, Any], other: dict[str, Any]
) -> dict[str, float]:
    """比较两个解在相同6 h表5行上的终止时刻和水分值。"""

    base_times, base_rows = table5_from_solution(base)
    other_times, other_rows = table5_from_solution(other)
    if base_times.size != other_times.size or not np.allclose(
        base_times, other_times, rtol=0.0, atol=1.0e-8
    ):
        raise ValueError("两个解的表5常规时间行不一致")
    common = min(base_rows.shape[0], other_rows.shape[0])
    if common == 0:
        table_difference = 0.0
    else:
        table_difference = float(np.max(np.abs(base_rows[:common] - other_rows[:common])))
    return {
        "termination_time_difference_h": float(
            other["termination_time_s"] - base["termination_time_s"]
        )
        / 3600.0,
        "table5_common_max_abs_difference_kgkg": table_difference,
        "common_table_rows": int(common),
    }


def run_grid_sensitivity(
    data_path: str | Path,
    max_time_s: float,
    progress: bool = False,
) -> tuple[dict[int, dict[str, Any]], dict[int, dict[str, Any]]]:
    """运行N=200/400/800/1600的同一CN模型空间加密检查。"""

    grid_solutions: dict[int, dict[str, Any]] = {}
    grid_summaries: dict[int, dict[str, Any]] = {}
    for n in (200, 400, 800, 1600):
        started = wall_time.perf_counter()
        candidate = solve_task3(
            data_path=data_path,
            n_intervals=n,
            time_step_long_s=TIME_STEP_LONG_DEFAULT_S,
            output_dt_s=OUTPUT_DT_S,
            max_time_s=max_time_s,
            progress=progress,
            collect_first3=False,
        )
        candidate_summary = validate_solution(candidate, require_first3=False)
        grid_solutions[n] = candidate
        grid_summaries[n] = {
            **candidate_summary,
            "runtime_s": wall_time.perf_counter() - started,
        }
        print(
            f"空间网格 N={n}: t_dry={candidate_summary['termination_time_h']:.8f} h, "
            f"耗时={grid_summaries[n]['runtime_s']:.2f} s"
        )
    return grid_solutions, grid_summaries


def print_grid_sensitivity(grid_summaries: dict[int, dict[str, Any]]) -> None:
    print("空间网格收敛（同一CN/调和平均/1 s+10 s模型）")
    previous = None
    for n in sorted(grid_summaries):
        summary = grid_summaries[n]
        difference = "—" if previous is None else f"{summary['termination_time_h'] - previous:.8f}"
        print(
            f"N={n:4d}, dr={summary['dr_cm']:.5f} cm, "
            f"t_dry={summary['termination_time_h']:.8f} h, "
            f"相对前一网格={difference} h"
        )
        previous = summary["termination_time_h"]


def main() -> None:
    parser = argparse.ArgumentParser(description="求解A题问题三CN长期烘干模型")
    parser.add_argument("--data-path", type=Path, default=DEFAULT_DATA_PATH)
    parser.add_argument("--task2-csv", type=Path, default=DEFAULT_TASK2_CSV_PATH)
    parser.add_argument("--n", type=int, default=N_DEFAULT)
    parser.add_argument("--time-step-long", type=float, default=TIME_STEP_LONG_DEFAULT_S)
    parser.add_argument("--max-time", type=float, default=DEFAULT_MAX_TIME_S)
    parser.add_argument("--no-write", action="store_true", help="不生成答案文件")
    parser.add_argument("--skip-sensitivity", action="store_true", help="跳过N/长期时间步敏感性")
    parser.add_argument("--progress", action="store_true", help="按6 h打印推进进度")
    args = parser.parse_args()

    started = wall_time.perf_counter()
    print(
        "问题三CN求解开始: "
        f"N={args.n}, dr={RADIUS_M / args.n:g}m, "
        f"dt_early={TIME_STEP_EARLY_S:g}s, dt_long={args.time_step_long:g}s, "
        f"output_dt={OUTPUT_DT_S:g}s, threshold C<{DRYING_THRESHOLD_KGKG:g}"
    )
    solution = solve_task3(
        data_path=args.data_path,
        n_intervals=args.n,
        time_step_long_s=args.time_step_long,
        output_dt_s=OUTPUT_DT_S,
        max_time_s=args.max_time,
        progress=args.progress,
        collect_first3=True,
    )
    summary = validate_solution(solution, require_first3=True)
    air_time, air_temperature, air_moisture = load_air_data(args.data_path)
    stable_T, stable_C, stable_count = stable_air_means(air_time, air_temperature, air_moisture)
    print(
        f"稳定段均值（附件实时计算，{STABLE_START_S:g}--{AIR_DATA_END_S:g}s，"
        f"{stable_count}点）: T_air={stable_T:.12f} C, C_air={stable_C:.12f} kg/kg"
    )
    print(f"基本数值检查: {summary}")
    comparison = compare_with_task2(solution, args.task2_csv)
    if (
        comparison["max_abs_difference_4dp_temperature_C"] > 1.0e-12
        or comparison["max_abs_difference_4dp_moisture_kgkg"] > 1.0e-12
    ):
        raise RuntimeError(
            "问题三0--3 h与问题二CSV的四位小数结果不一致: "
            f"dT={comparison['max_abs_difference_4dp_temperature_C']:.3e}, "
            f"dC={comparison['max_abs_difference_4dp_moisture_kgkg']:.3e}"
        )
    solution["task2_comparison"] = comparison
    print(f"0--3 h连续状态与既有task2 CSV全量比较: {comparison}")

    if not args.skip_sensitivity:
        _grid_solutions, grid_summaries = run_grid_sensitivity(
            args.data_path, args.max_time, progress=False
        )
        print_grid_sensitivity(grid_summaries)
        long5_started = wall_time.perf_counter()
        long5_solution = solve_task3(
            data_path=args.data_path,
            n_intervals=args.n,
            time_step_long_s=5.0,
            output_dt_s=OUTPUT_DT_S,
            max_time_s=args.max_time,
            progress=False,
            collect_first3=False,
        )
        long5_summary = validate_solution(long5_solution, require_first3=False)
        long5_difference = compare_sensitivity_tables(solution, long5_solution)
        print(
            "N={:d} 长期时间步敏感性: dt_long=10 s -> 5 s, "
            "t_dry={:.8f} h -> {:.8f} h, 差={:.8f} h, "
            "表5公共行最大差={:.8g} kg/kg, 耗时={:.2f} s".format(
                args.n,
                summary["termination_time_h"],
                long5_summary["termination_time_h"],
                long5_difference["termination_time_difference_h"],
                long5_difference["table5_common_max_abs_difference_kgkg"],
                wall_time.perf_counter() - long5_started,
            )
        )
    else:
        print("已跳过空间网格和长期时间步敏感性")

    if not args.no_write:
        template_hash_before = hashlib.sha256(DEFAULT_TEMPLATE_PATH.read_bytes()).hexdigest()
        write_result_xlsx(solution)
        write_result_csv(solution)
        output_summary = validate_written_outputs(solution)
        diagnostics_json, diagnostics_csv = write_diagnostics(
            solution,
            data_path=args.data_path,
            task2_csv_path=args.task2_csv,
            handoff_path=DEFAULT_HANDOFF_PATH,
            solver_path=Path(__file__),
            task2_comparison=comparison,
        )
        template_hash_after = hashlib.sha256(DEFAULT_TEMPLATE_PATH.read_bytes()).hexdigest()
        if template_hash_before != template_hash_after:
            raise RuntimeError("result3模板在输出过程中被修改")
        print(f"答案文件回读验收: {output_summary}")
        print(f"模板SHA-256（前后相同）: {template_hash_after}")
        print(f"已生成: {DEFAULT_XLSX_PATH}")
        print(f"已生成: {DEFAULT_CSV_PATH}")
        print(f"诊断日志: {diagnostics_json}, {diagnostics_csv}")
    else:
        print("--no-write: 未生成或覆盖答案文件")

    print_table5(solution)
    print(f"烘干结束时间 = {summary['termination_time_h']:.8f} h")
    print(f"总耗时: {wall_time.perf_counter() - started:.2f} s")


if __name__ == "__main__":
    main()
