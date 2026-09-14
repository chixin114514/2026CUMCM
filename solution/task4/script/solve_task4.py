"""A题问题四：材料坐标下的收缩圆柱热湿耦合求解器。

药材半径由附件2给出，材料坐标为 ``xi=r/R(t)``。在固定 ``xi`` 网格上，
温度和干基含水率采用守恒型径向有限体积离散；时间方向为冻结中点物性
Crank--Nicolson 型格式，每个时间步使用同步 Picard 迭代。

材料导数在材料坐标中已经等于固定 xi 导数，因此不再加入旧脚本中的
Rdot 迎风项；半径变化通过时间中点半径、空间尺度和移动表面边界进入方程。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import time as wall_time
from copy import copy
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from scipy.linalg import solve_banded


RADIUS_INITIAL_M = 0.02
LENGTH_M = 0.25
INITIAL_TEMPERATURE_C = 28.0
INITIAL_MOISTURE_KGKG = 2.55
HEAT_TRANSFER_W_M2_K = 25.0
MASS_TRANSFER_M_S = 8.0e-7

N_DEFAULT = 800
TIME_STEP_EARLY_S = 1.0
TIME_STEP_LONG_DEFAULT_S = 5.0
TIME_STEP_S_DEFAULT = TIME_STEP_LONG_DEFAULT_S
OUTPUT_DT_S = 60.0
STABLE_START_S = 9000.0
AIR_DATA_END_S = 14400.0
DRYING_THRESHOLD_KGKG = 0.15
EVENT_TOL_S = 0.1
DEFAULT_MAX_TIME_S = 3.0 * 24.0 * 3600.0
FIXED_MAX_TIME_S = 7.0 * 24.0 * 3600.0
PICARD_TOL_T_C = 1.0e-8
PICARD_TOL_C_KGKG = 1.0e-10
PICARD_MAX_ITER = 50
LINEAR_RESIDUAL_TOL = 1.0e-9
STATE_TOL = 1.0e-8

REPO_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_AIR_DATA_PATH = REPO_ROOT / "problem" / "附件" / "附件1.xlsx"
DEFAULT_RADIUS_DATA_PATH = REPO_ROOT / "problem" / "附件" / "附件2.xlsx"
DEFAULT_TEMPLATE_PATH = REPO_ROOT / "problem" / "附件" / "附件3" / "result4.xlsx"
DEFAULT_ANSWER_DIR = Path(__file__).resolve().parents[1] / "answer"
DEFAULT_XLSX_PATH = DEFAULT_ANSWER_DIR / "result4.xlsx"
DEFAULT_PAPER_PATH = DEFAULT_ANSWER_DIR / "task4_paper_table.xlsx"
DEFAULT_CSV_PATH = DEFAULT_ANSWER_DIR / "task4_all_answers.csv"
DEFAULT_DIAGNOSTICS_JSON = DEFAULT_ANSWER_DIR / "task4_diagnostics.json"
DEFAULT_DIAGNOSTICS_CSV = DEFAULT_ANSWER_DIR / "task4_diagnostics.csv"
DEFAULT_MECHANISM_JSON = DEFAULT_ANSWER_DIR / "task4_mechanisms.json"

OUTPUT_RADIUS_CM = np.arange(0.0, 2.0000001, 0.1)
CSV_COLUMNS = [
    "record_type", "time_s", "time_h", "radius_cm", "r_cm",
    "is_surface", "temperature_C", "moisture_kgkg",
]


def _file_sha256(path: str | Path) -> str:
    """Return the SHA-256 of an input file for reproducibility metadata."""
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _input_sources(
    air_data_path: str | Path,
    radius_data_path: str | Path,
    appendix: int,
    fixed_radius_m: float | None,
) -> dict[str, Any]:
    """Describe the only external inputs used by one Task4 solve.

    Fixed-radius mechanism A/B intentionally does not read Attachment 2.  The
    property formulas are the question-statement Appendix 3/4 formulas and
    are implemented only in :func:`material_properties` below.
    """
    air_path = Path(air_data_path)
    sources: dict[str, Any] = {
        "air_attachment": {
            "path": str(air_path.resolve()),
            "sha256": _file_sha256(air_path),
        },
        "property_formula": {
            "appendix": int(appendix),
            "source": f"题面附录{appendix}公式；solve_task4.py::material_properties",
        },
    }
    if fixed_radius_m is None:
        radius_path = Path(radius_data_path)
        sources["radius_attachment"] = {
            "path": str(radius_path.resolve()),
            "sha256": _file_sha256(radius_path),
        }
        sources["radius_model"] = "附件2分段线性插值 R(t)"
    else:
        sources["radius_attachment"] = None
        sources["radius_model"] = {
            "type": "fixed",
            "radius_m": float(fixed_radius_m),
            "source": "机制对照设定 R=2 cm",
        }
    return sources


def _required_columns(frame: pd.DataFrame, required: tuple[str, ...]) -> dict[str, Any]:
    mapping = {str(column).strip(): column for column in frame.columns}
    missing = [name for name in required if name not in mapping]
    if missing:
        raise ValueError(f"附件缺少必需列 {missing!r}；实际列为 {list(frame.columns)!r}")
    return mapping


def load_air_data(path: str | Path = DEFAULT_AIR_DATA_PATH) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """读取附件1并确认覆盖稳定段结束14400 s。"""
    frame = pd.read_excel(Path(path))
    columns = _required_columns(frame, ("时间", "温度", "水分浓度"))
    selected = frame[[columns["时间"], columns["温度"], columns["水分浓度"]]].copy()
    selected.columns = ["时间", "温度", "水分浓度"]
    for column in selected.columns:
        selected[column] = pd.to_numeric(selected[column], errors="raise")
    selected = selected.sort_values("时间").drop_duplicates("时间", keep="last")
    time_s = selected["时间"].to_numpy(dtype=float)
    temperature_C = selected["温度"].to_numpy(dtype=float)
    moisture_kgkg = selected["水分浓度"].to_numpy(dtype=float)
    values = np.column_stack((time_s, temperature_C, moisture_kgkg))
    if time_s.size < 2 or np.any(np.diff(time_s) <= 0.0):
        raise ValueError("附件1时间必须包含至少两个严格递增时间点")
    if time_s[0] > 1.0e-10 or time_s[-1] < AIR_DATA_END_S - 1.0e-10:
        raise ValueError("附件1未覆盖到14400 s")
    if not np.all(np.isfinite(values)):
        raise ValueError("附件1数据包含NaN或Inf")
    return time_s, temperature_C, moisture_kgkg


def stable_air_means(time_s: np.ndarray, temperature_C: np.ndarray, moisture_kgkg: np.ndarray) -> tuple[float, float, int]:
    stable = (time_s >= STABLE_START_S - 1.0e-10) & (time_s <= AIR_DATA_END_S + 1.0e-10)
    if not np.any(stable):
        raise ValueError("附件1没有9000--14400 s稳定段数据")
    return float(np.mean(temperature_C[stable])), float(np.mean(moisture_kgkg[stable])), int(np.count_nonzero(stable))


def air_at(time_value_s: float, time_s: np.ndarray, temperature_C: np.ndarray, moisture_kgkg: np.ndarray, stable_temperature_C: float, stable_moisture_kgkg: float) -> tuple[float, float]:
    if time_value_s <= AIR_DATA_END_S + 1.0e-10:
        clipped = min(max(float(time_value_s), float(time_s[0])), AIR_DATA_END_S)
        return float(np.interp(clipped, time_s, temperature_C)), float(np.interp(clipped, time_s, moisture_kgkg))
    return float(stable_temperature_C), float(stable_moisture_kgkg)


def load_radius_data(path: str | Path = DEFAULT_RADIUS_DATA_PATH) -> tuple[np.ndarray, np.ndarray]:
    frame = pd.read_excel(Path(path))
    columns = _required_columns(frame, ("时间", "半径"))
    selected = frame[[columns["时间"], columns["半径"]]].copy()
    selected.columns = ["时间", "半径"]
    for column in selected.columns:
        selected[column] = pd.to_numeric(selected[column], errors="raise")
    selected = selected.sort_values("时间").drop_duplicates("时间", keep="last")
    time_s = selected["时间"].to_numpy(dtype=float)
    radius_cm = selected["半径"].to_numpy(dtype=float)
    radius_m = radius_cm / 100.0
    values = np.column_stack((time_s, radius_cm, radius_m))
    if time_s.size < 2 or np.any(np.diff(time_s) <= 0.0) or time_s[0] > 1.0e-10 or not np.all(np.isfinite(values)):
        raise ValueError("附件2时间必须从0开始且严格递增、数据有限")
    if np.any(radius_m <= 0.0) or not np.isclose(radius_m[0], RADIUS_INITIAL_M, atol=1.0e-12, rtol=0.0):
        raise ValueError("附件2半径必须为正且初始值为2 cm")
    if np.any(np.diff(radius_m) > 1.0e-12):
        raise ValueError("附件2半径必须单调不增")
    return time_s, radius_m


def radius_at(time_s: float, radius_time_s: np.ndarray, radius_m: np.ndarray) -> float:
    if time_s < radius_time_s[0] - 1.0e-10 or time_s > radius_time_s[-1] + 1.0e-10:
        raise ValueError(f"求解时刻{time_s:g}s超出附件2覆盖范围，禁止外推")
    clipped = min(max(float(time_s), float(radius_time_s[0])), float(radius_time_s[-1]))
    return float(np.interp(clipped, radius_time_s, radius_m))


def fixed_xi_geometry(n_intervals: int) -> tuple[np.ndarray, np.ndarray, np.ndarray, float]:
    if n_intervals < 2:
        raise ValueError("xi区间数至少为2")
    dxi = 1.0 / float(n_intervals)
    xi = np.linspace(0.0, 1.0, n_intervals + 1)
    faces = 0.5 * (xi[:-1] + xi[1:])
    measure = np.empty_like(xi)
    measure[0] = 0.5 * faces[0] ** 2
    measure[1:-1] = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)
    measure[-1] = 0.5 * (1.0 - faces[-1] ** 2)
    if not np.isclose(np.sum(measure), 0.5, atol=1.0e-14, rtol=0.0):
        raise ValueError("固定xi控制体测度错误")
    return xi, faces, measure, dxi


def material_properties(temperature_C: np.ndarray, moisture_kgkg: np.ndarray, appendix: int = 4) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """附录3/4物性；D中的温度始终转换为Kelvin。"""
    temperature_C = np.asarray(temperature_C, dtype=float)
    moisture_kgkg = np.asarray(moisture_kgkg, dtype=float)
    if temperature_C.shape != moisture_kgkg.shape or not np.all(np.isfinite(temperature_C)) or not np.all(np.isfinite(moisture_kgkg)):
        raise ValueError("物性输入形状或有限性错误")
    if appendix not in (3, 4) or np.min(moisture_kgkg) < -1.0e-10:
        raise ValueError("物性附录或水分范围错误")
    c = np.maximum(moisture_kgkg, 1.0e-12)
    temperature_K = temperature_C + 273.15
    if np.any(temperature_K <= 0.0) or not np.all(np.isfinite(temperature_K)):
        raise ValueError("物性计算遇到非正或非有限Kelvin温度")
    if appendix == 3:
        rho = 650.0 + 128.0 * c
        cp = 1450.0 + 2736.0 * c / (c + 1.0)
        conductivity = 0.21 + 0.38 * c / (c + 1.0)
        diffusivity = 2.4e-3 * np.exp(-0.45 / c) * np.exp(-3850.0 / temperature_K)
    else:
        rho = 760.0 + 90.0 * c
        cp = 1850.0 + 2150.0 * c / (c + 1.0)
        conductivity = 0.12 + 0.20 * c / (c + 1.0)
        diffusivity = 4.2e-4 * np.exp(-0.30 / c) * np.exp(-3850.0 / temperature_K)
    return rho, cp, conductivity, diffusivity


def _harmonic_mean(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    left = np.asarray(left, dtype=float)
    right = np.asarray(right, dtype=float)
    denominator = left + right
    if left.shape != right.shape or np.any(left <= 0.0) or np.any(right <= 0.0) or not np.all(np.isfinite(denominator)):
        raise ValueError("界面物性必须为同形状正有限数组")
    result = 2.0 * left * right / denominator
    if not np.all(np.isfinite(result)):
        raise ValueError("调和平均得到NaN或Inf")
    return result


def _solve_tridiagonal(lower: np.ndarray, diagonal: np.ndarray, upper: np.ndarray, rhs: np.ndarray) -> np.ndarray:
    n = diagonal.size
    if lower.size != n - 1 or upper.size != n - 1 or rhs.size != n:
        raise ValueError("三对角系统长度错误")
    banded = np.zeros((3, n), dtype=float)
    banded[0, 1:] = upper
    banded[1] = diagonal
    banded[2, :-1] = lower
    result = solve_banded((1, 1), banded, rhs, check_finite=False)
    if not np.all(np.isfinite(result)):
        raise RuntimeError("三对角系统求解得到NaN或Inf")
    return result


def _normalized_tridiagonal_residual(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    rhs: np.ndarray,
    solution: np.ndarray,
) -> float:
    """Compute a scale-normalized infinity residual for a tridiagonal system."""
    residual = diagonal * solution - rhs
    if solution.size > 1:
        residual[:-1] += upper * solution[1:]
        residual[1:] += lower * solution[:-1]
    scale = max(
        float(np.max(np.abs(rhs))),
        float(np.max(np.abs(diagonal * solution))),
        float(np.max(np.abs(upper * solution[1:]))) if solution.size > 1 else 0.0,
        float(np.max(np.abs(lower * solution[:-1]))) if solution.size > 1 else 0.0,
        np.finfo(float).tiny,
    )
    return float(np.max(np.abs(residual)) / scale)


def _solve_cn_field(
    old_field: np.ndarray,
    storage: np.ndarray,
    conductance: np.ndarray,
    dt_s: float,
    boundary_conductance: float,
    air_value: float,
) -> tuple[np.ndarray, float]:
    """(S+dt*K/2)new=(S-dt*K/2)old+dt*b_mid。"""
    old_field = np.asarray(old_field, dtype=float)
    storage = np.asarray(storage, dtype=float)
    conductance = np.asarray(conductance, dtype=float)
    if storage.shape != old_field.shape or conductance.size != old_field.size - 1:
        raise ValueError("CN场/存储/导通数组形状错误")
    if (
        dt_s <= 0.0
        or boundary_conductance < 0.0
        or not np.isfinite(air_value)
        or np.any(storage <= 0.0)
        or np.any(conductance <= 0.0)
    ):
        raise ValueError("CN参数无效")
    diagonal_K = np.empty_like(storage)
    diagonal_K[0] = conductance[0]
    diagonal_K[1:-1] = conductance[:-1] + conductance[1:]
    diagonal_K[-1] = conductance[-1] + boundary_conductance
    half = 0.5 * dt_s
    diagonal = storage + half * diagonal_K
    lower = -half * conductance
    upper = -half * conductance
    rhs = (storage - half * diagonal_K) * old_field
    rhs[:-1] += half * conductance * old_field[1:]
    rhs[1:] += half * conductance * old_field[:-1]
    rhs[-1] += dt_s * boundary_conductance * float(air_value)
    result = _solve_tridiagonal(lower, diagonal, upper, rhs)
    normalized_residual = _normalized_tridiagonal_residual(
        lower, diagonal, upper, rhs, result
    )
    if (
        not np.isfinite(normalized_residual)
        or normalized_residual > LINEAR_RESIDUAL_TOL
    ):
        raise RuntimeError(
            "CN三对角系统归一化残差超限: "
            f"{normalized_residual:.3e}>{LINEAR_RESIDUAL_TOL:.3e}"
        )
    return result, normalized_residual


def _picard_step(old_temperature_C: np.ndarray, old_moisture_kgkg: np.ndarray, xi: np.ndarray, faces: np.ndarray, measure: np.ndarray, dxi: float, radius_mid_m: float, dt_s: float, air_temperature_C: float, air_moisture_kgkg: float, appendix: int = 4, context_time_s: float | None = None) -> tuple[np.ndarray, np.ndarray, int, float, float, float, float]:
    """同一组中点物性同步求解温度和水分，固定材料坐标无Rdot迎风项。"""
    if dt_s <= 0.0 or radius_mid_m <= 0.0:
        raise ValueError("Picard时间步和中点半径必须为正")
    if old_temperature_C.shape != xi.shape or old_moisture_kgkg.shape != xi.shape or measure.shape != xi.shape or faces.shape != (xi.size - 1,):
        raise ValueError("Picard网格形状错误")
    guess_T = old_temperature_C.copy()
    guess_C = old_moisture_kgkg.copy()
    dT = np.inf
    dC = np.inf
    for iteration in range(1, PICARD_MAX_ITER + 1):
        midpoint_T = 0.5 * (old_temperature_C + guess_T)
        midpoint_C = 0.5 * (old_moisture_kgkg + guess_C)
        rho, cp, conductivity, diffusivity = material_properties(midpoint_T, midpoint_C, appendix)
        face_k = _harmonic_mean(conductivity[:-1], conductivity[1:])
        face_D = _harmonic_mean(diffusivity[:-1], diffusivity[1:])
        scale = radius_mid_m * radius_mid_m * dxi
        new_T, temperature_linear_residual = _solve_cn_field(
            old_temperature_C,
            rho * cp * measure,
            faces * face_k / scale,
            dt_s,
            HEAT_TRANSFER_W_M2_K / radius_mid_m,
            air_temperature_C,
        )
        new_C, moisture_linear_residual = _solve_cn_field(
            old_moisture_kgkg,
            measure,
            faces * face_D / scale,
            dt_s,
            MASS_TRANSFER_M_S / radius_mid_m,
            air_moisture_kgkg,
        )
        if not np.all(np.isfinite(new_T)) or not np.all(np.isfinite(new_C)):
            raise RuntimeError("Picard步得到NaN或Inf")
        if np.min(new_C) < -1.0e-10:
            raise RuntimeError("Picard步得到明显负水分浓度")
        new_C = np.maximum(new_C, 0.0)
        dT = float(np.max(np.abs(new_T - guess_T)))
        dC = float(np.max(np.abs(new_C - guess_C)))
        if dT <= PICARD_TOL_T_C and dC <= PICARD_TOL_C_KGKG:
            return (
                new_T,
                new_C,
                iteration,
                dT,
                dC,
                float(temperature_linear_residual),
                float(moisture_linear_residual),
            )
        guess_T, guess_C = new_T, new_C
    location = "" if context_time_s is None else f", t={context_time_s:g}s"
    raise RuntimeError(f"CN Picard迭代未收敛{location}, dt={dt_s:g}s, R_mid={radius_mid_m:g}m, iter={PICARD_MAX_ITER}, max|dT|={dT:.3e}, max|dC|={dC:.3e}")


def _balance_residuals(
    old_temperature_C: np.ndarray,
    old_moisture_kgkg: np.ndarray,
    new_temperature_C: np.ndarray,
    new_moisture_kgkg: np.ndarray,
    measure: np.ndarray,
    radius_mid_m: float,
    dt_s: float,
    air_temperature_C: float,
    air_moisture_kgkg: float,
    appendix: int,
) -> tuple[float, float]:
    """Return global CN balance residuals in the fixed material coordinate.

    The spatial fluxes telescope under the radial finite-volume sum.  Thus the
    storage-rate minus midpoint Robin flux should be roundoff-sized for each
    accepted step.  The residual is intentionally written in the same
    material-coordinate measure as the handoff; no extra ``Rdot`` source is
    introduced.
    """
    midpoint_temperature = 0.5 * (old_temperature_C + new_temperature_C)
    midpoint_moisture = 0.5 * (old_moisture_kgkg + new_moisture_kgkg)
    rho, cp, _, _ = material_properties(
        midpoint_temperature, midpoint_moisture, appendix
    )
    midpoint_surface_temperature = float(
        0.5 * (old_temperature_C[-1] + new_temperature_C[-1])
    )
    midpoint_surface_moisture = float(
        0.5 * (old_moisture_kgkg[-1] + new_moisture_kgkg[-1])
    )
    temperature_boundary_flux = HEAT_TRANSFER_W_M2_K / radius_mid_m * (
        air_temperature_C - midpoint_surface_temperature
    )
    moisture_boundary_flux = MASS_TRANSFER_M_S / radius_mid_m * (
        air_moisture_kgkg - midpoint_surface_moisture
    )
    temperature_storage_rate = float(
        (rho * cp * measure) @ (new_temperature_C - old_temperature_C)
    ) / dt_s
    moisture_storage_rate = float(
        measure @ (new_moisture_kgkg - old_moisture_kgkg)
    ) / dt_s
    return (
        temperature_storage_rate - temperature_boundary_flux,
        moisture_storage_rate - moisture_boundary_flux,
    )


def _radius_for_time(time_s: float, radius_time_s: np.ndarray | None, radius_values_m: np.ndarray | None, fixed_radius_m: float | None) -> float:
    if fixed_radius_m is not None:
        return float(fixed_radius_m)
    if radius_time_s is None or radius_values_m is None:
        raise ValueError("动态半径求解缺少附件2数据")
    return radius_at(time_s, radius_time_s, radius_values_m)


def _advance_step(old_time_s: float, dt_s: float, old_temperature_C: np.ndarray, old_moisture_kgkg: np.ndarray, xi: np.ndarray, faces: np.ndarray, measure: np.ndarray, dxi: float, radius_time_s: np.ndarray | None, radius_values_m: np.ndarray | None, fixed_radius_m: float | None, air_time_s: np.ndarray, air_temperature_C: np.ndarray, air_moisture_kgkg: np.ndarray, stable_temperature_C: float, stable_moisture_kgkg: float, appendix: int = 4) -> tuple[Any, ...]:
    midpoint_time = float(old_time_s + 0.5 * dt_s)
    radius_mid = _radius_for_time(midpoint_time, radius_time_s, radius_values_m, fixed_radius_m)
    midpoint_air = air_at(midpoint_time, air_time_s, air_temperature_C, air_moisture_kgkg, stable_temperature_C, stable_moisture_kgkg)
    result = _picard_step(
        old_temperature_C,
        old_moisture_kgkg,
        xi,
        faces,
        measure,
        dxi,
        radius_mid,
        dt_s,
        midpoint_air[0],
        midpoint_air[1],
        appendix,
        midpoint_time,
    )
    new_temperature_C, new_moisture_kgkg = result[0], result[1]
    temperature_balance, moisture_balance = _balance_residuals(
        old_temperature_C,
        old_moisture_kgkg,
        new_temperature_C,
        new_moisture_kgkg,
        measure,
        radius_mid,
        dt_s,
        midpoint_air[0],
        midpoint_air[1],
        appendix,
    )
    return (*result, radius_mid, midpoint_air, temperature_balance, moisture_balance)


def _max_info(moisture_kgkg: np.ndarray, xi: np.ndarray, radius_m: float) -> tuple[float, int, float, float]:
    index = int(np.argmax(moisture_kgkg))
    return float(moisture_kgkg[index]), index, float(xi[index]), float(xi[index] * radius_m)


def _event_bisection(lower_time_s: float, lower_temperature_C: np.ndarray, lower_moisture_kgkg: np.ndarray, lower_max_kgkg: float, upper_time_s: float, upper_temperature_C: np.ndarray, upper_moisture_kgkg: np.ndarray, upper_max_kgkg: float, xi: np.ndarray, faces: np.ndarray, measure: np.ndarray, dxi: float, radius_time_s: np.ndarray | None, radius_values_m: np.ndarray | None, air_time_s: np.ndarray, air_temperature_C: np.ndarray, air_moisture_kgkg: np.ndarray, stable_temperature_C: float, stable_moisture_kgkg: float, fixed_radius_m: float | None = None, appendix: int = 4, event_tol_s: float = EVENT_TOL_S, temperature_bounds: tuple[float, float] | None = None) -> dict[str, Any]:
    """候选时刻均从同一个跨阈值前base状态重新推进。"""
    if lower_max_kgkg < DRYING_THRESHOLD_KGKG or upper_max_kgkg >= DRYING_THRESHOLD_KGKG or upper_time_s <= lower_time_s or event_tol_s <= 0.0:
        raise ValueError("事件二分上下端或精度错误")
    base_time = float(lower_time_s)
    base_T = lower_temperature_C.copy()
    base_C = lower_moisture_kgkg.copy()
    lower_time, lower_T, lower_C, lower_max = base_time, base_T.copy(), base_C.copy(), float(lower_max_kgkg)
    upper_time, upper_T, upper_C, upper_max = float(upper_time_s), upper_temperature_C.copy(), upper_moisture_kgkg.copy(), float(upper_max_kgkg)
    event_iterations = 0
    event_picard_iterations: list[int] = []
    event_temperature_balance: list[float] = []
    event_moisture_balance: list[float] = []
    while upper_time - lower_time > event_tol_s:
        midpoint = 0.5 * (lower_time + upper_time)
        values = _advance_step(base_time, midpoint - base_time, base_T, base_C, xi, faces, measure, dxi, radius_time_s, radius_values_m, fixed_radius_m, air_time_s, air_temperature_C, air_moisture_kgkg, stable_temperature_C, stable_moisture_kgkg, appendix)
        midpoint_T, midpoint_C = values[0], values[1]
        midpoint_picard_iterations = int(values[2])
        if temperature_bounds is not None:
            _check_state(midpoint_T, midpoint_C, temperature_bounds, f"事件候选t={midpoint:g}s")
        event_picard_iterations.append(midpoint_picard_iterations)
        event_temperature_balance.append(float(values[9]))
        event_moisture_balance.append(float(values[10]))
        midpoint_radius = _radius_for_time(midpoint, radius_time_s, radius_values_m, fixed_radius_m)
        midpoint_max, _, _, _ = _max_info(midpoint_C, xi, midpoint_radius)
        if midpoint_max < DRYING_THRESHOLD_KGKG:
            upper_time, upper_T, upper_C, upper_max = midpoint, midpoint_T, midpoint_C, midpoint_max
        else:
            lower_time, lower_T, lower_C, lower_max = midpoint, midpoint_T, midpoint_C, midpoint_max
        event_iterations += 1
    upper_radius = _radius_for_time(upper_time, radius_time_s, radius_values_m, fixed_radius_m)
    lower_radius = _radius_for_time(lower_time, radius_time_s, radius_values_m, fixed_radius_m)
    upper_max, upper_index, upper_xi, upper_position = _max_info(upper_C, xi, upper_radius)
    lower_max, lower_index, lower_xi, lower_position = _max_info(lower_C, xi, lower_radius)
    return {
        "termination_time_s": upper_time,
        "termination_temperature_internal_C": upper_T,
        "termination_moisture_internal_kgkg": upper_C,
        "termination_max_kgkg": upper_max,
        "termination_max_index": upper_index,
        "termination_max_xi": upper_xi,
        "termination_max_radius_m": upper_position,
        "termination_radius_m": upper_radius,
        "termination_lower_time_s": lower_time,
        "termination_lower_temperature_internal_C": lower_T,
        "termination_lower_moisture_internal_kgkg": lower_C,
        "termination_lower_max_kgkg": lower_max,
        "termination_lower_max_index": lower_index,
        "termination_lower_max_xi": lower_xi,
        "termination_lower_max_radius_m": lower_position,
        "termination_lower_radius_m": lower_radius,
        "termination_bracket_width_s": upper_time - lower_time,
        "termination_event_iterations": event_iterations,
        "termination_event_picard_max": int(max(event_picard_iterations, default=0)),
        "termination_event_picard_total": int(sum(event_picard_iterations)),
        "termination_event_candidate_count": int(len(event_picard_iterations)),
        "termination_event_temperature_balance_max_abs": float(
            max((abs(value) for value in event_temperature_balance), default=0.0)
        ),
        "termination_event_moisture_balance_max_abs": float(
            max((abs(value) for value in event_moisture_balance), default=0.0)
        ),
    }


def _state_bounds(air_temperature_C: np.ndarray, stable_temperature_C: float) -> tuple[float, float]:
    return float(min(INITIAL_TEMPERATURE_C, np.min(air_temperature_C), stable_temperature_C)), float(max(INITIAL_TEMPERATURE_C, np.max(air_temperature_C), stable_temperature_C))


def _check_state(temperature_C: np.ndarray, moisture_kgkg: np.ndarray, temperature_bounds: tuple[float, float], context: str) -> None:
    if not np.all(np.isfinite(temperature_C)) or not np.all(np.isfinite(moisture_kgkg)):
        raise RuntimeError(f"{context}包含NaN或Inf")
    if np.min(moisture_kgkg) < -1.0e-10 or np.max(moisture_kgkg) > INITIAL_MOISTURE_KGKG + STATE_TOL:
        raise RuntimeError(f"{context}水分超出[0,C0]范围")
    lower, upper = temperature_bounds
    if np.min(temperature_C) < lower - STATE_TOL or np.max(temperature_C) > upper + STATE_TOL:
        raise RuntimeError(f"{context}温度超出[{lower:g},{upper:g}]C范围")


def solve_task4(air_data_path: str | Path = DEFAULT_AIR_DATA_PATH, radius_data_path: str | Path = DEFAULT_RADIUS_DATA_PATH, max_time_s: float | None = None, time_step_s: float | None = None, output_dt_s: float = OUTPUT_DT_S, n_intervals: int = N_DEFAULT, progress: bool = False, *, time_step_early_s: float | None = None, time_step_long_s: float | None = None, event_tol_s: float = EVENT_TOL_S, appendix: int = 4, fixed_radius_m: float | None = None) -> dict[str, Any]:
    """从0开始求解，动态半径只在附件2覆盖范围内使用。"""
    if time_step_s is not None:
        if time_step_early_s is not None or time_step_long_s is not None:
            raise ValueError("time_step不能与early/long同时使用")
        time_step_early_s = time_step_long_s = float(time_step_s)
    time_step_early_s = TIME_STEP_EARLY_S if time_step_early_s is None else float(time_step_early_s)
    time_step_long_s = TIME_STEP_LONG_DEFAULT_S if time_step_long_s is None else float(time_step_long_s)
    max_time_s = (FIXED_MAX_TIME_S if fixed_radius_m is not None else DEFAULT_MAX_TIME_S) if max_time_s is None else float(max_time_s)
    if max_time_s <= 0.0 or output_dt_s <= 0.0 or time_step_early_s <= 0.0 or time_step_long_s <= 0.0 or event_tol_s <= 0.0:
        raise ValueError("时间、时间步和事件精度必须为正")
    if fixed_radius_m is not None and fixed_radius_m <= 0.0:
        raise ValueError("固定半径必须为正")
    for step_value in (time_step_early_s, time_step_long_s):
        ratio = output_dt_s / step_value
        if not np.isclose(ratio, round(ratio), atol=1.0e-10, rtol=0.0):
            raise ValueError("60 s输出间隔必须是两段内部时间步的整数倍")
    air_time_s, air_temperature, air_moisture = load_air_data(air_data_path)
    stable_T, stable_C, stable_count = stable_air_means(air_time_s, air_temperature, air_moisture)
    if fixed_radius_m is None:
        radius_time_s, radius_values_m = load_radius_data(radius_data_path)
        if max_time_s > radius_time_s[-1] + 1.0e-10:
            raise ValueError("动态半径求解超过附件2覆盖范围")
    else:
        radius_time_s = np.array([], dtype=float)
        radius_values_m = np.array([], dtype=float)
    xi, faces, measure, dxi = fixed_xi_geometry(n_intervals)
    temperature = np.full(xi.size, INITIAL_TEMPERATURE_C, dtype=float)
    moisture = np.full(xi.size, INITIAL_MOISTURE_KGKG, dtype=float)
    bounds = _state_bounds(air_temperature, stable_T)
    current_time = 0.0
    current_radius = _radius_for_time(0.0, None if fixed_radius_m is not None else radius_time_s, None if fixed_radius_m is not None else radius_values_m, fixed_radius_m)
    current_max, _, _, _ = _max_info(moisture, xi, current_radius)
    sampled_times: list[float] = [0.0]
    sampled_radius: list[float] = [current_radius]
    sampled_T: list[np.ndarray] = [temperature.copy()]
    sampled_C: list[np.ndarray] = [moisture.copy()]
    sampled_max: list[float] = [current_max]
    step_times: list[float] = []
    step_max: list[float] = []
    step_xi: list[float] = []
    step_position: list[float] = []
    step_radius: list[float] = []
    step_midpoint_radius: list[float] = []
    step_dt: list[float] = []
    step_midpoint_air_temperature: list[float] = []
    step_midpoint_air_moisture: list[float] = []
    radius_increments: list[float] = []
    iterations: list[int] = []
    delta_T: list[float] = []
    delta_C: list[float] = []
    temperature_linear_residuals: list[float] = []
    moisture_linear_residuals: list[float] = []
    temperature_balance_residuals: list[float] = []
    moisture_balance_residuals: list[float] = []
    temperature_boundary_fluxes: list[float] = []
    moisture_boundary_fluxes: list[float] = []
    min_T, max_T = float(np.min(temperature)), float(np.max(temperature))
    min_C, max_C = float(np.min(moisture)), float(np.max(moisture))
    early_steps = long_steps = 0
    progress_next = 0.0
    termination: dict[str, Any] | None = None
    guard = 0
    guard_limit = int(np.ceil(max_time_s / min(time_step_early_s, time_step_long_s))) + 1000

    def next_boundary(now: float, values: list[float]) -> float:
        valid = [value for value in values if value > now + 1.0e-9]
        return min(valid) if valid else float("inf")

    while current_time < max_time_s - 1.0e-9:
        guard += 1
        if guard > guard_limit:
            raise RuntimeError("超过时间步保护上限仍未达到烘干阈值")
        phase_dt = time_step_early_s if current_time < AIR_DATA_END_S - 1.0e-9 else time_step_long_s
        candidates = [current_time + phase_dt, max_time_s, (np.floor(current_time / output_dt_s + 1.0e-10) + 1.0) * output_dt_s, next_boundary(current_time, [AIR_DATA_END_S])]
        if fixed_radius_m is None:
            candidates.append(next_boundary(current_time, list(radius_time_s[1:])))
        time_new = min(value for value in candidates if value > current_time + 1.0e-9)
        time_new = min(time_new, max_time_s)
        dt_current = time_new - current_time
        values = _advance_step(current_time, dt_current, temperature, moisture, xi, faces, measure, dxi, None if fixed_radius_m is not None else radius_time_s, None if fixed_radius_m is not None else radius_values_m, fixed_radius_m, air_time_s, air_temperature, air_moisture, stable_T, stable_C, appendix)
        new_T, new_C, n_iter, dT, dC = values[:5]
        new_radius = _radius_for_time(time_new, None if fixed_radius_m is not None else radius_time_s, None if fixed_radius_m is not None else radius_values_m, fixed_radius_m)
        midpoint_radius = float(values[7])
        midpoint_air_temperature = float(values[8][0])
        midpoint_air_moisture = float(values[8][1])
        _check_state(new_T, new_C, bounds, f"t={time_new:g}s步")
        new_max, new_index, new_xi, new_position = _max_info(new_C, xi, new_radius)

        # 先把本次完整CN步纳入范围、Picard和半径诊断；即使它跨过阈值，
        # 也不能因为随后进入事件二分而遗漏这个全步状态。
        step_times.append(time_new)
        step_max.append(new_max)
        step_xi.append(new_xi)
        step_position.append(new_position)
        step_radius.append(new_radius)
        step_midpoint_radius.append(midpoint_radius)
        step_dt.append(float(dt_current))
        step_midpoint_air_temperature.append(midpoint_air_temperature)
        step_midpoint_air_moisture.append(midpoint_air_moisture)
        radius_increments.append(new_radius - current_radius)
        iterations.append(int(n_iter))
        delta_T.append(float(dT))
        delta_C.append(float(dC))
        temperature_linear_residuals.append(float(values[5]))
        moisture_linear_residuals.append(float(values[6]))
        temperature_balance_residuals.append(float(values[9]))
        moisture_balance_residuals.append(float(values[10]))
        temperature_boundary_fluxes.append(
            HEAT_TRANSFER_W_M2_K
            / midpoint_radius
            * (midpoint_air_temperature - 0.5 * (temperature[-1] + new_T[-1]))
        )
        moisture_boundary_fluxes.append(
            MASS_TRANSFER_M_S
            / midpoint_radius
            * (midpoint_air_moisture - 0.5 * (moisture[-1] + new_C[-1]))
        )
        min_T, max_T = min(min_T, float(np.min(new_T))), max(max_T, float(np.max(new_T)))
        min_C, max_C = min(min_C, float(np.min(new_C))), max(max_C, float(np.max(new_C)))
        if current_time < AIR_DATA_END_S - 1.0e-9:
            early_steps += 1
        else:
            long_steps += 1

        if current_max >= DRYING_THRESHOLD_KGKG and new_max < DRYING_THRESHOLD_KGKG:
            termination = _event_bisection(current_time, temperature, moisture, current_max, time_new, new_T, new_C, new_max, xi, faces, measure, dxi, None if fixed_radius_m is not None else radius_time_s, None if fixed_radius_m is not None else radius_values_m, air_time_s, air_temperature, air_moisture, stable_T, stable_C, fixed_radius_m, appendix, event_tol_s, bounds)
            _check_state(termination["termination_temperature_internal_C"], termination["termination_moisture_internal_kgkg"], bounds, "事件二分上端")
            _check_state(termination["termination_lower_temperature_internal_C"], termination["termination_lower_moisture_internal_kgkg"], bounds, "事件二分下端")
            event_temperatures = (
                termination["termination_temperature_internal_C"],
                termination["termination_lower_temperature_internal_C"],
            )
            event_moistures = (
                termination["termination_moisture_internal_kgkg"],
                termination["termination_lower_moisture_internal_kgkg"],
            )
            termination["event_min_temperature_C"] = float(min(np.min(state) for state in event_temperatures))
            termination["event_max_temperature_C"] = float(max(np.max(state) for state in event_temperatures))
            termination["event_min_moisture_kgkg"] = float(min(np.min(state) for state in event_moistures))
            termination["event_max_moisture_kgkg"] = float(max(np.max(state) for state in event_moistures))
            min_T = min(min_T, termination["event_min_temperature_C"])
            max_T = max(max_T, termination["event_max_temperature_C"])
            min_C = min(min_C, termination["event_min_moisture_kgkg"])
            max_C = max(max_C, termination["event_max_moisture_kgkg"])
            break
        if np.isclose(time_new / output_dt_s, round(time_new / output_dt_s), atol=1.0e-9, rtol=0.0):
            sampled_times.append(time_new)
            sampled_radius.append(new_radius)
            sampled_T.append(new_T.copy())
            sampled_C.append(new_C.copy())
            sampled_max.append(new_max)
        if progress and time_new >= progress_next - 1.0e-9:
            print(f"  已推进{time_new / 3600.0:.2f} h，R={new_radius * 100.0:.4f} cm，max C={new_max:.10f}，Picard={n_iter}")
            while progress_next <= time_new + 1.0e-9:
                progress_next += 6.0 * 3600.0
        current_time, current_radius = time_new, new_radius
        temperature, moisture, current_max = new_T, new_C, new_max
    if termination is None:
        raise RuntimeError(f"在{max_time_s / 3600.0:g} h内未达到全场严格阈值max C < {DRYING_THRESHOLD_KGKG:g}")
    return {
        "xi": xi, "dxi": dxi, "n_intervals": int(n_intervals),
        "time_step_s": time_step_long_s, "time_step_early_s": time_step_early_s,
        "time_step_long_s": time_step_long_s, "output_dt_s": output_dt_s,
        "event_tol_s": event_tol_s, "appendix": appendix,
        "fixed_radius_m": fixed_radius_m, "max_time_s": max_time_s,
        "temperature_bounds_C": bounds,
        "stable_temperature_C": stable_T, "stable_moisture_kgkg": stable_C, "stable_count": stable_count,
        "radius_data_time_s": radius_time_s.copy(), "radius_data_m": radius_values_m.copy(),
        "air_data_time_s": air_time_s.copy(), "air_temperature_C": air_temperature.copy(), "air_moisture_kgkg": air_moisture.copy(),
        "sampled_times_s": np.asarray(sampled_times), "sampled_radius_m": np.asarray(sampled_radius),
        "sampled_temperature_internal_C": np.asarray(sampled_T), "sampled_moisture_internal_kgkg": np.asarray(sampled_C), "sampled_max_kgkg": np.asarray(sampled_max),
        "step_times_s": np.asarray(step_times), "step_max_kgkg": np.asarray(step_max), "step_max_xi": np.asarray(step_xi), "step_max_radius_m": np.asarray(step_position), "step_radius_m": np.asarray(step_radius), "radius_increments_m": np.asarray(radius_increments),
        "step_midpoint_radius_m": np.asarray(step_midpoint_radius),
        "step_dt_s": np.asarray(step_dt),
        "step_midpoint_air_temperature_C": np.asarray(step_midpoint_air_temperature),
        "step_midpoint_air_moisture_kgkg": np.asarray(step_midpoint_air_moisture),
        "iteration_counts": np.asarray(iterations, dtype=int), "picard_delta_temperature": np.asarray(delta_T), "picard_delta_moisture": np.asarray(delta_C),
        "temperature_linear_residuals": np.asarray(temperature_linear_residuals),
        "moisture_linear_residuals": np.asarray(moisture_linear_residuals),
        "linear_residuals": np.maximum(
            np.asarray(temperature_linear_residuals),
            np.asarray(moisture_linear_residuals),
        ),
        "temperature_balance_residuals": np.asarray(temperature_balance_residuals),
        "moisture_balance_residuals": np.asarray(moisture_balance_residuals),
        "temperature_boundary_fluxes": np.asarray(temperature_boundary_fluxes),
        "moisture_boundary_fluxes": np.asarray(moisture_boundary_fluxes),
        "early_steps": early_steps, "long_steps": long_steps,
        "min_temperature_C": min_T, "max_temperature_C": max_T, "min_internal_moisture_kgkg": min_C, "max_internal_moisture_kgkg": max_C,
        "input_sources": _input_sources(
            air_data_path, radius_data_path, appendix, fixed_radius_m
        ),
        **termination,
    }


def validate_solution(solution: dict[str, Any]) -> dict[str, Any]:
    xi = np.asarray(solution["xi"], dtype=float)
    times = np.asarray(solution["sampled_times_s"], dtype=float)
    radii = np.asarray(solution["sampled_radius_m"], dtype=float)
    sampled_T = np.asarray(solution["sampled_temperature_internal_C"], dtype=float)
    sampled_C = np.asarray(solution["sampled_moisture_internal_kgkg"], dtype=float)
    lower_T = np.asarray(solution["termination_lower_temperature_internal_C"], dtype=float)
    lower_C = np.asarray(solution["termination_lower_moisture_internal_kgkg"], dtype=float)
    upper_T = np.asarray(solution["termination_temperature_internal_C"], dtype=float)
    upper_C = np.asarray(solution["termination_moisture_internal_kgkg"], dtype=float)
    if not all(np.all(np.isfinite(a)) for a in (xi, times, radii, sampled_T, sampled_C, lower_T, lower_C, upper_T, upper_C)):
        raise ValueError("结果包含NaN或Inf")
    if xi[0] != 0.0 or xi[-1] != 1.0 or np.any(np.diff(xi) <= 0.0):
        raise ValueError("固定xi网格错误")
    if times.size < 1 or not np.isclose(times[0], 0.0, atol=1.0e-12):
        raise ValueError("采样未从0开始")
    if times.size > 1 and not np.allclose(np.diff(times), OUTPUT_DT_S, atol=1.0e-8, rtol=0.0):
        raise ValueError("采样不是60 s间隔")
    if sampled_T.shape != (times.size, xi.size) or sampled_C.shape != sampled_T.shape or radii.shape != times.shape:
        raise ValueError("采样状态形状错误")
    fixed_radius = solution.get("fixed_radius_m")
    radius_time = np.asarray(solution["radius_data_time_s"], dtype=float)
    radius_values = np.asarray(solution["radius_data_m"], dtype=float)
    if fixed_radius is None:
        if radius_time.size < 2 or not np.isclose(radius_values[0], RADIUS_INITIAL_M, atol=1.0e-12, rtol=0.0) or np.any(np.diff(radius_values) > 1.0e-12):
            raise ValueError("附件2半径数据错误")
        expected_radii = np.array([radius_at(t, radius_time, radius_values) for t in times])
        known_error = float(np.max(np.abs(np.array([radius_at(t, radius_time, radius_values) for t in radius_time]) - radius_values)))
    else:
        expected_radii = np.full(times.shape, float(fixed_radius))
        known_error = 0.0
    if not np.allclose(radii, expected_radii, atol=1.0e-12, rtol=0.0):
        raise ValueError("R(t)采样错误")
    bounds = tuple(solution["temperature_bounds_C"])
    _check_state(lower_T, lower_C, bounds, "终止下端")
    _check_state(upper_T, upper_C, bounds, "终止上端")
    if np.min(sampled_C) < -1.0e-10 or np.max(sampled_C) > INITIAL_MOISTURE_KGKG + STATE_TOL:
        raise ValueError("采样水分范围错误")
    lower_time, upper_time = float(solution["termination_lower_time_s"]), float(solution["termination_time_s"])
    lower_max, upper_max = float(solution["termination_lower_max_kgkg"]), float(solution["termination_max_kgkg"])
    if lower_max < DRYING_THRESHOLD_KGKG or upper_max >= DRYING_THRESHOLD_KGKG or not (lower_time < upper_time <= lower_time + float(solution["event_tol_s"]) + 1.0e-10):
        raise ValueError("终止阈值括号错误")
    if not np.isclose(lower_max, np.max(lower_C), atol=1.0e-12, rtol=0.0) or not np.isclose(upper_max, np.max(upper_C), atol=1.0e-12, rtol=0.0):
        raise ValueError("终止max C不是全场最大值")
    upper_index = int(np.argmax(upper_C))
    if upper_index != int(solution["termination_max_index"]):
        raise ValueError("终止最大值索引错误")
    if not np.isclose(float(solution["termination_max_radius_m"]), xi[upper_index] * float(solution["termination_radius_m"]), atol=1.0e-14, rtol=0.0):
        raise ValueError("终止最大值物理位置错误")
    if fixed_radius is None and np.any(np.asarray(solution["radius_increments_m"], dtype=float) > 1.0e-12):
        raise ValueError("附件2半径出现非单调增大")
    if known_error > 1.0e-14:
        raise ValueError("附件2已知点回读错误")
    D_ref = material_properties(np.array([50.0]), np.array([0.10]), 4)[3][0]
    if not np.isclose(D_ref, 4.2e-4 * np.exp(-0.30 / 0.10) * np.exp(-3850.0 / 323.15), rtol=1.0e-13, atol=0.0):
        raise ValueError("Kelvin温度验证失败")
    iterations = np.asarray(solution["iteration_counts"], dtype=int)
    event_picard_max = int(solution.get("termination_event_picard_max", 0))
    event_candidate_count = int(solution.get("termination_event_candidate_count", 0))
    if iterations.size == 0 or np.max(iterations) > PICARD_MAX_ITER or event_picard_max > PICARD_MAX_ITER:
        raise ValueError("Picard记录错误")
    for key in (
        "step_times_s",
        "step_max_kgkg",
        "step_max_xi",
        "step_max_radius_m",
        "step_radius_m",
        "step_midpoint_radius_m",
        "step_dt_s",
        "step_midpoint_air_temperature_C",
        "step_midpoint_air_moisture_kgkg",
        "radius_increments_m",
        "picard_delta_temperature",
        "picard_delta_moisture",
        "temperature_linear_residuals",
        "moisture_linear_residuals",
        "linear_residuals",
        "temperature_balance_residuals",
        "moisture_balance_residuals",
        "temperature_boundary_fluxes",
        "moisture_boundary_fluxes",
    ):
        if np.asarray(solution[key]).size != iterations.size:
            raise ValueError(f"诊断字段{key}长度错误")
        if not np.all(np.isfinite(np.asarray(solution[key], dtype=float))):
            raise ValueError(f"诊断字段{key}包含NaN或Inf")
    if np.any(np.asarray(solution["picard_delta_temperature"]) > PICARD_TOL_T_C) or np.any(
        np.asarray(solution["picard_delta_moisture"]) > PICARD_TOL_C_KGKG
    ):
        raise ValueError("存在未达到收敛阈值的Picard时间步")
    temperature_linear = np.asarray(solution["temperature_linear_residuals"], dtype=float)
    moisture_linear = np.asarray(solution["moisture_linear_residuals"], dtype=float)
    combined_linear = np.asarray(solution["linear_residuals"], dtype=float)
    if not np.allclose(combined_linear, np.maximum(temperature_linear, moisture_linear), atol=0.0, rtol=0.0):
        raise ValueError("综合线性残差记录错误")
    if np.max(combined_linear) > LINEAR_RESIDUAL_TOL:
        raise ValueError("三对角线性系统残差超过容差")
    step_times = np.asarray(solution["step_times_s"], dtype=float)
    step_dt = np.asarray(solution["step_dt_s"], dtype=float)
    if not np.allclose(step_times, np.cumsum(step_dt), atol=1.0e-8, rtol=0.0):
        raise ValueError("时间步诊断记录与步长不一致")
    if np.any(np.asarray(solution["step_midpoint_radius_m"], dtype=float) <= 0.0):
        raise ValueError("中点半径诊断记录无效")
    return {
        "termination_time_s": upper_time, "termination_time_h": upper_time / 3600.0,
        "termination_radius_cm": float(solution["termination_radius_m"]) * 100.0,
        "termination_max_kgkg": upper_max, "termination_max_xi": float(solution["termination_max_xi"]),
        "termination_max_radius_cm": float(solution["termination_max_radius_m"]) * 100.0,
        "termination_lower_time_s": lower_time, "termination_lower_max_kgkg": lower_max,
        "termination_bracket_width_s": upper_time - lower_time,
        "radius_known_point_max_error_m": known_error, "radius_known_point_count": int(radius_values.size),
        "temperature_min_C": float(solution["min_temperature_C"]), "temperature_max_C": float(solution["max_temperature_C"]),
        "moisture_min_kgkg": float(solution["min_internal_moisture_kgkg"]), "moisture_max_kgkg": float(solution["max_internal_moisture_kgkg"]),
        "event_min_temperature_C": float(solution["event_min_temperature_C"]), "event_max_temperature_C": float(solution["event_max_temperature_C"]),
        "event_min_moisture_kgkg": float(solution["event_min_moisture_kgkg"]), "event_max_moisture_kgkg": float(solution["event_max_moisture_kgkg"]),
        "max_picard_iterations": int(max(np.max(iterations), event_picard_max)),
        "mean_picard_iterations": float(np.mean(iterations)),
        "event_picard_max": event_picard_max, "event_picard_total": int(solution.get("termination_event_picard_total", 0)),
        "event_candidate_count": event_candidate_count,
        "picard_iteration_count_total": int(iterations.size + event_candidate_count),
        "sample_count": int(times.size), "internal_node_count": int(xi.size), "center_boundary_zero_flux": True,
        "stable_temperature_C": float(solution["stable_temperature_C"]), "stable_moisture_kgkg": float(solution["stable_moisture_kgkg"]), "stable_count": int(solution["stable_count"]),
        "early_steps": int(solution["early_steps"]), "long_steps": int(solution["long_steps"]), "event_iterations": int(solution.get("termination_event_iterations", 0)), "kelvin_reference_D_m2_s": float(D_ref),
        "max_temperature_linear_residual": float(np.max(temperature_linear)),
        "max_moisture_linear_residual": float(np.max(moisture_linear)),
        "max_linear_residual": float(np.max(combined_linear)),
        "max_temperature_balance_residual": float(np.max(np.abs(np.asarray(solution["temperature_balance_residuals"], dtype=float)))),
        "max_moisture_balance_residual": float(np.max(np.abs(np.asarray(solution["moisture_balance_residuals"], dtype=float)))),
        "max_temperature_boundary_flux": float(np.max(np.abs(np.asarray(solution["temperature_boundary_fluxes"], dtype=float)))),
        "max_moisture_boundary_flux": float(np.max(np.abs(np.asarray(solution["moisture_boundary_fluxes"], dtype=float)))),
    }


def _regular_sample_indices(solution: dict[str, Any]) -> np.ndarray:
    times = np.asarray(solution["sampled_times_s"], dtype=float)
    last_time = np.floor(float(solution["termination_time_s"]) / OUTPUT_DT_S + 1.0e-12) * OUTPUT_DT_S
    indices = np.flatnonzero((times >= OUTPUT_DT_S - 1.0e-8) & (times <= last_time + 1.0e-8))
    if indices.size and not np.allclose(times[indices], np.arange(OUTPUT_DT_S, last_time + 0.1, OUTPUT_DT_S), atol=1.0e-8):
        raise ValueError("缺少连续60 s采样")
    return indices


def _paper_regular_hours(solution: dict[str, Any]) -> list[int]:
    last = int(np.floor(float(solution["termination_time_s"]) / 3600.0 / 6.0 + 1.0e-12) * 6)
    return list(range(6, last + 1, 6))


def _sample_index_at_time(solution: dict[str, Any], time_s: float) -> int:
    matches = np.flatnonzero(np.isclose(np.asarray(solution["sampled_times_s"]), time_s, atol=1.0e-8, rtol=0.0))
    if matches.size != 1:
        raise ValueError(f"缺少{time_s:g}s完整状态")
    return int(matches[0])


def _fixed_state_values(xi: np.ndarray, temperature_C: np.ndarray, moisture_kgkg: np.ndarray, radius_m: float) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    fixed_T = np.full(OUTPUT_RADIUS_CM.size, np.nan)
    fixed_C = np.full(OUTPUT_RADIUS_CM.size, np.nan)
    valid = OUTPUT_RADIUS_CM / 100.0 <= radius_m + 1.0e-12
    for index, radius_cm in enumerate(OUTPUT_RADIUS_CM):
        if valid[index]:
            target = min(1.0, radius_cm / 100.0 / radius_m)
            fixed_T[index] = np.interp(target, xi, temperature_C)
            fixed_C[index] = np.interp(target, xi, moisture_kgkg)
    return fixed_T, fixed_C, valid


def _surface_values(temperature_C: np.ndarray, moisture_kgkg: np.ndarray) -> tuple[float, float]:
    return float(temperature_C[-1]), float(moisture_kgkg[-1])


def write_result_xlsx(solution: dict[str, Any], output_path: str | Path = DEFAULT_XLSX_PATH, template_path: str | Path = DEFAULT_TEMPLATE_PATH) -> None:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(Path(template_path))
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError("result4模板工作表错误")
    sheet = workbook["Sheet1"]
    header_style = copy(sheet.cell(1, 2)._style)
    time_style = copy(sheet.cell(2, 1)._style)
    value_style = copy(sheet.cell(2, 2)._style)
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    sheet.cell(1, 1).value = "时间\\到药材中心的距离"
    for column, radius in enumerate(OUTPUT_RADIUS_CM, start=2):
        cell = sheet.cell(1, column, float(radius)); cell._style = copy(header_style); cell.number_format = "0.0"
    cell = sheet.cell(1, OUTPUT_RADIUS_CM.size + 2, "药材表面"); cell._style = copy(header_style)
    xi = solution["xi"]
    indices = _regular_sample_indices(solution)
    times = solution["sampled_times_s"]
    radii = solution["sampled_radius_m"]
    states_T = solution["sampled_temperature_internal_C"]
    states_C = solution["sampled_moisture_internal_kgkg"]
    for row, state_index in enumerate(indices, start=2):
        cell = sheet.cell(row, 1, int(round(times[state_index]))); cell._style = copy(time_style); cell.number_format = "0"
        _, fixed_C, valid = _fixed_state_values(xi, states_T[state_index], states_C[state_index], float(radii[state_index]))
        for column, (value, is_valid) in enumerate(zip(fixed_C, valid), start=2):
            cell = sheet.cell(row, column); cell._style = copy(value_style); cell.value = float(f"{value:.4f}") if is_valid else None; cell.number_format = "0.0000"
        cell = sheet.cell(row, OUTPUT_RADIUS_CM.size + 2, float(f"{states_C[state_index, -1]:.4f}")); cell._style = copy(value_style); cell.number_format = "0.0000"
    workbook.save(output_path)


def _csv_state_rows(time_s: float, radius_m: float, temperature_C: np.ndarray, moisture_kgkg: np.ndarray, xi: np.ndarray, sample_time: bool) -> list[dict[str, str]]:
    fixed_T, fixed_C, valid = _fixed_state_values(xi, temperature_C, moisture_kgkg, radius_m)
    token = str(int(round(time_s))) if sample_time else f"{time_s:.6f}"
    fixed_type = "sample_fixed" if sample_time else "drying_end_fixed"
    surface_type = "sample_surface" if sample_time else "drying_end_surface"
    radius_token = f"{radius_m * 100.0:.4f}"
    rows: list[dict[str, str]] = []
    for index, radius_cm in enumerate(OUTPUT_RADIUS_CM):
        if valid[index]:
            rows.append({"record_type": fixed_type, "time_s": token, "time_h": f"{time_s / 3600.0:.8f}", "radius_cm": radius_token, "r_cm": f"{radius_cm:.4f}", "is_surface": "0", "temperature_C": f"{fixed_T[index]:.4f}", "moisture_kgkg": f"{fixed_C[index]:.4f}"})
    surface_T, surface_C = _surface_values(temperature_C, moisture_kgkg)
    rows.append({"record_type": surface_type, "time_s": token, "time_h": f"{time_s / 3600.0:.8f}", "radius_cm": radius_token, "r_cm": radius_token, "is_surface": "1", "temperature_C": f"{surface_T:.4f}", "moisture_kgkg": f"{surface_C:.4f}"})
    return rows


def write_result_csv(solution: dict[str, Any], output_path: str | Path = DEFAULT_CSV_PATH) -> None:
    output_path = Path(output_path); output_path.parent.mkdir(parents=True, exist_ok=True)
    rows: list[dict[str, str]] = []
    indices = _regular_sample_indices(solution)
    for index in indices:
        rows.extend(_csv_state_rows(solution["sampled_times_s"][index], solution["sampled_radius_m"][index], solution["sampled_temperature_internal_C"][index], solution["sampled_moisture_internal_kgkg"][index], solution["xi"], True))
    rows.extend(_csv_state_rows(solution["termination_time_s"], solution["termination_radius_m"], solution["termination_temperature_internal_C"], solution["termination_moisture_internal_kgkg"], solution["xi"], False))
    with output_path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=CSV_COLUMNS); writer.writeheader(); writer.writerows(rows)


def _paper_rows(solution: dict[str, Any]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    xi = solution["xi"]
    for hour in _paper_regular_hours(solution):
        index = _sample_index_at_time(solution, hour * 3600.0)
        _, fixed_C, valid = _fixed_state_values(xi, solution["sampled_temperature_internal_C"][index], solution["sampled_moisture_internal_kgkg"][index], solution["sampled_radius_m"][index])
        rows.append([float(hour), fixed_C[0] if valid[0] else None, fixed_C[5] if valid[5] else None, fixed_C[10] if valid[10] else None, solution["sampled_moisture_internal_kgkg"][index, -1]])
    _, fixed_C, valid = _fixed_state_values(xi, solution["termination_temperature_internal_C"], solution["termination_moisture_internal_kgkg"], solution["termination_radius_m"])
    rows.append([f"烘干结束时间（{solution['termination_time_s'] / 3600.0:.4f} h）", fixed_C[0] if valid[0] else None, fixed_C[5] if valid[5] else None, fixed_C[10] if valid[10] else None, solution["termination_moisture_internal_kgkg"][-1]])
    return rows


def write_paper_table(solution: dict[str, Any], output_path: str | Path = DEFAULT_PAPER_PATH) -> None:
    output_path = Path(output_path); output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(); sheet = workbook.active; sheet.title = "表6"
    for column, header in enumerate(["时间/h", "0 cm", "0.5 cm", "1.0 cm", "药材表面"], start=1):
        sheet.cell(1, column, header).font = Font(bold=True)
    rows = _paper_rows(solution)
    for row_index, values in enumerate(rows, start=2):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            if column == 1 and isinstance(value, (int, float)): cell.number_format = "0.0"
            elif column > 1 and value is not None: cell.value = float(f"{float(value):.4f}"); cell.number_format = "0.0000"
    summary = 2 + len(rows) + 1
    sheet.cell(summary, 1, "最终烘干时长/h")
    cell = sheet.cell(summary, 2, float(f"{solution['termination_time_s'] / 3600.0:.4f}")); cell.number_format = "0.0000"
    workbook.save(output_path)


def _fixed_valid_count(radius_m: float) -> int:
    return int(np.count_nonzero(OUTPUT_RADIUS_CM / 100.0 <= radius_m + 1.0e-12))


def _format_pattern_ok(value: str, decimals: int) -> bool:
    return re.fullmatch(rf"\d+\.\d{{{decimals}}}", value) is not None


def validate_written_outputs(solution: dict[str, Any], xlsx_path: str | Path = DEFAULT_XLSX_PATH, paper_path: str | Path = DEFAULT_PAPER_PATH, csv_path: str | Path = DEFAULT_CSV_PATH) -> dict[str, Any]:
    """独立回读三个文件；特别检查移动表面之外的固定位置为空。"""
    summary = validate_solution(solution)
    indices = _regular_sample_indices(solution); xi = solution["xi"]
    workbook = load_workbook(Path(xlsx_path), read_only=True, data_only=True); sheet = workbook["Sheet1"]; rows = list(sheet.iter_rows(values_only=True))
    if workbook.sheetnames != ["Sheet1"] or (sheet.max_row, sheet.max_column) != (indices.size + 1, OUTPUT_RADIUS_CM.size + 2): raise ValueError("result4结构错误")
    if rows[0][0] != "时间\\到药材中心的距离" or rows[0][-1] != "药材表面" or not np.allclose(np.asarray(rows[0][1:-1], dtype=float), OUTPUT_RADIUS_CM): raise ValueError("result4表头错误")
    xlsx_fixed: dict[tuple[int, str], float] = {}; xlsx_surface: dict[int, float] = {}; blank_count = 0
    for row_offset, state_index in enumerate(indices, start=1):
        time_value = int(round(solution["sampled_times_s"][state_index]))
        _, fixed_C, valid = _fixed_state_values(xi, solution["sampled_temperature_internal_C"][state_index], solution["sampled_moisture_internal_kgkg"][state_index], solution["sampled_radius_m"][state_index])
        if rows[row_offset][0] != time_value: raise ValueError("result4时间错误")
        for index, is_valid in enumerate(valid):
            value = rows[row_offset][index + 1]
            if not is_valid:
                blank_count += 1
                if value is not None: raise ValueError("result4域外位置存在外推值")
            else:
                expected = float(f"{fixed_C[index]:.4f}")
                if value is None or float(value) != expected: raise ValueError("result4固定值不一致")
                xlsx_fixed[(time_value, f"{OUTPUT_RADIUS_CM[index]:.4f}")] = float(value)
        expected_surface = float(f"{solution['sampled_moisture_internal_kgkg'][state_index, -1]:.4f}")
        if float(rows[row_offset][-1]) != expected_surface: raise ValueError("result4表面值不一致")
        xlsx_surface[time_value] = float(rows[row_offset][-1])
    with Path(csv_path).open("r", newline="", encoding="utf-8") as stream:
        csv_rows = list(csv.DictReader(stream))
    if not csv_rows or list(csv_rows[0]) != CSV_COLUMNS:
        raise ValueError("CSV字段错误")
    expected_count = sum(_fixed_valid_count(solution["sampled_radius_m"][index]) + 1 for index in indices) + _fixed_valid_count(solution["termination_radius_m"]) + 1
    if len(csv_rows) != expected_count:
        raise ValueError("CSV行数错误")
    expected_record_counts = {
        "sample_fixed": int(sum(_fixed_valid_count(solution["sampled_radius_m"][index]) for index in indices)),
        "sample_surface": int(indices.size),
        "drying_end_fixed": _fixed_valid_count(solution["termination_radius_m"]),
        "drying_end_surface": 1,
    }
    record_counts = {record_type: sum(row["record_type"] == record_type for row in csv_rows) for record_type in expected_record_counts}
    if record_counts != expected_record_counts:
        raise ValueError(f"CSV记录类型数量错误: {record_counts}")
    allowed_record_types = set(expected_record_counts)
    sample_last_time = int(np.floor(float(solution["termination_time_s"]) / OUTPUT_DT_S + 1.0e-12) * OUTPUT_DT_S)
    termination_time = float(solution["termination_time_s"])
    termination_radius_cm = float(solution["termination_radius_m"]) * 100.0
    seen_records: set[tuple[str, str, str]] = set()
    _, termination_fixed_C, termination_valid = _fixed_state_values(
        xi,
        solution["termination_temperature_internal_C"],
        solution["termination_moisture_internal_kgkg"],
        solution["termination_radius_m"],
    )
    for row in csv_rows:
        record_type = row["record_type"]
        if record_type not in allowed_record_types:
            raise ValueError(f"CSV包含未知record_type: {record_type}")
        expected_surface_flag = "1" if record_type.endswith("_surface") else "0"
        if row["is_surface"] != expected_surface_flag:
            raise ValueError("CSV record_type与is_surface不一致")
        for field, decimals in (("time_h", 8), ("radius_cm", 4), ("r_cm", 4), ("temperature_C", 4), ("moisture_kgkg", 4)):
            if not _format_pattern_ok(row[field], decimals):
                raise ValueError(f"CSV字段{field}格式错误")
        time_s_value = float(row["time_s"])
        if not np.isclose(float(row["time_h"]), time_s_value / 3600.0, atol=5.0e-9, rtol=0.0):
            raise ValueError("CSV time_h与time_s不一致")
        value = float(row["moisture_kgkg"])
        if value < 0.0 or not np.isfinite(value):
            raise ValueError("CSV水分无效")
        record_key = (record_type, row["time_s"], row["r_cm"])
        if record_key in seen_records:
            raise ValueError("CSV存在重复记录")
        seen_records.add(record_key)
        if record_type.startswith("sample_"):
            if re.fullmatch(r"\d+", row["time_s"]) is None:
                raise ValueError("CSV采样time_s必须为整数秒")
            time_value = int(row["time_s"])
            if time_value < int(OUTPUT_DT_S) or time_value > sample_last_time or time_value % int(OUTPUT_DT_S) != 0:
                raise ValueError("CSV采样time_s不是严格60 s网格")
            matches = np.flatnonzero(np.isclose(solution["sampled_times_s"], time_value, atol=1.0e-8, rtol=0.0))
            if matches.size != 1:
                raise ValueError("CSV采样时间无法对应求解状态")
            state_index = int(matches[0])
            expected_sample_radius_cm = float(solution["sampled_radius_m"][state_index]) * 100.0
            if not np.isclose(float(row["radius_cm"]), expected_sample_radius_cm, atol=5.0e-5, rtol=0.0):
                raise ValueError("CSV采样radius_cm与R(t)不一致")
            if record_type == "sample_surface":
                if not np.isclose(float(row["r_cm"]), expected_sample_radius_cm, atol=5.0e-5, rtol=0.0):
                    raise ValueError("CSV采样表面r_cm与R(t)不一致")
                expected = xlsx_surface[time_value]
            else:
                radius_index = int(round(float(row["r_cm"]) * 10.0))
                if radius_index < 0 or radius_index >= OUTPUT_RADIUS_CM.size or not np.isclose(float(row["r_cm"]), OUTPUT_RADIUS_CM[radius_index], atol=5.0e-5, rtol=0.0):
                    raise ValueError("CSV采样固定r_cm格式或范围错误")
                _, _, valid = _fixed_state_values(
                    xi,
                    solution["sampled_temperature_internal_C"][state_index],
                    solution["sampled_moisture_internal_kgkg"][state_index],
                    solution["sampled_radius_m"][state_index],
                )
                if not valid[radius_index]:
                    raise ValueError("CSV采样固定位置位于移动表面之外")
                expected = xlsx_fixed[(time_value, row["r_cm"])]
        else:
            if re.fullmatch(r"\d+\.\d{6}", row["time_s"]) is None:
                raise ValueError("CSV终止time_s必须为六位小数")
            if not np.isclose(time_s_value, termination_time, atol=5.0e-7, rtol=0.0):
                raise ValueError("CSV终止time_s与solution不一致")
            if not np.isclose(float(row["radius_cm"]), termination_radius_cm, atol=5.0e-5, rtol=0.0):
                raise ValueError("CSV终止radius_cm与solution不一致")
            if record_type == "drying_end_surface":
                if not np.isclose(float(row["r_cm"]), termination_radius_cm, atol=5.0e-5, rtol=0.0):
                    raise ValueError("CSV终止表面r_cm与真实表面不一致")
                expected = float(f"{solution['termination_moisture_internal_kgkg'][-1]:.4f}")
            else:
                radius_index = int(round(float(row["r_cm"]) * 10.0))
                if radius_index < 0 or radius_index >= OUTPUT_RADIUS_CM.size or not np.isclose(float(row["r_cm"]), OUTPUT_RADIUS_CM[radius_index], atol=5.0e-5, rtol=0.0) or not termination_valid[radius_index]:
                    raise ValueError("CSV终止固定位置域外或r_cm错误")
                expected = float(f"{termination_fixed_C[radius_index]:.4f}")
        if value != expected:
            raise ValueError("CSV与result4值不一致")
    paper_workbook = load_workbook(Path(paper_path), read_only=True, data_only=True); paper = paper_workbook["表6"]; paper_rows = list(paper.iter_rows(values_only=True)); hours = _paper_regular_hours(solution)
    if paper_workbook.sheetnames != ["表6"] or (paper.max_row, paper.max_column) != (len(hours) + 4, 5): raise ValueError("论文表结构错误")
    if list(paper_rows[0]) != ["时间/h", "0 cm", "0.5 cm", "1.0 cm", "药材表面"]: raise ValueError("论文表表头错误")
    for row_index, hour in enumerate(hours, start=1):
        index = _sample_index_at_time(solution, hour * 3600.0); _, fixed_C, valid = _fixed_state_values(xi, solution["sampled_temperature_internal_C"][index], solution["sampled_moisture_internal_kgkg"][index], solution["sampled_radius_m"][index])
        expected_values = [fixed_C[0], fixed_C[5], fixed_C[10], solution["sampled_moisture_internal_kgkg"][index, -1]]
        for column, expected in enumerate(expected_values, start=1):
            if column == 4 or valid[[0, 5, 10][column - 1]]:
                if float(paper_rows[row_index][column]) != float(f"{expected:.4f}"): raise ValueError("论文表常规值不一致")
    termination_row = paper_rows[len(hours) + 1]
    if termination_row[0] != f"烘干结束时间（{solution['termination_time_s'] / 3600.0:.4f} h）":
        raise ValueError("论文表终止时间标签错误")
    termination_expected = [termination_fixed_C[0], termination_fixed_C[5], termination_fixed_C[10], solution["termination_moisture_internal_kgkg"][-1]]
    for column, expected in enumerate(termination_expected, start=1):
        radius_index = (0, 5, 10, None)[column - 1]
        if radius_index is not None and not termination_valid[radius_index]:
            if termination_row[column] is not None:
                raise ValueError("论文表终止固定位置域外不应有数值")
            continue
        if termination_row[column] is None or float(termination_row[column]) != float(f"{expected:.4f}"):
            raise ValueError("论文表终止行水分值不一致")
    summary_row = paper_rows[-1]
    expected_duration_h = float(f"{termination_time / 3600.0:.4f}")
    if summary_row[0] != "最终烘干时长/h":
        raise ValueError("论文表最终时长标签错误")
    if summary_row[1] is None or float(summary_row[1]) != expected_duration_h:
        raise ValueError("论文表最终时长数值不一致")
    return {**summary, "xlsx_rows": sheet.max_row, "xlsx_columns": sheet.max_column, "csv_rows": len(csv_rows), "paper_rows": paper.max_row, "paper_columns": paper.max_column, "fixed_blank_count": blank_count, "csv_record_counts": record_counts}


def diagnostics_summary(solution: dict[str, Any]) -> dict[str, Any]:
    """Summarize Picard, physical-constraint, and conservation diagnostics."""
    step_times = np.asarray(solution["step_times_s"], dtype=float)
    step_dt = np.asarray(solution["step_dt_s"], dtype=float)
    iterations = np.asarray(solution["iteration_counts"], dtype=int)
    delta_temperature = np.asarray(solution["picard_delta_temperature"], dtype=float)
    delta_moisture = np.asarray(solution["picard_delta_moisture"], dtype=float)
    temperature_linear = np.asarray(solution["temperature_linear_residuals"], dtype=float)
    moisture_linear = np.asarray(solution["moisture_linear_residuals"], dtype=float)
    linear_residuals = np.asarray(solution["linear_residuals"], dtype=float)
    temperature_balance = np.asarray(solution["temperature_balance_residuals"], dtype=float)
    moisture_balance = np.asarray(solution["moisture_balance_residuals"], dtype=float)
    temperature_flux = np.asarray(solution["temperature_boundary_fluxes"], dtype=float)
    moisture_flux = np.asarray(solution["moisture_boundary_fluxes"], dtype=float)
    midpoint_radius = np.asarray(solution["step_midpoint_radius_m"], dtype=float)
    midpoint_air_temperature = np.asarray(
        solution["step_midpoint_air_temperature_C"], dtype=float
    )
    midpoint_air_moisture = np.asarray(
        solution["step_midpoint_air_moisture_kgkg"], dtype=float
    )
    if step_times.size == 0:
        raise ValueError("缺少时间步诊断记录")

    def _relative(residuals: np.ndarray, fluxes: np.ndarray) -> tuple[float, float]:
        # Use one characteristic boundary-flux scale for the whole run.  A
        # per-step denominator would report a misleading relative value of 1
        # whenever both the flux and roundoff residual are near zero.
        scale = max(
            float(np.max(np.abs(fluxes))),
            float(np.max(np.abs(residuals))),
            np.finfo(float).tiny,
        )
        relative = np.abs(residuals) / scale
        index = int(np.argmax(np.abs(residuals)))
        return float(np.max(relative)), float(step_times[index])

    moisture_relative, moisture_worst_time = _relative(
        moisture_balance, moisture_flux
    )
    temperature_relative, temperature_worst_time = _relative(
        temperature_balance, temperature_flux
    )
    temperature_driver = temperature_flux * midpoint_radius / HEAT_TRANSFER_W_M2_K
    moisture_driver = moisture_flux * midpoint_radius / MASS_TRANSFER_M_S
    radius_increments = np.asarray(solution["radius_increments_m"], dtype=float)
    temperature_bounds = tuple(float(v) for v in solution["temperature_bounds_C"])
    return {
        "task": 4,
        "caliber": {
            "coordinate": "材料坐标 xi=r/R(t)",
            "spatial_scheme": "守恒型径向有限体积（中心半控制体）",
            "temporal_scheme": "Crank-Nicolson",
            "interface_property": "调和平均",
            "nonlinear_scheme": "同步中点物性 Picard",
            "n_intervals": int(solution["n_intervals"]),
            "dt_early_s": float(solution["time_step_early_s"]),
            "dt_long_s": float(solution["time_step_long_s"]),
            "output_dt_s": float(solution["output_dt_s"]),
            "event_tol_s": float(solution["event_tol_s"]),
            "appendix": int(solution["appendix"]),
            "step_count": int(iterations.size),
            "sample_count": int(np.asarray(solution["sampled_times_s"]).size),
            "picard_tol_temperature_C": float(PICARD_TOL_T_C),
            "picard_tol_moisture_kgkg": float(PICARD_TOL_C_KGKG),
            "linear_residual_tol": float(LINEAR_RESIDUAL_TOL),
        },
        "input_sources": solution["input_sources"],
        "picard": {
            "max_iterations": int(np.max(iterations)),
            "mean_iterations": float(np.mean(iterations)),
            "max_final_delta_temperature_C": float(np.max(delta_temperature)),
            "max_final_delta_moisture_kgkg": float(np.max(delta_moisture)),
            "converged_steps": int(
                np.sum(
                    (delta_temperature <= PICARD_TOL_T_C)
                    & (delta_moisture <= PICARD_TOL_C_KGKG)
                )
            ),
            "failed_steps": int(
                np.sum(
                    (delta_temperature > PICARD_TOL_T_C)
                    | (delta_moisture > PICARD_TOL_C_KGKG)
                )
            ),
            "event_candidate_count": int(
                solution.get("termination_event_candidate_count", 0)
            ),
            "event_max_iterations": int(
                solution.get("termination_event_picard_max", 0)
            ),
        },
        "linear_solver": {
            "max_temperature_normalized_inf_residual": float(
                np.max(temperature_linear)
            ),
            "max_moisture_normalized_inf_residual": float(
                np.max(moisture_linear)
            ),
            "max_normalized_inf_residual": float(np.max(linear_residuals)),
            "tol": float(LINEAR_RESIDUAL_TOL),
        },
        "physics": {
            "finite": True,
            "temperature_min_C": float(solution["min_temperature_C"]),
            "temperature_max_C": float(solution["max_temperature_C"]),
            "temperature_bounds_C": list(temperature_bounds),
            "temperature_range_ok": bool(
                solution["min_temperature_C"] >= temperature_bounds[0] - STATE_TOL
                and solution["max_temperature_C"] <= temperature_bounds[1] + STATE_TOL
            ),
            "moisture_min_kgkg": float(solution["min_internal_moisture_kgkg"]),
            "moisture_max_kgkg": float(solution["max_internal_moisture_kgkg"]),
            "moisture_nonnegative": bool(
                solution["min_internal_moisture_kgkg"] >= -STATE_TOL
            ),
            "moisture_bounded_by_initial": bool(
                solution["max_internal_moisture_kgkg"]
                <= INITIAL_MOISTURE_KGKG + STATE_TOL
            ),
            "center_boundary_zero_flux": True,
            "radius_min_m": float(np.min(np.asarray(solution["step_radius_m"]))),
            "radius_max_m": float(np.max(np.asarray(solution["step_radius_m"]))),
            "radius_monotone_nonincreasing": bool(
                np.all(radius_increments <= 1.0e-12)
            ),
            "max_radius_increase_m": float(np.max(radius_increments)),
            "temperature_boundary_driver_min_C": float(np.min(temperature_driver)),
            "temperature_inward_fraction": float(
                np.mean(temperature_driver >= -STATE_TOL)
            ),
            "moisture_boundary_driver_max_kgkg": float(np.max(moisture_driver)),
            "moisture_outward_fraction": float(
                np.mean(moisture_driver <= STATE_TOL)
            ),
        },
        "conservation": {
            "coordinate_measure": "Σ m_i C_i / Σ m_i and Σ m_i (rho cp)_mid T_i; 2πL canceled",
            "moisture": {
                "max_abs_balance_residual": float(np.max(np.abs(moisture_balance))),
                "max_relative_balance_residual": moisture_relative,
                "worst_time_s": moisture_worst_time,
                "boundary_flux_integral": float(np.sum(moisture_flux * step_dt)),
                "event_max_abs_balance_residual": float(
                    solution.get("termination_event_moisture_balance_max_abs", 0.0)
                ),
            },
            "energy": {
                "max_abs_balance_residual": float(
                    np.max(np.abs(temperature_balance))
                ),
                "max_relative_balance_residual": temperature_relative,
                "worst_time_s": temperature_worst_time,
                "boundary_flux_integral": float(np.sum(temperature_flux * step_dt)),
                "event_max_abs_balance_residual": float(
                    solution.get("termination_event_temperature_balance_max_abs", 0.0)
                ),
            },
        },
        "termination": {
            "time_s": float(solution["termination_time_s"]),
            "lower_time_s": float(solution["termination_lower_time_s"]),
            "upper_max_kgkg": float(solution["termination_max_kgkg"]),
            "lower_max_kgkg": float(solution["termination_lower_max_kgkg"]),
            "bracket_width_s": float(solution["termination_bracket_width_s"]),
            "max_xi": float(solution["termination_max_xi"]),
            "max_radius_m": float(solution["termination_max_radius_m"]),
        },
    }


def write_diagnostics(
    solution: dict[str, Any],
    json_path: str | Path = DEFAULT_DIAGNOSTICS_JSON,
    csv_path: str | Path = DEFAULT_DIAGNOSTICS_CSV,
) -> tuple[Path, Path]:
    """Persist per-step Picard, physics, and conservation logs."""
    summary = diagnostics_summary(solution)
    json_destination = Path(json_path)
    json_destination.parent.mkdir(parents=True, exist_ok=True)
    json_destination.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    step_times = np.asarray(solution["step_times_s"], dtype=float)
    step_dt = np.asarray(solution["step_dt_s"], dtype=float)
    midpoint_radius = np.asarray(solution["step_midpoint_radius_m"], dtype=float)
    midpoint_air_temperature = np.asarray(
        solution["step_midpoint_air_temperature_C"], dtype=float
    )
    midpoint_air_moisture = np.asarray(
        solution["step_midpoint_air_moisture_kgkg"], dtype=float
    )
    iterations = np.asarray(solution["iteration_counts"], dtype=int)
    delta_temperature = np.asarray(solution["picard_delta_temperature"], dtype=float)
    delta_moisture = np.asarray(solution["picard_delta_moisture"], dtype=float)
    temperature_linear = np.asarray(solution["temperature_linear_residuals"], dtype=float)
    moisture_linear = np.asarray(solution["moisture_linear_residuals"], dtype=float)
    linear_residuals = np.asarray(solution["linear_residuals"], dtype=float)
    temperature_balance = np.asarray(solution["temperature_balance_residuals"], dtype=float)
    moisture_balance = np.asarray(solution["moisture_balance_residuals"], dtype=float)
    temperature_flux = np.asarray(solution["temperature_boundary_fluxes"], dtype=float)
    moisture_flux = np.asarray(solution["moisture_boundary_fluxes"], dtype=float)
    step_max = np.asarray(solution["step_max_kgkg"], dtype=float)
    step_xi = np.asarray(solution["step_max_xi"], dtype=float)
    step_position = np.asarray(solution["step_max_radius_m"], dtype=float)
    radius = np.asarray(solution["step_radius_m"], dtype=float)
    csv_destination = Path(csv_path)
    csv_destination.parent.mkdir(parents=True, exist_ok=True)
    with csv_destination.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(
            (
                "step",
                "time_s",
                "dt_s",
                "midpoint_radius_m",
                "radius_m",
                "midpoint_air_temperature_C",
                "midpoint_air_moisture_kgkg",
                "picard_iterations",
                "picard_delta_temperature_C",
                "picard_delta_moisture_kgkg",
                "temperature_linear_normalized_inf_residual",
                "moisture_linear_normalized_inf_residual",
                "linear_normalized_inf_residual",
                "temperature_balance_residual",
                "moisture_balance_residual",
                "temperature_boundary_flux",
                "moisture_boundary_flux",
                "max_moisture_kgkg",
                "max_xi",
                "max_radius_m",
            )
        )
        for index in range(iterations.size):
            writer.writerow(
                (
                    index + 1,
                    f"{step_times[index]:.6f}",
                    f"{step_dt[index]:.6f}",
                    f"{midpoint_radius[index]:.12e}",
                    f"{radius[index]:.12e}",
                    f"{midpoint_air_temperature[index]:.12e}",
                    f"{midpoint_air_moisture[index]:.12e}",
                    int(iterations[index]),
                    f"{delta_temperature[index]:.6e}",
                    f"{delta_moisture[index]:.6e}",
                    f"{temperature_linear[index]:.6e}",
                    f"{moisture_linear[index]:.6e}",
                    f"{linear_residuals[index]:.6e}",
                    f"{temperature_balance[index]:.6e}",
                    f"{moisture_balance[index]:.6e}",
                    f"{temperature_flux[index]:.6e}",
                    f"{moisture_flux[index]:.6e}",
                    f"{step_max[index]:.12e}",
                    f"{step_xi[index]:.12e}",
                    f"{step_position[index]:.12e}",
                )
            )
    return json_destination, csv_destination


def mechanism_summary(mechanisms: dict[str, Any]) -> dict[str, Any]:
    """Create a JSON-safe A/B/C mechanism split with provenance."""
    cases: dict[str, Any] = {}
    for label in ("A", "B", "C"):
        solution = mechanisms[label]
        summary = validate_solution(solution)
        cases[label] = {
            "description": {
                "A": "固定 R=2 cm + 题面附录3物性",
                "B": "固定 R=2 cm + 题面附录4物性",
                "C": "附件2 R(t) + 题面附录4物性",
            }[label],
            "termination_time_s": float(summary["termination_time_s"]),
            "termination_time_h": float(summary["termination_time_h"]),
            "termination_bracket_width_s": float(summary["termination_bracket_width_s"]),
            "termination_max_kgkg": float(summary["termination_max_kgkg"]),
            "input_sources": solution["input_sources"],
            "configuration": {
                "n_intervals": int(solution["n_intervals"]),
                "time_step_early_s": float(solution["time_step_early_s"]),
                "time_step_long_s": float(solution["time_step_long_s"]),
                "output_dt_s": float(solution["output_dt_s"]),
                "event_tol_s": float(solution["event_tol_s"]),
                "appendix": int(solution["appendix"]),
                "fixed_radius_m": solution.get("fixed_radius_m"),
            },
        }
    A_time = cases["A"]["termination_time_s"]
    B_time = cases["B"]["termination_time_s"]
    C_time = cases["C"]["termination_time_s"]
    return {
        "contract": "同一N=800、1/5 s、同步Picard-CN和事件二分平台；只改变物性公式或半径",
        "common_input_sources": {
            "air_attachment": cases["A"]["input_sources"]["air_attachment"],
            "property_source": "题面附录3/4公式；唯一实现位置 solve_task4.py::material_properties",
        },
        "cases": cases,
        "split": {
            "property_effect_B_minus_A_s": float(B_time - A_time),
            "property_effect_B_minus_A_h": float((B_time - A_time) / 3600.0),
            "shrink_effect_C_minus_B_s": float(C_time - B_time),
            "shrink_effect_C_minus_B_h": float((C_time - B_time) / 3600.0),
        },
    }


def write_mechanism_report(
    mechanisms: dict[str, Any], path: str | Path = DEFAULT_MECHANISM_JSON
) -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(mechanism_summary(mechanisms), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return destination


def compare_sensitivity_tables(base: dict[str, Any], other: dict[str, Any]) -> dict[str, Any]:
    """仅比较共同真实时间/绝对半径的常规场，终止时间单独报告。"""
    base_times = np.asarray(base["sampled_times_s"]); other_times = np.asarray(other["sampled_times_s"])
    common = sorted(set(np.round(base_times, 8)).intersection(set(np.round(other_times, 8))))
    max_difference = 0.0; max_location: tuple[float, float] | None = None; surface_difference = 0.0
    for time_value in common:
        i, j = _sample_index_at_time(base, time_value), _sample_index_at_time(other, time_value)
        r_base, r_other = base["sampled_radius_m"][i], other["sampled_radius_m"][j]
        for radius_cm in OUTPUT_RADIUS_CM:
            radius_m = radius_cm / 100.0
            if radius_m > min(r_base, r_other) + 1.0e-12: continue
            v_base = np.interp(min(1.0, radius_m / r_base), base["xi"], base["sampled_moisture_internal_kgkg"][i])
            v_other = np.interp(min(1.0, radius_m / r_other), other["xi"], other["sampled_moisture_internal_kgkg"][j])
            difference = abs(float(v_base - v_other))
            if difference > max_difference: max_difference, max_location = difference, (float(time_value), float(radius_cm))
        surface_difference = max(surface_difference, abs(float(base["sampled_moisture_internal_kgkg"][i, -1] - other["sampled_moisture_internal_kgkg"][j, -1])))
    return {"common_time_count": len(common), "common_max_abs_difference_kgkg": max_difference, "common_max_difference_time_s": None if max_location is None else max_location[0], "common_max_difference_radius_cm": None if max_location is None else max_location[1], "surface_max_abs_difference_kgkg": surface_difference, "termination_time_base_s": float(base["termination_time_s"]), "termination_time_other_s": float(other["termination_time_s"]), "termination_time_difference_s": float(other["termination_time_s"] - base["termination_time_s"]), "termination_time_difference_h": float(other["termination_time_s"] - base["termination_time_s"]) / 3600.0}


def run_sensitivity(base_solution: dict[str, Any], progress: bool = False) -> dict[str, Any]:
    common = {"air_data_path": DEFAULT_AIR_DATA_PATH, "radius_data_path": DEFAULT_RADIUS_DATA_PATH, "max_time_s": base_solution["max_time_s"], "output_dt_s": base_solution["output_dt_s"], "progress": progress, "event_tol_s": base_solution["event_tol_s"], "appendix": base_solution["appendix"], "fixed_radius_m": base_solution.get("fixed_radius_m")}
    n400 = solve_task4(n_intervals=400, time_step_early_s=base_solution["time_step_early_s"], time_step_long_s=base_solution["time_step_long_s"], **common)
    half = solve_task4(n_intervals=base_solution["n_intervals"], time_step_early_s=base_solution["time_step_early_s"] / 2.0, time_step_long_s=base_solution["time_step_long_s"] / 2.0, **common)
    validate_solution(n400); validate_solution(half)
    return {"space": compare_sensitivity_tables(base_solution, n400), "time": compare_sensitivity_tables(base_solution, half), "n400": n400, "dt_half": half}


def run_coarse_comparison(*args: Any, **kwargs: Any) -> dict[str, Any]:
    kwargs.setdefault("n_intervals", 400)
    return validate_solution(solve_task4(*args, **kwargs))


def run_mechanisms(base_solution: dict[str, Any], progress: bool = False) -> dict[str, Any]:
    common = {"air_data_path": DEFAULT_AIR_DATA_PATH, "radius_data_path": DEFAULT_RADIUS_DATA_PATH, "n_intervals": base_solution["n_intervals"], "time_step_early_s": base_solution["time_step_early_s"], "time_step_long_s": base_solution["time_step_long_s"], "output_dt_s": base_solution["output_dt_s"], "event_tol_s": base_solution["event_tol_s"], "max_time_s": FIXED_MAX_TIME_S, "progress": progress, "fixed_radius_m": RADIUS_INITIAL_M}
    A = solve_task4(appendix=3, **common); B = solve_task4(appendix=4, **common)
    validate_solution(A); validate_solution(B)
    return {"A": A, "B": B, "C": base_solution}


def main() -> None:
    parser = argparse.ArgumentParser(description="求解A题问题四材料坐标移动边界模型")
    parser.add_argument("--air-data-path", type=Path, default=DEFAULT_AIR_DATA_PATH)
    parser.add_argument("--radius-data-path", type=Path, default=DEFAULT_RADIUS_DATA_PATH)
    parser.add_argument("--max-time", type=float, default=None)
    parser.add_argument("--time-step", type=float, default=None, help="兼容参数：同时设置两段步长")
    parser.add_argument("--time-step-early", type=float, default=None)
    parser.add_argument("--time-step-long", type=float, default=None)
    parser.add_argument("--event-tol", type=float, default=EVENT_TOL_S)
    parser.add_argument("--output-dt", type=float, default=OUTPUT_DT_S)
    parser.add_argument("--n", type=int, default=N_DEFAULT)
    parser.add_argument("--appendix", type=int, choices=(3, 4), default=4)
    parser.add_argument("--fixed-radius", type=float, default=None, help="固定半径/m，仅用于等价核对或机制")
    parser.add_argument("--no-write", action="store_true")
    parser.add_argument("--skip-sensitivity", action="store_true")
    parser.add_argument("--mechanism", action="store_true")
    parser.add_argument("--progress", action="store_true")
    args = parser.parse_args()
    started = wall_time.perf_counter()
    solution = solve_task4(air_data_path=args.air_data_path, radius_data_path=args.radius_data_path, max_time_s=args.max_time, time_step_s=args.time_step, output_dt_s=args.output_dt, n_intervals=args.n, progress=args.progress, time_step_early_s=args.time_step_early, time_step_long_s=args.time_step_long, event_tol_s=args.event_tol, appendix=args.appendix, fixed_radius_m=args.fixed_radius)
    summary = validate_solution(solution)
    print(f"问题四材料坐标CN求解: N={args.n}, 早/晚dt={solution['time_step_early_s']:g}/{solution['time_step_long_s']:g}s, event_tol={args.event_tol:g}s")
    print(f"稳定段均值: T_air={summary['stable_temperature_C']:.12f} C, C_air={summary['stable_moisture_kgkg']:.12f} kg/kg ({summary['stable_count']}点)")
    print(f"终止证据: lower t={summary['termination_lower_time_s']:.9f}s max C={summary['termination_lower_max_kgkg']:.12f}; upper t={summary['termination_time_s']:.9f}s max C={summary['termination_max_kgkg']:.12f}; R={summary['termination_radius_cm']:.9f}cm, r_max={summary['termination_max_radius_cm']:.9f}cm")
    print(f"数值检查: {summary}")
    if not args.skip_sensitivity and args.fixed_radius is None:
        sensitivity = run_sensitivity(solution)
        print(f"空间敏感性N400: {sensitivity['space']}")
        print(f"时间敏感性早/晚步长减半: {sensitivity['time']}")
    elif args.skip_sensitivity:
        print("已跳过敏感性检查")
    if args.mechanism:
        mechanisms = run_mechanisms(solution)
        mechanism_path = write_mechanism_report(mechanisms)
        mechanism_report = mechanism_summary(mechanisms)
        print(
            "机制A/B/C时长(s): "
            f"{mechanism_report['cases']['A']['termination_time_s']:.9f}, "
            f"{mechanism_report['cases']['B']['termination_time_s']:.9f}, "
            f"{mechanism_report['cases']['C']['termination_time_s']:.9f}"
        )
        print(
            "机制分解: "
            f"物性(B-A)={mechanism_report['split']['property_effect_B_minus_A_s']:.9f}s, "
            f"收缩(C-B)={mechanism_report['split']['shrink_effect_C_minus_B_s']:.9f}s"
        )
        print(f"机制输入来源摘要: {mechanism_path}")
    if args.no_write:
        print("--no-write: 未生成或覆盖answer文件")
    else:
        template_hash = hashlib.sha256(DEFAULT_TEMPLATE_PATH.read_bytes()).hexdigest()
        write_result_xlsx(solution); write_result_csv(solution); write_paper_table(solution)
        output_summary = validate_written_outputs(solution)
        if hashlib.sha256(DEFAULT_TEMPLATE_PATH.read_bytes()).hexdigest() != template_hash: raise RuntimeError("result4模板被修改")
        print(f"答案文件回读验收: {output_summary}")
        diagnostics_json, diagnostics_csv = write_diagnostics(solution)
        diagnostics = diagnostics_summary(solution)
        print(f"诊断日志: {diagnostics_json}, {diagnostics_csv}")
        print(
            "Picard/守恒: "
            f"max/mean={diagnostics['picard']['max_iterations']}/"
            f"{diagnostics['picard']['mean_iterations']:.6f}, "
            f"max线性残差={diagnostics['linear_solver']['max_normalized_inf_residual']:.3e}, "
            f"max水分守恒残差={diagnostics['conservation']['moisture']['max_abs_balance_residual']:.3e}, "
            f"max能量守恒残差={diagnostics['conservation']['energy']['max_abs_balance_residual']:.3e}"
        )
    print(f"烘干结束时间={summary['termination_time_h']:.12f} h")
    print(f"总耗时: {wall_time.perf_counter() - started:.2f}s")


if __name__ == "__main__":
    main()
